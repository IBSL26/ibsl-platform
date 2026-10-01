"""Add the Professional insights block to the presenter notes of a unit deck.

Usage: python3 enrich.py <deck.pptx> <enrich/uNN.json>
Run after building a deck with unitNN.js. Safe to re-run: an existing
"Professional insights:" block is removed and written again.
Slide text is not touched.
"""
import json, sys
from pptx import Presentation
from pptx.oxml.ns import qn

MARK = "Professional insights:"

def shape_text(sh):
    """Read a shape's text without creating an empty text body (keeps slide XML untouched)."""
    tb = sh._element.find(qn("p:txBody"))
    if tb is None:
        return ""
    return "\n".join("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in tb.iter(qn("a:p")))

def main(deck, spec):
    d = json.load(open(spec, encoding="utf-8"))
    chk = d.pop("_check", {})
    p = Presentation(deck)
    slides = list(p.slides)
    for key, entries in d.items():
        s = slides[int(key) - 1]
        text = " | ".join(t for t in (shape_text(sh) for sh in s.shapes) if t.strip())
        want = chk.get(key, "")
        if want and not text.replace("\n", " / ").startswith(want):
            sys.exit(f"Slide {key} does not match: expected '{want}', found '{text[:40]}'")
        tf = s.notes_slide.notes_text_frame
        paras = list(tf.paragraphs)
        # remove an earlier block (from the blank line before the marker onwards)
        for i, para in enumerate(paras):
            if para.text.strip() == MARK:
                start = i - 1 if i > 0 and not paras[i - 1].text.strip() else i
                for q in paras[start:]:
                    q._p.getparent().remove(q._p)
                break
        lines = ["", MARK]
        for e in entries:
            lines += ["", f"{e['label']}: {e['text']}", f"Source: {e['source']}"]
        for line in lines:
            tf.add_paragraph().text = line
    p.save(deck)
    print(f"{deck}: {len(d)} slides enriched")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
