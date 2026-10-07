"""Final three-way scan: the Unit 4 deck against the two Unit 4 pages. Usage: python3 final_scan.py <deck> <p.html> <f.html>
1. Key learning outcomes and section outcomes word for word on both pages. 2. Every piece of slide text of 15 characters or more is on a page
(deck-only labels are listed in DECK_ONLY). 3. No slide text under 24pt; every notes page opens with a bold heading.
4. Leftovers of the old unit and of rules (Step 5, Save to Portfolio, times, "lens", old case) in pages, slides and notes. 5. The two pages carry the same part titles and sub-lines."""
import sys, re, html
from pptx import Presentation
from pptx.util import Pt
deck, P, F = sys.argv[1:4]
def norm(t): return re.sub(r'\s+([?.,;:])', r'\1', re.sub(r'\s+', ' ', t).replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"').replace('–', '-')).strip()
def plain(path, js=True):
    t = open(path, encoding='utf-8').read()
    t = re.sub(r'<style.*?</style>', ' ', t, flags=re.S)
    if not js: t = re.sub(r'<script.*?</script>', ' ', t, flags=re.S)
    return norm(html.unescape(re.sub(r'<[^>]+>', ' ', t)).replace("\\'", "'"))
tp, tf = plain(P), plain(F); vp, vf = plain(P, False), plain(F, False)
low = (tp + ' ' + tf).lower()
prs = Presentation(deck); bad = 0
DECK_ONLY = [r'^Key learning outcomes$', r'^The unit journey$', r'^Section \d', r'^Section learning outcomes$', r'Facilitator deck$', r'^Module 2 · Direction · Unit 4$',
  r'^Testing whether the chosen direction', r'^On a Key Result from Unit 3$', r'^Must be true:', r'^Role-based blind spots: the ABCV checkpoint each role most often overlooks$',
  r'^Unit Summary · The closing question$', r'^The case: The Future of a University\. Three rounds, then The Reveal\.$', r'^A client services company · from Unit 3, part 2\.2$',
  r'^Four failed strategies\. Defuse the mine before it detonates\.$', r'^\d\.\d[ab]? · ', r'^Overview · ', r'^The Checkpoint Reveals$', r'^Cause of Death · ', r'^Bringing ABCV Together$', r'^Direction Integrity',
  # the four working lines of every group-work slide (same wording as the Unit 2 and Unit 3 decks)
  r'^the group agrees each entry$', r'^one member types the agreed wording$', r'^every member types the agreed entries, in the session or after it$', r'^every Key Result, through the four steps$',
  # short forms of the step instructions (the full sentences are in the notes and on the pages)
  r'^which customer need the Key Result serves$', r'^what could slow, restrict or prevent delivery$', r'^who or what else can give the customer what it delivers$', r'^what about it will make customers choose you$',
  r'^Choose the investigation question the scenario answers: Familiarity, Exclusion or Complacency\.$', r'^Choose the ABCV checkpoint the illusion hid from leaders: Arena, Boundaries, Competition or Value Proposition\.$',
  r'^\d · (Find the illusion|Find where it sits|Define the Arena|Understand the Boundaries|See the Competition|Establish the Value Proposition)$']
print('1. Outcomes')
texts = []
for i, s in enumerate(prs.slides, 1):
    runs = []
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.font.size is not None and r.font.size < Pt(24): bad += 1; print(f'  slide {i}: {r.font.size.pt}pt "{r.text[:40]}"')
                if r.text.strip(): runs.append(norm(r.text))
    texts.append(runs)
    nt = s.notes_slide.notes_text_frame.paragraphs[0]
    if not (nt.runs and nt.runs[0].font.bold): bad += 1; print(f'  slide {i}: notes do not open with a bold heading')
def need(label, t, where='pf'):
    global bad
    miss = [w for w in where if norm(t) not in (tp if w == 'p' else tf)]
    if miss: bad += 1; print(f'  DIFFERS · {label} · not on {"/".join("participant" if m == "p" else "facilitator" for m in miss)} page: {t[:110]}')
for t in texts[1][1:]: need('KLO', t)
for i, sl in enumerate(texts, 1):
    if 'Section learning outcomes' in sl:
        k = sl.index('Section learning outcomes'); [need(f'section outcome, slide {i}', t) for t in sl[k + 1:k + 3]]
print('2. Slide text on the pages')
seen = 0
for i, sl in enumerate(texts, 1):
    for t in sl:
        if len(t) < 15 or any(re.search(p, t) for p in DECK_ONLY): continue
        seen += 1
        k = re.sub(r'^Objective: ', '', t).rstrip('.').lower()
        if k not in low and k.strip('"') not in low:
            bad += 1; print(f'  slide {i}: not on a page: {t[:120]}')
print('   strings checked:', seen)
print('3. Leftovers')
pats = [r'Step 5', r'[Ff]ive steps', r'Save to Portfolio', r'Go to Submit', r'\blens(es)?\b', r'Meridian', r'\b\d+\s*(–|-|to)?\s*\d*\s*(minutes|mins)\b', r'[Ww]hiteboard|flip chart', r'pre-?work',
        r'Leadership Ownership', r'Supporting Partners', r'[Cc]ritical conditions and', r'Strategy2Results(?!®)', r'\bS2R(?!®)', r'\bSIP\b']
for name, t in (('participant page', vp), ('facilitator page', vf), ('deck slides', ' | '.join(' | '.join(x) for x in texts)), ('deck notes', '\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides))):
    hits = [m.group(0) for pat in pats for m in re.finditer(r'.{0,40}' + pat + r'.{0,30}', t)]
    print(f'   {name}: {len(hits)}', hits[:6]); bad += len(hits)
print('4. Page pair')
for lab, rx in (('part titles', r'class="acc-t[^"]*">(.*?)<'), ('sub-lines', r'class="acc-meta">(.*?)<')):
    a = [html.unescape(x) for x in re.findall(rx, open(P, encoding='utf-8').read())]; b = [html.unescape(x) for x in re.findall(rx, open(F, encoding='utf-8').read())]
    a, b = a[:11], b[:11]   # the eleven parts of Sections 1 to 4; what follows is the Unit Summary and the portal's own panels
    same = a == b; print('   ', lab, 'identical on both pages:', same, '' if same else (a, b)); bad += (not same)
print('slides:', len(texts), '· problems:', bad)
