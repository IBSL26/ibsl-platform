"""Final scan: Carol's deck against the two Unit 3 pages. Usage: python3 -I final_scan.py <deck> <p.html> <f.html>"""
import sys, re, html
from pptx import Presentation
deck, P, F = sys.argv[1:4]
def plain(path, js=True):
    t = open(path, encoding='utf-8').read()
    if not js: t = re.sub(r'<script.*?</script>|<style.*?</style>', ' ', t, flags=re.S)
    else: t = re.sub(r'<style.*?</style>', ' ', t, flags=re.S)
    t = html.unescape(re.sub(r'<[^>]+>', ' ', t)).replace("\\'", "'")
    return re.sub(r'\s+', ' ', t).replace('’', "'")
tp, tf = plain(P), plain(F)
vp, vf = plain(P, False), plain(F, False)
prs = Presentation(deck)
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        else: yield sh
slides = []
for s in prs.slides:
    slides.append([re.sub(r'\s+', ' ', sh.text_frame.text.replace('\x0b', ' ')).strip().replace('’', "'") for sh in walk(s.shapes) if sh.has_text_frame and sh.text_frame.text.strip()])
notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
bad = 0
def need(label, text, where='pf'):
    global bad
    k = re.sub(r'\s+', ' ', text).replace('’', "'")
    miss = [w for w in where if k not in (tp if w == 'p' else tf)]
    if miss: bad += 1; print(f'  DIFFERS · {label} · not on {"/".join("participant" if m=="p" else "facilitator" for m in miss)} page: {text[:110]}')
print('1. Key learning outcomes (slide 2)');  [need('KLO', t) for t in slides[1][1:]]
print('2. Section outcomes (section slides)')
for i, sl in enumerate(slides, 1):
    if 'Section learning outcomes' in sl:
        k = sl.index('Section learning outcomes'); [need(f'slide {i}', t) for t in sl[k + 1:k + 3]]
print('3. Five step names');  [need('step', t) for t in ('Find Themes', 'Inspiring Objectives', 'Define Key Results', 'Alignment Test', 'Priority Matrix')]
print('4. Worked example')
for t in slides[14]:
    if t.startswith('We will be'): [need('SiP statement', x.strip()) for x in re.split(r'(?<=\.)\s+', t) if x.strip()]
kiss = [t for t in slides[15] if t not in ('KEEP', 'IMPROVE', 'START', 'STOP') and not t.isupper()]
[need('KISS entry', t) for t in kiss]
[need('theme', t) for t in slides[16] if t in ('Customer responsiveness', 'Execution speed', 'Collaboration culture', 'Value creation')]
print('   themes on slide 17:', [t for t in slides[16] if not t.startswith(('STEP', 'FROM', 'Reading')) and not t.isdigit()])
for t in slides[17]:
    m = re.match(r'(\d): (.+)', t)
    if m: need('Objective ' + m.group(1) + ' (slide 18)', m.group(2))
for t in slides[18]:
    if t.startswith(('Reduce', 'Increase', 'Improve')): need('Key Result (slide 19)', t)
print('   matrix on slide 21:', [t for t in slides[20] if re.match(r'\d · ', t)])
print('5. 3.1 content');  [need('3.1b risk', t) for t in slides[24][1:5]]; [need('CEO card', t) for t in slides[25][1:] if len(t) > 25]
print('6. Leftovers')
pats = [r'\bsix\b', r'Step 6', r'Less is More', r'Enterprise Priority\b', r'[Ii]ndividual work', r'\bSIP\b', r'Key account', r'\bvotes?\b', r'up to six', r'[Ss]ix-step', r'whiteboard|flip chart', r'\b\d+\s*(minutes|mins)\b']
for name, t in (('participant page', vp), ('facilitator page', vf), ('deck slides', ' | '.join(' | '.join(x) for x in slides)), ('deck notes', '\n'.join(notes))):
    hits = []
    for pat in pats:
        hits += [m.group(0) for m in re.finditer(r'.{0,40}' + pat + r'.{0,30}', t)]
    print(f'   {name}: {len(hits)}', hits[:6])
    bad += len(hits)
print('7. Page pair'); 
for lab, rx in (('part titles', r'class="acc-t[^"]*">(.*?)<'), ('sub-lines', r'class="acc-meta">(.*?)<')):
    a = re.findall(rx, open(P, encoding='utf-8').read())[:11]; b = re.findall(rx, open(F, encoding='utf-8').read())[:11]
    print('   ', lab, 'identical on both pages:', a == b)
    bad += a != b
print('slides:', len(slides), '· problems:', bad)
