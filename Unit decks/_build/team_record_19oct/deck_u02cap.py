# -*- coding: utf-8 -*-
"""Unit 2 deck after the capstone rework (19 October 2026), built on Carol's deck.
Section 4 teaches the worked example (Kestrel Point Services); Section 5 is the group's own build on the shared
group record; the 1.2 reflection is removed; 5.1 Role Contribution and 5.2 are gone.
usage: python3 deck_u02cap.py <deck in> <deck out> <portal_5_1.png> <ref_5.1.png>"""
import sys, copy
from pptx import Presentation
from pptx.util import Emu, Pt
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
import notes_new as N

src, out, portal, ref = sys.argv[1:5]
p = Presentation(src)
S = lambda n: p.slides[n - 1]
def shape(s, name):
    for sh in s.shapes:
        if sh.name == name: return sh
    raise KeyError(name)
def runs(s, name, *texts):
    r = shape(s, name).text_frame.paragraphs[0].runs
    assert len(r) == len(texts), (name, len(r))
    for x, t in zip(r, texts): x.text = t
def drop(s, name):
    sh = shape(s, name); sh._element.getparent().remove(sh._element)
import re
LABELS = ('Deepen, say: ', 'Say: ', 'Ask: ', 'Add: ', 'Read from the slide: ', 'Tell participants: ', 'Say how the group checks: ')
def set_notes(s, text):
    tf = s.notes_slide.notes_text_frame
    tf.text = text
    for para in tf.paragraphs:
        t = ''.join(r.text for r in para.runs)
        if not t: continue
        head = (t.strip() == t.strip().upper() and re.search('[A-Z]', t)) or re.match(r'^\d+\. ', t) or t.endswith(':') and t.upper() == t
        if head:
            for r in para.runs: r.font.bold = True
            continue
        for lab in LABELS:
            if t.startswith(lab):
                r0 = para.runs[0]; r0.text = lab; r0.font.bold = True
                rest = para.add_run(); rest.text = t[len(lab):]; rest.font.bold = None
                for r in para.runs[1:-1]: r.text = ''
                break

# 1 and 8 and 36: partial notes changes
for n, pairs in N.PART.items():
    norm = lambda x: '\n'.join(l.rstrip() for l in x.split('\n'))
    t = norm(S(n).notes_slide.notes_text_frame.text)
    for a, b in pairs:
        a = norm(a)
        assert a in t, (n, a[:60])
        t = t.replace(a, b)
    set_notes(S(n), t)

# 10: the 1.2 reflection comes off
drop(S(10), 'Reflection 1.2'); drop(S(10), 'Part 1.2')

# 25: Section 4 subtitle
runs(S(25), 'Text 3', 'One Worked Example: Architecture · Intent · Success in Practice')
# 26: the company of the worked example
runs(S(26), 'Text 0', '4.1a · The worked example: Kestrel Point Services')
runs(S(26), 'Text 2', 'Company:  ', 'facilities management and technical maintenance')
runs(S(26), 'Text 4', 'Size:  ', 'about 1,200 employees, eight regional service hubs')
runs(S(26), 'Text 6', 'Problem:  ', 'service varies, requests stall between functions')
runs(S(26), 'Text 8', 'Period:  ', 'the 2027–2029 strategy')
# 27: how each element works
runs(S(27), 'Text 0', '4.1b · Step 1: how each element works')
runs(S(27), 'Text 2', 'Answer  ', 'the four questions of the element')
runs(S(27), 'Text 4', 'Write  ', 'what the element is saying, in one or two sentences')
runs(S(27), 'Text 6', 'Ask  ', '“Is this what we mean?”')
runs(S(27), 'Text 8', 'Revise  ', 'until it says what the group means')
runs(S(27), 'Text 10', '14 of 14:  ', 'the elements come together under five headings')
runs(S(29), 'Text 0', '4.1d · Six checks on the Architecture')
runs(S(30), 'Text 0', '4.2 · Step 2: the Strategy Intent Statement')
runs(S(31), 'Text 0', '4.3a · Step 3: Success in Practice')
runs(S(32), 'Text 0', '4.3b · Three connected outputs')
runs(S(32), 'Text 14', 'In Section 5 each group builds the three for its own company.')
# 33: Section 5 subtitle
runs(S(33), 'Text 3', 'Build Your Company’s Strategy')
# 34: 5.1 Your Company's Strategy, with the portal picture
s34 = S(34)
runs(s34, 'Text 0', '5.1 · Your Company’s Strategy')
for k in range(3, 23):
    drop(s34, 'Shape %d' % k if k % 2 else 'Text %d' % k)
q = shape(s34, 'Shape 1'); qt = shape(s34, 'Text 2')
q.height = qt.height = Emu(int(76 * 12700))
runs(s34, 'Text 2', 'One scribe per assignment · every member agrees')
pic = s34.shapes.add_picture(portal, Emu(int(43 * 12700)), Emu(int(212 * 12700)), width=Emu(int(873 * 12700)))
if pic.height > Emu(int(300 * 12700)):
    r = 300 * 12700 / pic.height; pic.height = Emu(int(300 * 12700)); pic.width = Emu(int(pic.width * r))
pic.name = 'Portal 5.1'
# 35: Section 5 reflection is now in 5.1
s35 = S(35)
old = shape(s35, 'Reflection 5.2'); L, T, W = old.left, old.top, old.width
old._element.getparent().remove(old._element)
newp = s35.shapes.add_picture(ref, L, T, width=W); newp.name = 'Reflection 5.1'
oval = shape(s35, 'Part 5.2'); oval.name = 'Part 5.1'
oval.text_frame.paragraphs[0].runs[0].text = '5.1'
# full notes
for n, text in N.FULL.items():
    set_notes(S(n), text)
p.save(out)
print('saved', out)
