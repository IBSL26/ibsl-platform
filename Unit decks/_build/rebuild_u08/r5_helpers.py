A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
def plain(path):
    t = open(path, encoding='utf-8').read()
    t = re.sub(r'<style.*?</style>|<script.*?</script>|<!--.*?-->', ' ', t, flags=re.S)
    t = re.sub(r'</?(?:em|strong|b|i)\b[^>]*>', '', t)
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', t)))
def tx(s): return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', '', s))).strip()
def on(s, where='pf'):
    s = tx(s)
    for w in where: assert s in (vp if w == 'p' else vf), ('not on the %s page: ' % w) + s
    return s
def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def first(s): t = texts(s); return t[0].text_frame.text if t else ''
def set_text(shape, *run_texts):
    ps = shape.text_frame.paragraphs; runs = ps[0].runs
    assert len(runs) == len(run_texts), (shape.name, [r.text for r in runs], run_texts)
    for r, t in zip(runs, run_texts): r.text = t
def find(slide, start):
    m = [sh for sh in texts(slide) if sh.text_frame.text.startswith(start)]; assert len(m) == 1, (start, [x.text_frame.text for x in texts(slide)]); return m[0]
BAD = re.compile(r'\b(minutes?|mins?|Timing|working days|five-day|Listen for|Challenge any|Check that|Push back)\b')
def set_notes(slide, lines, watch_ok=False):
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in lines:
        assert not BAD.search(line), line
        if not watch_ok: assert 'Watch for' not in line, line
        p = etree.SubElement(body, '{%s}p' % A)
        bold = line.startswith('§'); text = line[1:] if bold else line
        if not text.strip():
            e = etree.SubElement(p, '{%s}endParaRPr' % A); e.set('lang', 'en-GB'); e.set('dirty', '0'); continue
        r = etree.SubElement(p, '{%s}r' % A); rp = etree.SubElement(r, '{%s}rPr' % A); rp.set('lang', 'en-GB')
        if bold: rp.set('b', '1')
        rp.set('dirty', '0'); etree.SubElement(r, '{%s}t' % A).text = text
def note_replace(slide, pairs):
    """Replace text inside Carol's own notes, paragraph by paragraph; each old text must be found."""
    found = {a: 0 for a, b in pairs}
    for p in slide.notes_slide.notes_text_frame.paragraphs:
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in pairs:
            if a in new: new = new.replace(a, b); found[a] += 1
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]: r.text = ''
    for a, n in found.items(): assert n, 'not in the notes: ' + a
def duplicate(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId] = dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID + 'embed', RID + 'link', RID + 'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    nt, st = dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'), qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag == qn('p:sp')]: nt.append(copy.deepcopy(el))
    return dst
def drop_shapes(slide, names):
    for sh in [s for s in slide.shapes if s.name in names]: sh._element.getparent().remove(sh._element)
OVAL = None
def oval(slide, label, left, top):
    global OVAL
    from pptx.oxml import parse_xml
    el = parse_xml(OVAL); nid = max(int(e.get('id')) for e in slide._element.iter(qn('p:cNvPr'))) + 1
    el.find('.//' + qn('p:cNvPr')).set('id', str(nid)); slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]; sh.left, sh.top = left, top; sh.name = 'Part ' + label
    sh.text_frame.paragraphs[0].runs[0].text = label
    return sh
def reflection_slide(title, items, notes):
    """items: [(part label, picture file)] stacked; each with a gold oval at the left."""
    s = duplicate(S[6])  # a slide with the standard title box
    drop_shapes(s, [sh.name for sh in s.shapes if sh.name != 'Text 0'])
    set_text(texts(s)[0], title)
    left, right, top0, bottom = Inches(1.85), Inches(12.85), Inches(1.6), Inches(7.2)
    gap = Inches(0.22); width = right - left
    hs = []
    for _, f in items:
        w, h = Image.open(os.path.join(PICS, f)).size; hs.append(width * h / w)
    scale = min(1, (bottom - top0 - gap * (len(items) - 1)) / sum(hs)); print(title, 'picture scale', round(scale, 2))
    y = top0
    for (lab, f), h in zip(items, hs):
        pw = int(width * scale); ph = int(h * scale)
        pic = s.shapes.add_picture(os.path.join(PICS, f), left, int(y), width=pw); pic.name = 'Reflection ' + lab
        oval(s, lab, Inches(0.6), int(y + ph / 2 - Inches(0.475)))
        y += ph + gap
    assert y - gap <= bottom + 10
    set_notes(s, notes)
    return s

# the gold oval of the reflection slides (as in Carol's Unit 7 deck)
OVAL = ('''<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="900" name="Part"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="868680" cy="868680"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="C6A24C"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="C6A24C"/></a:solidFill></a:ln><a:effectLst/></p:spPr><p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-GB" sz="2400" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:latin typeface="Calibri"/></a:rPr><a:t>1.1</a:t></a:r></a:p></p:txBody></p:sp>''')
