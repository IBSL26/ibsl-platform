"""Add "Section N Reflections" slides to a unit deck (as in Carol's Unit 1 deck).

Usage: python3 add_reflection_slides.py <deck_in.pptx> <spec.json> <deck_out.pptx>

Each new slide shows, for one section, the reflection boxes of the participant page as pictures,
with the part number in a gold circle beside each. The slide goes straight after the slide
whose title starts with "after" in the spec. Existing slides and their notes are not touched.
Run once only: the script stops if a "Section N Reflections" slide is already in the deck.
"""
import copy, json, os, sys
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

GREEN, GOLD, WHITE = RGBColor(0x1A, 0x5C, 0x2C), RGBColor(0xC6, 0xA2, 0x4C), RGBColor(0xFF, 0xFF, 0xFF)

def shape_text(sh):
    """Text of a shape read from its XML (asking python-pptx for text_frame would add an empty one)."""
    tb = sh._element.find(qn("p:txBody"))
    if tb is None:
        return ""
    return "\n".join("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in tb.iter(qn("a:p"))).strip()

def first_text(slide):
    for sh in slide.shapes:
        if shape_text(sh):
            return shape_text(sh)
    return ""

def all_text(slide):
    return " | ".join(shape_text(sh) for sh in slide.shapes if shape_text(sh))

def main(deck, spec_path, out):
    spec = json.load(open(spec_path, encoding="utf-8"))
    base = os.path.dirname(os.path.abspath(spec_path))
    prs = Presentation(deck)
    slides = list(prs.slides)
    if any("Reflections" in first_text(s) and first_text(s).startswith("Section") for s in slides):
        sys.exit("This deck already has reflection slides.")
    model = next(s for s in slides if first_text(s) == spec["title_model"])     # slide whose title style is copied
    model_title = next(sh for sh in model.shapes if shape_text(sh) == spec["title_model"])
    bg = model._element.cSld.find(qn("p:bg"))
    lst = prs.slides._sldIdLst
    placed = []
    for item in spec["slides"]:
        hits = [i for i, s in enumerate(slides) if all_text(s).startswith(item["after"])]
        if len(hits) != 1:
            sys.exit(f"'{item['after']}' matches {len(hits)} slides")
        s = prs.slides.add_slide(model.slide_layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        if bg is not None:
            s._element.cSld.insert(0, copy.deepcopy(bg))
        t = copy.deepcopy(model_title._element)
        s.shapes._spTree.append(t)
        runs = t.findall(".//" + qn("a:r"))
        runs[0].find(qn("a:t")).text = item["title"]
        for r in runs[1:]:
            r.getparent().remove(r)
        top = Inches(1.72)
        pic_left, pic_w, dia, gap = Inches(1.85), Inches(10.85), Inches(0.95), Inches(0.32)
        for pic in item["pictures"]:
            path = os.path.join(base, pic["file"])
            w, h = Image.open(path).size
            ph_ = int(pic_w * h / w)
            p = s.shapes.add_picture(path, pic_left, top, width=pic_w)
            p._element.nvPicPr.cNvPr.set("descr", pic["alt"])
            p.name = "Reflection " + pic["part"]
            o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), top + int((ph_ - dia) / 2), dia, dia)
            o.name = "Part " + pic["part"]
            o.fill.solid(); o.fill.fore_color.rgb = GOLD; o.line.color.rgb = GOLD
            o.shadow.inherit = False
            tf = o.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
            par = tf.paragraphs[0]; par.alignment = PP_ALIGN.CENTER
            r = par.add_run(); r.text = pic["part"]
            r.font.size = Pt(24); r.font.bold = True; r.font.name = "Calibri"; r.font.color.rgb = WHITE
            top += ph_ + gap
        s.notes_slide.notes_text_frame.text = item["notes"]
        placed.append((hits[0], s))
    # move each new slide to sit after its anchor (work from the back so positions stay valid)
    ids = list(lst)
    new_ids = ids[len(slides):]
    for (pos, s), el in sorted(zip(placed, new_ids), key=lambda x: -x[0][0]):
        lst.remove(el)
        lst.insert(pos + 1, el)
    prs.save(out)
    print(f"{len(placed)} reflection slides added -> {out} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    main(*sys.argv[1:4])
