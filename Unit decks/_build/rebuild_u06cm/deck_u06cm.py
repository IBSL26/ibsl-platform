# -*- coding: utf-8 -*-
"""Unit 6 deck, rework of 14 October 2026 (Carol's "Amended content.docx"), built on top of Carol's own deck.
Usage: python3 deck_u06cm.py <her deck (28 slides, md5 19c48b04…)> <facilitator page> <pictures folder> <out.pptx>
Her design is kept: slides that still carry the content are edited; new slides are copies of her slide layouts with new text.
Every presenter note follows the approved format (HOW TO TEACH IT, numbered steps with Say / Ask, ON THE PORTAL, AFTER THE TEACHING)
and carries all of the facilitator page's content for that part, so nothing on the facilitator page is left out of the notes."""
import sys, os, re, copy, html, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import u6cm_content as C

deck, FPAGE, PICS, out = sys.argv[1:5]
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def tx(s):
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', s.replace('<br>', ' ')))).replace(' ,', ',').replace(' .', '.').replace(' :', ':').replace(' ;', ';').strip()


def plain(path):
    t = open(path, encoding='utf-8').read()
    t = re.sub(r'<style.*?</style>|<script.*?</script>|<!--.*?-->', ' ', t, flags=re.S)
    return tx(t)


VF = plain(FPAGE)


def F(s):
    """A sentence taken from the facilitator page: it must be on the page word for word."""
    s = tx(s)
    assert s in VF, 'not on the facilitator page: ' + s[:120]
    return s


prs = Presentation(deck)
S = list(prs.slides)
assert len(S) == 28, len(S)


def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def shp(s, name): m = [x for x in s.shapes if x.name == name]; assert len(m) == 1, (name, [x.name for x in s.shapes]); return m[0]


def set_runs(shape, *run_texts):
    runs = shape.text_frame.paragraphs[0].runs
    assert len(runs) >= len(run_texts), (shape.name, [r.text for r in runs], run_texts)
    for r, t in zip(runs, run_texts): r.text = t
    for r in runs[len(run_texts):]: r.text = ''


def set_paras(shape, lines):
    """Bulleted text box: one paragraph per line, in the format of the first paragraph."""
    txb = shape.text_frame._txBody
    ps = txb.findall(qn('a:p')); tpl = ps[0]
    for p in ps: txb.remove(p)
    for line in lines:
        p = copy.deepcopy(tpl)
        rs = p.findall(qn('a:r'))
        for r in rs[1:]: p.remove(r)
        rs[0].find(qn('a:t')).text = line
        txb.append(p)


def drop(slide, names):
    for sh in [x for x in slide.shapes if x.name in names]: sh._element.getparent().remove(sh._element)


def move(slide, names, dy):
    for sh in [x for x in slide.shapes if x.name in names]: sh.top = sh.top + int(dy)


def duplicate(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId] = dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID + 'embed', RID + 'link', RID + 'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    nt, st = dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'), qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag == qn('p:sp')]: nt.append(copy.deepcopy(el))
    return dst


BAD = re.compile(r'\b((?-i:Timing)|Listen for|Challenge any|Check that|Push back|Watch for)\b|rather than|instead of|in place of|automat(?!ically establish)|\blens\b', re.I)


def set_notes(slide, lines):
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in lines:
        assert not BAD.search(line), line
        p = etree.SubElement(body, '{%s}p' % A)
        bold = line.startswith('§'); text = line[1:] if bold else line
        if not text.strip():
            e = etree.SubElement(p, '{%s}endParaRPr' % A); e.set('lang', 'en-GB'); e.set('dirty', '0'); continue
        r = etree.SubElement(p, '{%s}r' % A); rp = etree.SubElement(r, '{%s}rPr' % A); rp.set('lang', 'en-GB')
        if bold: rp.set('b', '1')
        rp.set('dirty', '0'); etree.SubElement(r, '{%s}t' % A).text = text


OVAL = ('''<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="900" name="Part"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="868680" cy="868680"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="C6A24C"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="C6A24C"/></a:solidFill></a:ln><a:effectLst/></p:spPr><p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-GB" sz="2400" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:latin typeface="Calibri"/></a:rPr><a:t>1.1</a:t></a:r></a:p></p:txBody></p:sp>''')


def oval(slide, label, left, top):
    from pptx.oxml import parse_xml
    el = parse_xml(OVAL); nid = max(int(e.get('id')) for e in slide._element.iter(qn('p:cNvPr'))) + 1
    el.find('.//' + qn('p:cNvPr')).set('id', str(nid)); slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]; sh.left, sh.top = left, top; sh.name = 'Part ' + label
    sh.text_frame.paragraphs[0].runs[0].text = label


def refill(slide, title, items):
    drop(slide, [sh.name for sh in slide.shapes if sh.name != 'Text 0'])
    set_runs(shp(slide, 'Text 0'), title)
    left, right, top0, bottom = Inches(1.85), Inches(12.85), Inches(1.6), Inches(7.2)
    gap = Inches(0.3); width = right - left; hs = []
    for _, f in items:
        w, h = Image.open(os.path.join(PICS, f)).size; hs.append(width * h / w)
    scale = min(1, (bottom - top0 - gap * (len(items) - 1)) / sum(hs))
    y = top0
    for (lab, f), h in zip(items, hs):
        pw = int(width * scale); ph = int(h * scale)
        pic = slide.shapes.add_picture(os.path.join(PICS, f), left, int(y), width=pw); pic.name = 'Reflection ' + lab
        oval(slide, lab, Inches(0.6), int(y + ph / 2 - Inches(0.475)))
        y += ph + gap


def opener(slide, title, outcomes):
    set_runs(shp(slide, 'Text 3'), title)
    set_runs(shp(slide, 'Text 6'), outcomes[0]); set_runs(shp(slide, 'Text 8'), outcomes[1])
    if len(outcomes) == 3:
        # three outcomes: the rows move up a little and a third row is copied from the second
        for nm, dy in (('Shape 5', -0.15), ('Image 0', -0.15), ('Text 6', -0.15), ('Shape 7', -0.4), ('Image 1', -0.4), ('Text 8', -0.4)):
            move(slide, [nm], Inches(dy))
        for nm in ('Shape 7', 'Image 1', 'Text 8'):
            el = copy.deepcopy(shp(slide, nm)._element); slide.shapes._spTree.append(el)
            sh = slide.shapes[-1]; sh.name = nm + ' b'; sh.top = sh.top + Inches(1.15)
        set_runs(shp(slide, 'Text 8 b'), outcomes[2])
        for nm in ('Text 6', 'Text 8', 'Text 8 b'): shp(slide, nm).height = Inches(1.05)


def icon_rows(slide, rows, n_max=5):
    """Slides laid out as icon rows (Text 2, 4, 6, 8, 10 with two runs: label, text)."""
    names = ['Text 2', 'Text 4', 'Text 6', 'Text 8', 'Text 10']
    for i, nm in enumerate(names):
        if not [x for x in slide.shapes if x.name == nm]: continue
        if i < len(rows): set_runs(shp(slide, nm), *rows[i])
        else: drop(slide, ['Shape %d' % (2 * i + 1), 'Image %d' % i, nm])


def cards3(slide, heads, bodies):
    for nm, t in zip(('Text 3', 'Text 7', 'Text 11'), heads): set_runs(shp(slide, nm), t)
    for nm, t in zip(('Text 4', 'Text 8', 'Text 12'), bodies): set_runs(shp(slide, nm), t)


def cards2(slide, h1, s1, b1, h2, s2, b2, foot=None):
    set_runs(shp(slide, 'Text 3'), h1); set_runs(shp(slide, 'Text 4'), s1); set_paras(shp(slide, 'Text 5'), b1)
    set_runs(shp(slide, 'Text 8'), h2); set_runs(shp(slide, 'Text 9'), s2); set_paras(shp(slide, 'Text 10'), b2)
    for h, sub, body in ((h1, 'Text 4', 'Text 5'), (h2, 'Text 9', 'Text 10')):
        if len(h) > 22:   # a two-line card heading: the sub-line and the list move down
            for nm in (sub, body): move(slide, [nm], Inches(0.45))
            b = shp(slide, body); b.height = b.height - Inches(0.45)
    for sb, body in ((s1, 'Text 5'), (s2, 'Text 10')):
        if len(sb) > 28:  # a two-line sub-line: the list moves down again
            move(slide, [body], Inches(0.32)); b = shp(slide, body); b.height = b.height - Inches(0.32)
    if foot is not None: set_runs(shp(slide, 'Text 12'), foot)


def add_footer(slide, text):
    """The dark footer bar of Carol's slide 27 (Shape 13 + Text 14), copied onto a slide."""
    for nm in ('Shape 13', 'Text 14'):
        slide.shapes._spTree.append(copy.deepcopy(shp(S[26], nm)._element))
    set_runs(shp(slide, 'Text 14'), text)


def title(slide, t, name='Text 0'): set_runs(shp(slide, name), t)


def say(s): return 'Say: "%s"' % F(s)
def says(*ss): return 'Say: "%s"' % ' '.join(F(s) for s in ss)
def ask(s): return 'Ask: "%s"' % s


PORTAL = '§ON THE PORTAL, AFTER THE TEACHING:'
TEACH = '§HOW TO TEACH IT:'
REFQ = {k: v[1] for k, v in C.REFS.items()}
SR = 'S2R® Contribution Map™'

# ═══ templates kept aside before any edit ═══
T_TWO = S[4]      # two cards with bullets and a footer (1.1 measurement / management)
T_ROWS = S[6]     # five icon rows (FACES)
T_THREE = S[7]    # three cards (Three Sights)
T_TEAM = S[14]    # two cards without a footer (team / leader)
T_REFL = S[10]    # Section 1 Reflections
new = {}
new['two_sys'] = duplicate(T_TWO)
new['esrg'] = duplicate(T_ROWS)
new['interact'] = duplicate(T_TEAM)
new['cm_logic'] = duplicate(T_THREE)
new['cm_connect'] = duplicate(T_TWO)
new['faces_two'] = duplicate(T_TEAM)
new['faces_short'] = duplicate(T_ROWS)
new['bridges_weak'] = duplicate(S[12])
new['acc_two'] = duplicate(T_TEAM)
new['hot'] = duplicate(T_ROWS)
new['refl3'] = duplicate(T_REFL)
new['map_ex'] = duplicate(T_ROWS)
new['map_read'] = duplicate(T_THREE)
new['faces_ex'] = duplicate(T_ROWS)

# ═══ 1 · title, 2 · Key learning outcomes, 3 · journey ═══
set_notes(S[0], ['§UNIT 6 · PERFORMANCE MANAGEMENT SETUP', 'Module 3 · Influence · Unit 6', '',
                 '§HOW TO USE THIS SLIDE:', '',
                 '§1. Open the unit with its intent.',
                 'Unit 6 sets up enterprise performance management on two operating systems: what people deliver (the Human Operating System) and what the organisation enables (the Organisational Operating System, examined through ESRG).',
                 'Participants learn the ' + SR + ', which connects an enterprise Key Result to the contributions of business functions, people and organisational conditions, and FACES, the five functions through which those contributions are managed. In Section 5 each group builds both for its own Capstone strategy.', '',
                 '§2. Set the tone.',
                 'Avoid: turning the Contribution Map or FACES into a compliance checklist. They are design and management tools.',
                 'Keep in view: a missed target is evidence of a performance gap; it does not establish its cause.',
                 'Probe leaders who equate "we have a scorecard" with "we have a PM system": "What happens after you look at the scorecard? What changes? Who decides what?"'])
set_notes(S[1], ['§KEY LEARNING OUTCOMES (identical in the participant and facilitator files)', '',
                 '§HOW TO USE THIS SLIDE:', '',
                 '§1. Read the four outcomes.'] + [t for t in [
                     '1. Evaluate whether organisational performance arrangements support the achievement of strategic priorities.',
                     '2. Define clear relationships between enterprise outcomes, leadership accountability and performance evidence.',
                     '3. Exercise consistent executive judgement on performance status and its strategic significance.',
                     '4. Interpret performance insight in ways that strengthen accountability, decision quality and organisational responsiveness.']] +
          ['', 'Each section slide carries that section’s learning outcomes.'])
set_notes(S[2], ['§THE UNIT JOURNEY', '', '§HOW TO USE THIS SLIDE:', '',
                 '§1. Walk the five sections.',
                 'Section 1 · Awareness: measurement and management, the two operating systems, the ' + SR + ' and FACES.',
                 'Section 2 · Intelligence: performance management as the execution bridge, and compliance and commitment architecture.',
                 'Section 3 · Extrapolating: executive performance perspectives and the continuous performance management shift.',
                 'Section 4 · Integration: one worked example that maps the contributions to an enterprise Key Result and manages them through FACES.',
                 'Section 5 · Application: each group builds its own Contribution Map and FACES application, as Capstone work.', '',
                 '§2. Name the questions the unit keeps asking.',
                 '"If your organisation stopped producing its performance reports tomorrow, which decisions would leaders no longer be able to make?"',
                 '"When a target is missed, how do you distinguish an individual contribution problem from an organisational delivery problem?"',
                 '"If employees are meeting their performance commitments but system outages continue to delay customer resolution, who must take responsibility for the performance gap?"'])


def opener_notes(n, lab, ttl, hero, slo, intro_lines):
    return (['§SECTION %d · %s' % (n, lab.upper()), ttl, hero, '', '§Section learning outcomes:'] +
            ['%d. %s' % (i + 1, o) for i, o in enumerate(slo)] + ['', '§HOW TO TEACH IT:', ''] + intro_lines)


# ═══ Section 1 ═══
opener(S[3], C.S1_TITLE, C.S1_SLO)
set_notes(S[3], opener_notes(1, 'Awareness — What', C.S1_TITLE, C.S1_HERO, C.S1_SLO,
                             ['§1. Frame the section.',
                              'Say: "This section sets the foundations: what performance management is, the two operating systems behind enterprise performance, the ' + SR + ' that connects them to an enterprise result, and FACES, the five functions that manage them."']))

s = S[4]  # 1.1 measurement and management (Carol's slide kept)
set_notes(s, ['§1.1 — MEASUREMENT VS MANAGEMENT (The core distinction)', '', TEACH, '',
              '§1. Reframe performance management before you define it.',
              'Ask leaders to clear their minds of the words "surveillance" and "pressure" and to picture the dashboard of a high-performance vehicle. A dashboard gives the driver accurate, real-time information so that they can make better decisions: whether to accelerate, adjust course, or pull over and refuel. It passes no judgement on the driver. The performance management system is the organisational equivalent: the visibility infrastructure of strategy execution.', '',
              '§2. Teach the distinction.',
              says('Performance management (PM) and performance measurement are distinct disciplines.', 'Measurement answers: What happened? It provides data, results, scores and dashboards.', 'Management answers: What do we do about it? It establishes the conversations, decisions, actions and accountability through which leaders respond to that evidence.'), '',
              '§3. Give the example.',
              says('An organisation can measure a missed service target accurately without managing the cause.', 'A dashboard may show that customer requests take twelve days against a standard of five.', 'Effective management asks whether the delay reflects unclear individual expectations, insufficient capacity, a broken hand-off, slow approval rights, or a combination.', 'Each explanation calls for a different response.'), '',
              '§4. Ask, now that the distinction is taught.',
              ask('If your organisation stopped producing its performance reports tomorrow, which decisions would leaders no longer be able to make?'),
              'Follow with: "When a target is missed, how do you distinguish an individual contribution problem from an organisational delivery problem?" Take two or three answers in the chat.'])

s = S[5]  # 1.1 the design sequence
title(s, '1.1 · The design sequence')
icon_rows(s, [('1 · Define success', ''), ('2 · Design how performance will be managed', ''),
              ('3 · Identify the evidence decisions require', ''), ('4 · Use the evidence  ', 'recognise, correct, improve')], 4)
set_notes(s, ['§1.1 — THE DESIGN SEQUENCE', '', TEACH, '',
              '§1. Teach the four steps in order.',
              'Say: "The design sequence is deliberate: define success; design how performance will be managed; identify the evidence that management decisions require; then use that evidence to recognise contribution, correct constraints and improve delivery."',
              'Step 1 · Define success: ' + F('Agree the result the organisation must achieve.'),
              'Step 2 · Design how performance will be managed: ' + F('Agree the conversations, decisions, actions and accountability that will respond to the evidence.'),
              'Step 3 · Identify the evidence that management decisions require: ' + F('Collect the data the management decisions need.'),
              'Step 4 · Use the evidence: ' + F('Recognise contribution, correct constraints and improve delivery.'), '',
              '§2. Close the part.',
              say('Reports become useful because the organisation has decided who will act on them and how.')])

s = new['two_sys']
title(s, '1.2 · Two operating systems')
cards2(s, 'Human Operating System', 'What people deliver', ['Results against KPIs and OKRs', 'Behavioural expectations'],
       'Organisational Operating System', 'What the organisation enables', ['Environment and Systems', 'Resources and Governance (ESRG)'],
       'Both must be visible, measured, reviewed and managed.')
set_notes(s, ['§1.2 — TWO OPERATING SYSTEMS BEHIND ENTERPRISE PERFORMANCE (Understanding what people deliver and what the organisation enables)', '', TEACH, '',
              '§1. Open with the question.',
              'Begin by asking participants: "When an employee or department misses a performance target, how does your organisation determine whether the problem is with individual performance or the conditions under which the work is performed?" Take two or three answers in the chat.', '',
              '§2. Teach why there are two systems.',
              says('Enterprise performance is the combined result of individual and collective contributions, supported by the organisational conditions required to deliver those results.'),
              says('An organisation may have capable employees who consistently work towards their targets, yet still struggle to achieve its strategic outcomes.', 'Equally, an organisation may have strong systems, sufficient resources and well-designed processes, but fail to achieve results because individuals are not delivering the contributions expected of them.'),
              says('This distinction is fundamental to Enterprise Performance Management (EPM).', 'A complete performance management system must evaluate what people contribute and whether the organisation provides the conditions necessary for those contributions to produce the intended results.', 'We refer to these as two interconnected operating systems.'), '',
              '§3. Teach the Human Operating System (left card).',
              says('The Human Operating System concerns the measurable performance of people within the organisation.', 'It examines whether individuals and teams deliver the results expected of them and demonstrate the behaviours required to achieve those results.', 'Human performance has two primary dimensions.'),
              'Performance against KPIs and objectives. ' + says('These establish what an individual or team is expected to deliver.', 'For example, a customer service employee may be expected to resolve 90% of enquiries within an agreed turnaround time, while a sales manager may be accountable for achieving a quarterly revenue target.', 'The performance management system must establish clear targets, evidence of achievement and a consistent basis for evaluating actual performance.'),
              'They may include strategic contributions expressed through OKRs, operational outputs measured through key performance indicators (KPIs), and specific performance commitments.',
              'Performance against behavioural expectations. ' + says('These establish how individuals are expected to perform their responsibilities.', 'For example, a manager may achieve departmental targets while consistently withholding information from other departments, undermining collaboration and delaying enterprise-wide delivery.', 'A complete assessment must therefore consider both the results achieved and the behaviours demonstrated.'),
              'They include observable behaviours such as collaboration, accountability, adherence to organisational values, responsiveness, ethical conduct and leadership effectiveness.',
              'The Human Operating System answers: "Are individuals and teams delivering their expected results and demonstrating the behaviours required for sustained organisational performance?"', '',
              '§4. Introduce the Organisational Operating System (right card).',
              says('The Organisational Operating System concerns the conditions that enable or constrain performance.'),
              'Say: "Within S2R®, these conditions are examined through ESRG: Environment, Systems, Resources and Governance. ESRG provides the structure for evaluating whether the organisation has the operating capability necessary to support its strategic ambition." The next slide takes the four elements one at a time.'])

s = new['esrg']
title(s, '1.2 · The Organisational Operating System · ESRG')
icon_rows(s, [('E  ', 'Environment: the conditions people work in'), ('S  ', 'Systems: processes, workflows, technology'),
              ('R  ', 'Resources: people, skills, capacity, funding'), ('G  ', 'Governance: decision rights, escalation')], 4)
esrg_notes = []
for i, (letter, name, body) in enumerate([('E', 'Environment', 'The conditions within which people and organisational processes operate. These include the working environment, organisational climate, cross-functional relationships and external operating conditions that influence delivery.|Example: a customer service team may have clear targets, but persistent conflict between departments prevents timely resolution of customer enquiries.|Performance management consideration: Are the operating conditions supporting or obstructing the achievement of expected results?'),
                                          ('S', 'Systems', 'The processes, workflows, technologies and operating mechanisms through which work is performed. Systems determine how efficiently and reliably the organisation converts effort and resources into outputs.|Example: employees may be expected to process 100 applications daily, but an unreliable processing system limits actual throughput to 60 applications.|Performance management consideration: Are the systems and processes capable of supporting the required performance standards?'),
                                          ('R', 'Resources', 'The people, skills, capacity, funding, equipment, information and other resources required to deliver results. Resource adequacy must be evaluated against the level of performance the organisation expects.|Example: a department may be required to double its service output while operating with the same capacity and equipment.|Performance management consideration: Does the organisation have sufficient and appropriately deployed resources to achieve its targets?'),
                                          ('G', 'Governance', 'The decision rights, accountability arrangements, policies, controls and escalation mechanisms that guide organisational execution. Governance determines whether decisions can be made at the appropriate level and whether constraints can be resolved in time.|Example: a department may have the resources to complete a project, but repeated approval delays prevent it from meeting agreed milestones.|Performance management consideration: Do governance arrangements enable timely decisions, clear accountability and effective resolution of performance constraints?')]):
    a, ex, pmc = body.split('|')
    esrg_notes += ['%s · %s. Say: "%s"' % (letter, name, F(a)), F(ex), F(pmc), '']
set_notes(s, ['§1.2 — THE ORGANISATIONAL OPERATING SYSTEM · ESRG', '', TEACH, '',
              '§1. Teach the four elements, one row at a time.'] + esrg_notes +
          ['§2. Close the slide.', 'The Organisational Operating System answers: "Does the organisation have the environment, systems, resources and governance necessary to deliver its expected results consistently?"'])

s = new['interact']
title(s, '1.2 · How the two operating systems interact')
cards2(s, 'Human Operating System', 'asks', ['Individual targets met?', 'Procedures followed?', 'Accountable and responsive?', 'Supervisory commitments met?'],
       'ESRG', 'asks', ['Working conditions support it?', 'System reliable?', 'Capacity sufficient?', 'Decisions on time?'])
rows = ['%s | %s' % (tx(a), tx(b)) for a, b in [
    ['Are employees meeting their individual resolution targets?', 'Environment: Are cross-functional working conditions supporting timely resolution?'],
    ['Are employees following the required service procedures?', 'Systems: Is the customer management system functioning reliably?'],
    ['Are employees demonstrating accountability and responsiveness?', 'Resources: Is staffing capacity sufficient for the volume of enquiries?'],
    ['Are team leaders meeting their supervisory commitments?', 'Governance: Are decisions and escalations being resolved within the required time?']]]
set_notes(s, ['§1.2 — HOW THE TWO OPERATING SYSTEMS INTERACT', '', TEACH, '',
              '§1. Teach the interdependence.',
              says('The two systems are interdependent.', 'Human performance measures what individuals and teams contribute.', 'Organisational performance examines whether ESRG provides the conditions required for those contributions to produce enterprise results.'), '',
              '§2. Use the customer service example.',
              'Say: "Consider a customer service department with a target of resolving 90% of enquiries within 48 hours. At the end of the month, the department achieves only 65%. A conventional performance review may conclude that the employees are underperforming. An integrated EPM system examines both operating systems before determining the appropriate management response."',
              'Read the questions on each side (Human Operating System | ESRG):'] + rows + ['',
              '§3. Show the two possible findings.',
              says('The evidence may reveal that employees are performing their responsibilities correctly, but system failures and approval delays are preventing the department from achieving its targets.', 'Alternatively, ESRG conditions may be adequate, while employees are consistently failing to meet agreed performance expectations.'),
              says('Both situations produce disappointing results, but they require different management responses.', 'This is why an enterprise performance management system must distinguish between human performance shortfalls and organisational operating constraints.', 'It enables leaders to establish accountability fairly, direct interventions appropriately and avoid repeatedly holding individuals responsible for conditions outside their control.'),
              F('Use the customer service example to demonstrate that a missed target is evidence of a performance gap, but does not automatically establish its cause.'),
              F('Emphasise that leaders must examine both the contribution expected from people and the ESRG conditions required to enable delivery.'), '',
              '§4. Teach the Enterprise Performance Principle.',
              says('Organisational performance emerges from the interaction between what individuals and teams contribute and how effectively the organisation’s Environment, Systems, Resources and Governance enable those contributions to produce results.'),
              'Say: "An integrated EPM system must therefore establish two connected lines of performance accountability. Human Performance: individual and team KPIs, OKRs and behavioural expectations. Organisational Performance: ESRG effectiveness in enabling the achievement of enterprise outcomes. Both must be visible, measured, reviewed and managed."',
              F('The central teaching point is that performance accountability must extend to those responsible for organisational operating conditions as well as those responsible for individual and team targets.'), '',
              PORTAL, 'Participants write a personal reflection at the foot of part 1.2 (Learning Portfolio): ' + REFQ['ref13']])

s = new['cm_logic']
title(s, '1.3 · The ' + SR)
cards3(s, ['Step 1 · Enterprise OKR', 'Step 2 · Functions', 'Step 3 · People and ESRG'],
       ['What must the organisation achieve?', 'What must each function deliver?', 'What is required to deliver each contribution?'])
set_notes(s, ['§1.3 — THE S2R® CONTRIBUTION MAP™ (Connecting enterprise strategic outcomes to functional contributions)', '', TEACH, '',
              '§1. Introduce the map with the customer service example.',
              F('Introduce the Contribution Map using the customer service example.') + ' ' + F('Begin with the enterprise key result of improving 48-hour customer enquiry resolution from 65% to 90%.'),
              F('Explain that achieving this outcome requires contributions from several functions, each with different responsibilities.') + ' ' + F('Walk participants through the map, showing how Customer Service, IT, Operations and Workforce Planning contribute to the same enterprise result.'), '',
              '§2. Teach what the map is.',
              says('The ' + SR + ' is a strategy-to-execution tool that identifies the interconnected contributions required from business functions, individuals and teams, and organisational operating conditions to achieve an enterprise strategic outcome.'),
              'Say: "It establishes a clear line of sight between the organisation’s strategic intent and the contributions required across its two operating systems: the Human Operating System, the results and behaviours individuals and teams must deliver; and the Organisational Operating System, the Environment, Systems, Resources and Governance (ESRG) necessary to enable those contributions."',
              says('The map identifies the relationships and dependencies between functions, making visible how their combined contributions support enterprise results.', 'Where appropriate, it also establishes measurable performance commitments or evidence that the required organisational conditions are in place.'),
              'Say: "The central principle: enterprise results are achieved through connected human and organisational contributions, supported by clear functional accountability."', '',
              '§3. Teach why the map matters.',
              says('An organisation achieves its strategic results through the combined contributions of its business functions.', 'Each function performs different activities, possesses different capabilities and carries different responsibilities.', 'However, these contributions must ultimately support the same enterprise strategic outcomes.'),
              says('A common weakness in performance management is that departments establish their own KPIs and measure their performance independently.', 'Finance measures financial performance, HR measures people-related outcomes, IT measures technology performance, and Operations measures operational efficiency.', 'Each department may report satisfactory performance while the organisation continues to miss its strategic targets.', 'This happens when functional performance is managed without sufficient clarity about how the contributions connect to enterprise results.'),
              says('For example, an organisation may have an enterprise objective to improve customer service.', 'Customer Service may focus on resolving enquiries, IT on maintaining system availability, Operations on processing requests, and HR on ensuring adequate staffing.', 'All four functions may have relevant KPIs.', 'However, unless their contributions are connected to the enterprise objective, leaders may struggle to establish whether the combined performance of these functions is sufficient to achieve the intended result.'),
              'Say: "The ' + SR + ' provides that connection."', '',
              '§4. Teach the logic, one card at a time.',
              'Step 1 · Enterprise OKR: what must the organisation achieve? ' + F('Identify the agreed enterprise strategic objective and its measurable key result.') + ' ' + F('This establishes the collective outcome towards which the organisation is working.'),
              'Step 2 · Required functional contributions: what must each function deliver? ' + F('Identify the specific contribution each business function must make towards achieving the enterprise outcome.') + ' ' + F('These contributions must reflect the interdependencies between functions.'),
              'Step 3 · Human and organisational contributions: what is required to deliver each functional contribution? For each function, identify the results and behaviours required from individuals and teams, together with measurable performance commitments wherever practical; and the specific ESRG conditions necessary to enable delivery, together with appropriate performance measures or evidence that those conditions are established and functioning.',
              F('Only the ESRG elements relevant to the required contribution are included.') + ' ' + F('Not every ESRG requirement needs a numerical metric.'), '',
              '§5. Emphasise the three teaching points.',
              '1. ' + F('Functional KPIs must have a clear connection to enterprise outcomes.') + ' ' + F('A department can perform well against its own measures without making the contribution the enterprise requires.'),
              '2. ' + F('Functional contributions are interdependent.') + ' ' + F('A weakness in one function can undermine the results of several others.'),
              '3. ' + F('Accountability must follow the contribution.') + ' ' + F('Each function must be accountable for the performance commitments it controls, while collective enterprise results remain a shared leadership responsibility.')])

s = new['cm_connect']
title(s, '1.3 · How the map connects the two systems')
cards2(s, 'Human Performance', 'What must people deliver?', ['Results and behaviours', 'Measurable commitments'],
       'Organisational Performance', 'What must the organisation enable?', ['ESRG requirements', 'Measures or evidence of readiness'],
       'Connected dimensions that require distinct management attention.')
set_notes(s, ['§1.3 — HOW THE CONTRIBUTION MAP CONNECTS TO THE TWO OPERATING SYSTEMS', '', TEACH, '',
              '§1. Teach the two dimensions.',
              says('The ' + SR + ' integrates the two operating systems into the design of enterprise performance.', 'For every required functional contribution, the map identifies two interconnected dimensions.'),
              'Human Performance: what must people deliver? ' + F('The map establishes the expected results and behaviours of individuals and teams, supported by measurable performance commitments wherever practical.'),
              'Organisational Performance: what must the organisation enable? ' + F('The map identifies the specific Environment, Systems, Resources and Governance (ESRG) requirements necessary for people to deliver their contributions.') + ' ' + F('These requirements may be expressed through numerical performance standards or other evidence of organisational readiness and effectiveness.'), '',
              '§2. Give the IT example.',
              says('For example, maintaining reliable customer management systems requires IT employees to monitor and resolve incidents promptly.', 'It also requires reliable infrastructure and sufficient technical support capacity.', 'The human contribution can be assessed through incident-resolution performance.', 'Organisational capability can be assessed through system availability and evidence of adequate technical support arrangements.'),
              says('The two dimensions are connected but require distinct management attention.', 'The Contribution Map therefore establishes what must be delivered by people, what must be enabled by the organisation, and how both contribute to the enterprise result.'), '',
              '§3. Bridge to FACES.',
              says('The ' + SR + ' establishes what each function must contribute, what people must deliver and which organisational conditions must support their performance.', 'It also identifies the measures and evidence required to assess these contributions.'),
              says('The next challenge is to ensure that these commitments are consistently delivered and that performance gaps trigger appropriate management action.', 'The FACES Framework establishes the five management functions required to achieve this across the Human and Organisational Operating Systems.'), '',
              PORTAL, 'Participants write a personal reflection at the foot of part 1.3 (Learning Portfolio): ' + REFQ['ref14']])

s = S[6]  # 1.4 FACES (Carol's slide)
title(s, '1.4 · The FACES framework')
icon_rows(s, [('F  ', 'Foster a Results Culture'), ('A  ', 'Align Performance'), ('C  ', 'Create Measurement Transparency'),
              ('E  ', 'Embed Shared Accountability'), ('S  ', 'Support Continuous Development')])
set_notes(s, ['§1.4 — THE FACES FRAMEWORK (Five functions of an integrated Enterprise Performance Management system)', '', TEACH, '',
              '§1. Teach what FACES is.',
              says('The FACES Framework defines five interconnected functions that an effective Enterprise Performance Management (EPM) system must fulfil.'),
              'Say: "Each function applies to both operating systems. Human Performance: what individuals and teams deliver through their KPIs, OKRs and demonstrated behaviours. Organisational Performance: how effectively the organisation’s Environment, Systems, Resources and Governance (ESRG) enable the achievement of intended results."', '',
              '§2. Name the five functions on the slide.',
              'F · Foster a Results Culture · A · Align Performance · C · Create Measurement Transparency · E · Embed Shared Accountability · S · Support Continuous Development.',
              'The next slide sets out each function for the people and for the ESRG conditions.'])

s = new['faces_two']
title(s, '1.4 · FACES across the two operating systems')
cards2(s, 'Human performance', 'KPIs, OKRs and behaviours', ['F · Ownership of commitments', 'A · Line of sight to outcomes', 'C · Clear, credible criteria', 'E · Clear responsibility', 'S · Capability development'],
       'Organisational performance', 'ESRG', ['F · Leaders fix constraints', 'A · ESRG fit to strategy', 'C · ESRG made visible', 'E · Named ESRG owners', 'S · Operating capability'])
fl = []
for r in C.FACES_ROWS:
    letter, name = re.match(r'<strong>([A-Z]) &mdash; ([^<]+)</strong>', r[0]).groups(); hum, org = r[1], r[2]
    fl += ['%s · %s.' % (letter, name), 'Human performance. Say: "%s"' % F(hum), 'Organisational performance. Say: "%s"' % F(org), '']
set_notes(s, ['§1.4 — FACES: COMPARING HUMAN AND ORGANISATIONAL PERFORMANCE', '', TEACH, '',
              '§1. Take each function across both columns, the people first, then the ESRG conditions.', ''] + fl)

s = new['faces_short']
title(s, '1.4 · When the enterprise result falls short')
icon_rows(s, [('Customer Service  ', 'follow-through below standard'), ('IT  ', 'incidents handled, availability below standard'),
              ('Operations  ', 'approval delays and processing gaps'), ('Workforce Planning  ', 'schedules done, coverage short')], 4)
add_footer(s, '48-hour resolution: 65% against a year-end target of 90%')
SH = [('Customer Service', 'Follow-through on outstanding enquiries is below the agreed standard.', 'Outstanding cases are reviewed against assigned responsibilities. Employees receive feedback on missed commitments, recurring obstacles are escalated, and subsequent follow-through results establish whether performance is improving.'),
      ('IT', 'Incident-response commitments are being achieved, but system availability remains below the required standard.', 'Incident-response performance and system availability are examined separately. IT leadership investigates recurring outages, assigns responsibility for correcting the underlying infrastructure problems and reviews subsequent availability results.'),
      ('Operations', 'Processing performance is affected by approval delays and some employee processing gaps.', 'Processing records distinguish employee-related delays from approval bottlenecks. Employees receive support where performance gaps are identified, while the appropriate decision owners review approval arrangements that are slowing delivery.'),
      ('HR / Workforce Planning', 'Staffing schedules are completed as required, but actual staffing coverage remains insufficient.', 'Completed schedules are compared with actual staffing requirements. Workforce Planning coordinates available staffing adjustments, while unresolved capacity requirements are submitted to the appropriate resource decision-makers.')]
set_notes(s, ['§1.4 — HOW FACES WORKS WHEN THE ENTERPRISE RESULT FALLS SHORT', '', TEACH, '',
              '§1. Set the scene.',
              says('Assume the organisation reviews its performance and finds that only 65% of customer enquiries are being resolved within 48 hours, against the year-end target of 90%.', 'The review reveals different performance conditions across the four functions.'), '',
              '§2. Walk through the four functions.',
              F('Walk participants through the four functions, showing how the same FACES principles are applied differently according to the contributions required.')] +
          sum([['%s. Evidence: %s How FACES is applied: %s' % (a, F(b), F(c))] for a, b, c in SH], []) + ['',
              '§3. Draw the distinction.',
              F('Draw attention to the distinction between human performance accountability and organisational enablement accountability.'), '',
              '§4. Conclude by reinforcing the connection.',
              says('The ' + SR + ' establishes the contributions, performance measures and enabling conditions required for enterprise results.', 'FACES establishes the management disciplines through which those contributions are aligned, evaluated, acted upon and continuously strengthened.'), '',
              PORTAL, 'Participants write a personal reflection at the foot of part 1.4 (Learning Portfolio): ' + REFQ['ref15']])

s = S[10]
refill(s, 'Section 1 Reflections', [('1.2', 'ref_1.2.png'), ('1.3', 'ref_1.3.png'), ('1.4', 'ref_1.4.png')])
set_notes(s, ['§SECTION 1 REFLECTIONS', '', '§HOW TO USE THIS SLIDE:', '',
              '§1. Show the slide and read the three prompts aloud.', '1.2 Reflection — Two Operating Systems: "' + REFQ['ref13'] + '"',
              '1.3 Reflection — Contributions in Your Organisation: "' + REFQ['ref14'] + '"',
              '1.4 Reflection — FACES in Your Organisation: "' + REFQ['ref15'] + '"', '',
              '§2. Say what the reflections are for.', 'Say: "These reflections are personal. Name one real, recurring shortfall and look at whether your system covers the people and the ESRG conditions. Then map one of your own organisation’s Key Results the way the Contribution Map does, and look at how your organisation manages those contributions through FACES."', '',
              '§3. Bridge to Section 2.', 'Say: "You now know what must be managed. Section 2 shows how performance management keeps strategy connected to performance throughout delivery."', '',
              PORTAL, 'Participants write these personal reflections at the foot of parts 1.2, 1.3 and 1.4 of the participant file. They go into their Learning Portfolio when they submit the unit.'])

# ═══ Section 2 ═══
opener(S[11], C.S2_TITLE, C.S2_SLO)
set_notes(S[11], opener_notes(2, 'Intelligence — Why', C.S2_TITLE, C.S2_HERO, C.S2_SLO,
                              ['§1. Frame the section.',
                               'Say: "Section 1 set out what must be managed. Section 2 explains why performance management is the bridge that keeps strategy connected to performance throughout delivery, and how the architecture you choose shapes accountability."']))

s = S[12]  # 2.1 bridges
title(s, '2.1 · Performance management as the execution bridge')
drop(s, ['Shape 9', 'Image 4', 'Text 10'])
add_footer(s, 'Keeping strategy connected to performance throughout delivery.')
set_notes(s, ['§2.1 — PERFORMANCE MANAGEMENT AS THE EXECUTION BRIDGE (Four execution bridges)', '', TEACH, '',
              '§1. Name the challenge.',
              says('The ' + SR + ' establishes the contributions required to achieve enterprise results.', 'FACES defines the five functions through which human and organisational performance are managed.', 'The next challenge is sustaining the connection between these contributions and the intended results throughout execution.'),
              says('Strategy is delivered over time.', 'Performance expectations must remain clear, progress must be visible, and emerging gaps must lead to timely management responses.', 'Without these disciplines, the connection between strategic intent and actual delivery gradually weakens.'),
              'Without this bridge, organisations generate what might be called random acts of excellence: genuine effort, real capability, strong intentions, none of which reliably move the needle.', '',
              '§2. Teach the bridge.',
              says('This is why performance management serves as an execution bridge: it provides the disciplines that keep strategy connected to performance throughout delivery.'), '',
              '§3. Teach the four bridges, one row at a time.',
              'Bridge 1 · Decomposing the Big Picture: ' + F('Connecting strategic outcomes to actionable performance expectations.'),
              'Bridge 2 · Defining the Logic of Success: ' + F('Establishing how successful performance is understood and evaluated.'),
              'Bridge 3 · Calibrating the Rhythm: ' + F('Maintaining visibility of progress throughout the execution period.'),
              'Bridge 4 · Structuring the Dialogue: ' + F('Converting performance evidence into decisions and action.'),
              says('Together, these bridges explain why performance management is essential to sustaining the connection between strategy, human contribution, organisational capability and enterprise results.')])

s = new['bridges_weak']
title(s, '2.1 · How the four bridges work together')
drop(s, ['Shape 9', 'Image 4', 'Text 10'])
icon_rows(s, [('Bridge 1  ', 'line of sight and fit of conditions'), ('Bridge 2  ', 'a shared logic of success'), ('Bridge 3  ', 'visibility early enough to act'), ('Bridge 4  ', 'evidence turned into decisions')], 4)
bl = []
for r in C.BRIDGES:
    nm = tx(r[0]); why, hum, org, weak = r[1:]
    bl += ['Bridge ' + nm + '.', 'Why it matters. Say: "%s"' % F(why), 'Human Operating System: ' + F(hum), 'Organisational Operating System (ESRG): ' + F(org), 'What happens when the bridge weakens: ' + F(weak), '']
set_notes(s, ['§2.1.1 — HOW THE FOUR EXECUTION BRIDGES WORK TOGETHER', '', TEACH, '',
              '§1. Frame the slide.',
              says('The four bridges work together to sustain the connection between strategy and results.', 'Each addresses a particular risk that can weaken execution across the Human and Organisational Operating Systems.'),
              F('Take the bridges one at a time through the table: why each matters, what it sustains for the people, what it sustains for the ESRG conditions, and what happens when it weakens.'), '',
              '§2. Teach each bridge.', ''] + bl +
          ['§3. Ask, now that the bridges are taught.', ask('Which of the four bridges is weakest in your organisation, and what do you see happening as a result?') + ' Take two or three answers in the chat.'])

s = S[13]  # 2.2 table
title(s, '2.2 · Compliance and commitment architecture')
drop(s, ['Rectangle 49', 'Rectangle 51', 'Rectangle 53', 'Rectangle 55', 'Rectangle 57', 'Rectangle 59', 'Rectangle 61', 'Rectangle 63'])
move(s, ['Rectangle 65', 'Rectangle 67', 'Rectangle 69'], Inches(-1.6)); move(s, ['Rectangle 71'], Inches(-1.3))
ROWS22 = [('Rectangle 23', 'Rectangle 25', 'Rectangle 27', 'Rectangle 29'), ('Rectangle 31', 'Rectangle 33', 'Rectangle 35', 'Rectangle 37'),
          ('Rectangle 39', 'Rectangle 41', 'Rectangle 43', 'Rectangle 45'), ('Rectangle 47', 'Rectangle 65', 'Rectangle 67', 'Rectangle 69')]
for k, (ln, a, b_, c) in enumerate(ROWS22):   # four rows at 24pt, spaced over the slide (Rule 16)
    y = Inches(2.28 + 1.05 * k)
    shp(s, ln).top = y
    for nm in (a, b_, c):
        x = shp(s, nm); x.top = y + Inches(0.12); x.height = Inches(0.85)
        if nm == b_: x.width = Inches(4.7); x.text_frame.word_wrap = True
        for p in x.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(24)
shp(s, 'Rectangle 71').top = Inches(2.28 + 1.05 * 4)
for el in s._element.iter():
    if el.tag in (qn('a:rPr'), qn('a:endParaRPr'), qn('a:defRPr')) and el.get('sz') and int(el.get('sz')) < 2400: el.set('sz', '2400')
set_runs(shp(s, 'Rectangle 67'), 'Anxiety-producing; done to the individual')
set_runs(shp(s, 'Rectangle 69'), 'Energy-generating; done with the individual')
DIM = [('Purpose', 'Ensure minimum standards are met and deviations corrected.', 'Generate shared ownership of results and continuous improvement.'),
       ('Ownership', 'Leader holds accountability; team is subject to it.', 'Team co-owns accountability; leader facilitates.'),
       ('Tone', 'Assessment and judgement; hierarchical.', 'Inquiry and problem-solving; collaborative.'),
       ('PM Experience', 'Anxiety-producing; done to the individual.', 'Energy-generating; done with the individual.')]
set_notes(s, ['§2.2 — COMPLIANCE VS COMMITMENT ARCHITECTURE (The design choice that shapes performance behaviour)', '', TEACH, '',
              '§1. Name the design choice.',
              F('Teach the two architectures first, then the four dimensions, then how each applies to the people and to the ESRG conditions.'),
              says('The four execution bridges establish the disciplines required to sustain the connection between strategy and results.', 'However, their effectiveness is influenced by the way an organisation approaches performance expectations, accountability and performance conversations.', 'This introduces an important distinction between compliance-driven and commitment-driven performance management architectures.'), '',
              '§2. Teach the two architectures.',
              'A compliance-driven architecture emphasises adherence to prescribed expectations, monitoring performance and correcting deviations.',
              'A commitment-driven architecture emphasises shared understanding of expected results, ownership of contributions, constructive performance dialogue and collective responsibility for improvement.',
              says('Both approaches recognise the importance of standards and accountability.', 'Their differences become visible in how performance is experienced and how the organisation responds when results fall short.'), '',
              '§3. Teach the four dimensions, one row at a time.'] +
          ['%s. Compliance-driven: %s Commitment-driven: %s' % (a, F(b), F(c)) for a, b, c in DIM])

s = new['acc_two']
title(s, '2.2 · Accountability across the two operating systems')
cards2(s, 'Human Operating System', 'compliance and commitment', ['Compliance: targets and standards met?', 'Commitment: contribution understood and owned?'],
       'Organisational Operating System', 'compliance and commitment', ['Compliance: arrangements followed?', 'Commitment: arrangements effective, constraints addressed?'])
set_notes(s, ['§2.2.2 — ACCOUNTABILITY ACROSS THE TWO OPERATING SYSTEMS', '', TEACH, '',
              '§1. Apply the distinction to both systems.',
              says('The distinction between compliance and commitment applies to both human and organisational performance.'),
              'Human Operating System. ' + says('A compliance-driven approach tends to emphasise whether employees have met prescribed KPIs, OKRs and behavioural standards.', 'A commitment-driven approach maintains these expectations while strengthening employees’ understanding of their contribution, ownership of results and participation in performance improvement.'),
              'Organisational Operating System. ' + says('A compliance-driven approach may establish requirements for systems, resources, processes and governance, but concentrate primarily on whether prescribed arrangements have been followed.', 'A commitment-driven approach also examines whether those arrangements are effective in supporting the intended results and whether emerging constraints are being addressed.'), '',
              '§2. Show why it matters when performance falls short.',
              says('This distinction becomes particularly important when performance falls short.', 'An employee may miss a target because of inadequate follow-through, an organisational constraint or a combination of both.', 'A credible performance management system distinguishes between these causes before determining the appropriate response.'),
              says('Commitment strengthens accountability when employees remain responsible for their contributions and organisational leaders remain equally responsible for the enabling conditions within their authority.'), '',
              '§3. Ask, now that the architectures are taught.',
              ask('When a target is missed in your organisation, which of the two architectures does the review conversation sound like?') + ' Take two or three answers in the chat.'])

# ═══ Section 3 ═══
opener(S[16], C.S3_TITLE, C.S3_SLO)
set_notes(S[16], opener_notes(3, 'Extrapolating — Where', C.S3_TITLE, C.S3_HERO, C.S3_SLO,
                              ['§1. Frame the section.',
                               says('Section 2 established the four execution bridges that sustain the connection between strategy and results.', 'It also examined how compliance-driven and commitment-driven approaches influence accountability, performance behaviour and management responses.'),
                               'Say: "Section 3 examines where enterprise performance management becomes vulnerable, through two areas: Unity in Diversity, how executive performance perspectives influence what leaders notice, prioritise and interpret; and the Continuous Performance Management Shift, how performance management practices are evolving to strengthen timely feedback, development and collective performance."']))

s = S[17]  # 3.1 picture slide
set_runs(shp(s, 'Text 1'), '3.1 · Unity in Diversity: executive performance perspectives')
set_notes(s, ['§3.1 — UNITY IN DIVERSITY: EXECUTIVE PERFORMANCE PERSPECTIVES (Different executive perspectives, one enterprise result)', '', TEACH, '',
              '§1. Teach functional bias first.',
              F('Teach functional bias first: what executives notice, prioritise and interpret.'),
              says('Enterprise performance management depends on the collective contribution of executive leaders who hold different responsibilities.', 'Each executive function develops a particular understanding of performance based on its mandate, expertise, operating priorities and the evidence it routinely examines.', 'These perspectives are necessary because enterprise performance is multidimensional.'),
              says('This reflects the principle of functional bias, introduced earlier in CESI: leaders may interpret organisational situations primarily through the priorities and assumptions associated with their professional responsibilities.'),
              'Within enterprise performance management, functional bias can influence what executives notice (the performance indicators and developments that attract their attention), what executives prioritise (the results, risks and improvements they consider most important) and how executives interpret evidence (the conclusions they draw about organisational performance and its underlying causes).', '',
              '§2. Name Unity in Diversity.',
              says('The challenge for executive teams is to bring these perspectives together so that functional performance information supports a coherent understanding of enterprise results.'),
              'Say: "Unity in Diversity means integrating different executive performance perspectives into one collective understanding of enterprise progress."', '',
              '§3. Teach what the executive team must establish.',
              says('The executive team must establish whether the performance reported by individual functions is sufficient to produce the intended collective outcome.', 'This requires attention to both operating systems.'),
              'Human Operating System. ' + says('Executives examine whether individuals and teams are delivering their agreed KPIs, OKRs and behavioural expectations.', 'They also consider whether the required contributions are being coordinated effectively across functions.'),
              'Organisational Operating System. ' + says('Executives examine whether the relevant Environment, Systems, Resources and Governance conditions are enabling those contributions.', 'Where organisational constraints affect delivery, the appropriate executive owners must take responsibility for resolving them.'),
              'Enterprise Integration. The executive team brings the evidence together to determine whether functional performance is supporting the intended enterprise result; where human contribution or organisational capability is weakening delivery; and which executive decisions or coordinated interventions are required.',
              says('The ' + SR + ' provides the established relationships between contributions.', 'The four execution bridges provide the disciplines for sustaining those relationships.', 'Executive integration ensures that the resulting performance evidence is interpreted and acted upon collectively.'),
              'Key insight. Say: "Executive performance integration requires leaders to recognise the limits of their functional perspectives and use collective performance evidence to make decisions in the interests of enterprise results."'])

s = new['hot']
title(s, '3.1 · The ten executive performance hot zones')
drop(s, [sh.name for sh in s.shapes if sh.name != 'Text 0'])
tpl = {nm: shp(S[6], nm)._element for nm in ('Shape 1', 'Image 0', 'Text 2')}
img_rel = {}
HOTS = [('CEO  ', 'combined result unchecked'), ('CFO  ', 'savings read as success'), ('COO  ', 'outputs without outcomes'), ('CHRO  ', 'conditions left unchecked'),
        ('CIO/CTO  ', 'uptime without enablement'), ('CMO  ', 'promise beyond delivery'), ('CCO  ', 'sales beyond capacity'), ('CPO  ', 'savings over supply'),
        ('CRO  ', 'control slowing execution'), ('CSO  ', 'milestones without outcomes')]
srcrel = S[6].part.rels
for k, (role, spot) in enumerate(HOTS):
    col, row = k // 5, k % 5
    dx, dy = Inches(6.15 * col), Inches(1.07 * row)
    for nm, el in tpl.items():
        e = copy.deepcopy(el)
        if nm == 'Image 0':
            blip = e.find('.//' + qn('a:blip')); old = blip.get(RID + 'embed')
            if old not in img_rel: img_rel[old] = s.part.relate_to(srcrel[old]._target, srcrel[old].reltype)
            blip.set(RID + 'embed', img_rel[old])
        s.shapes._spTree.append(e)
        sh = s.shapes[-1]; sh.name = '%s r%d' % (nm, k); sh.left = sh.left + dx; sh.top = sh.top + dy
        if nm == 'Text 2':
            sh.width = Inches(5.0); set_runs(sh, role, spot)
nid = 2
for e in s._element.iter(qn('p:cNvPr')): e.set('id', str(nid)); nid += 1
hz = []
for r in C.HOTZ:
    hz += [tx(r[0]) + '.', 'Natural performance focus: ' + F(r[1]), 'Potential performance blind spot: ' + F(r[2]), 'Enterprise integration requirement: ' + F(r[3]), '']
set_notes(s, ['§3.1.1 — THE TEN EXECUTIVE PERFORMANCE HOT ZONES', '', TEACH, '',
              '§1. Frame the slide.',
              'Say: "The ten executive performance perspectives show the blind spots each may create and the integration required to support enterprise results."',
              F('Then open the role cards, starting with the CEO, and read the natural focus, the blind spot and the integration requirement of each.'), '',
              '§2. Teach each role, one row at a time.', ''] + hz +
          [F('Executive titles and responsibilities vary between organisations.') + ' ' + F('The table illustrates functional perspectives.') + ' ' + F('Each organisation reads it against its own executive structure.'), '',
           '§3. Ask, now that the hot zones are taught.',
           ask('Which blind spot do you recognise in your own executive team, and what does it cost the enterprise result?') + ' Take two or three answers in the chat.'])

s = S[19]  # 3.2 continuous shift
title(s, '3.2 · The continuous PM shift')
set_runs(shp(s, 'Text 2'), '3')
set_runs(shp(s, 'Text 3'), 'shifts in performance management practice')
set_runs(shp(s, 'Text 5'), 'Shift 1  ', 'From Annual to Continuous')
set_runs(shp(s, 'Text 7'), 'Shift 2  ', 'From Rating-Centric to Development-Centric')
set_runs(shp(s, 'Text 9'), 'Shift 3  ', 'From Individual to Team Performance')
SHIFTS = [('1 · From Annual to Continuous', 'Performance is formally evaluated at annual or biannual intervals.', 'Regular check-ins, periodic reviews and timely feedback support performance throughout execution.', 'Human contributions and relevant ESRG conditions are reviewed frequently enough to identify emerging gaps and influence results.'),
          ('2 · From Rating-Centric to Development-Centric', 'Performance conversations concentrate heavily on ratings and historical achievement.', 'Performance evidence informs development, capability strengthening and future improvement.', 'Individual development is considered alongside improvements to organisational processes, systems, resources and governance arrangements.'),
          ('3 · From Individual to Team Performance', 'Performance assessment concentrates primarily on individual achievement.', 'Individual accountability is complemented by attention to team and cross-functional contributions.', 'Human collaboration and organisational dependencies become more visible within the performance management system.')]
set_notes(s, ['§3.2 — THE CONTINUOUS PERFORMANCE MANAGEMENT SHIFT (How performance management practices are evolving)', '', TEACH, '',
              '§1. Name the shift.',
              says('Organisations are increasingly recognising the limitations of performance management systems built primarily around periodic assessments and individual ratings.', 'Three important shifts are influencing how performance management is approached.'), '',
              '§2. Walk through the three shifts, then the implication of each for the two operating systems.'] +
          sum([['Shift %s. Traditional orientation: %s Evolving orientation: %s Implication for the two operating systems: %s' % (a, F(b), F(c), F(d))] for a, b, c, d in SHIFTS], []) + ['',
              '§3. Broaden the conversation.',
              'Say: "These developments reinforce the importance of managing performance while execution is unfolding. They also broaden the performance management conversation."',
              F('Employee development remains important, but organisations must also improve the operating conditions that influence performance.'),
              F('Individual accountability remains necessary, but enterprise results often require coordinated contributions across teams and functions.'),
              F('Periodic assessments remain relevant, but they must be supported by timely performance conversations that allow leaders to respond to emerging challenges.'),
              'Key insight. Say: "Continuous performance management strengthens the organisation’s ability to recognise, interpret and respond to performance conditions while there is still an opportunity to influence results."', '',
              '§4. Ask, now that the shifts are taught.',
              ask('Where does your organisation sit on each of the three shifts today?') + ' Take two or three answers in the chat.', '',
              PORTAL, 'Participants write a personal reflection at the foot of part 3.2 (Learning Portfolio): ' + REFQ['ref12']])

s = new['refl3']
refill(s, 'Section 3 Reflections', [('3.2', 'ref_3.2.png')])
set_notes(s, ['§SECTION 3 REFLECTIONS', '', '§HOW TO USE THIS SLIDE:', '',
              '§1. Show the slide and read the prompt aloud.', '3.2 Reflection — Your Current PM System: "' + REFQ['ref12'] + '"', '',
              '§2. Say what the reflection is for.', 'Say: "Describe the architecture you actually have today, against the three shifts, and name the system your organisation uses."', '',
              '§3. Bridge to Section 4.', 'Say: "Section 4 brings the unit together in one worked example: the contributions an enterprise Key Result requires, and how FACES manages them."', '',
              PORTAL, 'Participants write this personal reflection at the foot of part 3.2 of the participant file. It goes into their Learning Portfolio when they submit the unit.'])

# ═══ Section 4 ═══
opener(S[20], C.S4_TITLE, C.S4_SLO)
set_notes(S[20], opener_notes(4, 'Integration — Collective', C.S4_TITLE, C.S4_HERO, C.S4_SLO,
                              ['§1. Frame the section.',
                               says('The previous sections established the foundations of enterprise performance management, the disciplines that sustain execution and the importance of integrating executive performance perspectives.', 'We now bring these concepts together through a worked example.')]))

s = new['map_ex']
title(s, '4.1 · Mapping the contributions: customer service')
icon_rows(s, [('Customer Service  ', 'resolve enquiries within 48 hours'), ('IT  ', 'reliable customer management systems'),
              ('Operations  ', 'processing and approvals in time'), ('Workforce Planning  ', 'enough staffing capacity')], 4)
add_footer(s, 'Key Result: 48-hour resolution from 65% to 90% by year-end')
ml = []
for r in C.MAP_ROWS:
    ml += [tx(r[0]) + '.', 'Required functional contribution: ' + F(r[1]), 'Human Operating System, required contribution: ' + F(r[2]), 'HOS performance metric: ' + tx(r[3]), 'Required ESRG contribution: ' + tx(r[4]), 'ESRG measure or evidence of readiness: ' + tx(r[5]), '']
set_notes(s, ['§4.1 — MAPPING THE CONTRIBUTIONS REQUIRED FOR ENTERPRISE RESULTS (Worked example: improving customer service)', '', TEACH, '',
              '§1. Frame the worked example.',
              says('The ' + SR + ' establishes the interconnected contributions required to achieve an enterprise result.', 'It identifies what each function must deliver, the performance expected from individuals and teams, and the organisational conditions necessary to enable those contributions.', 'The following example demonstrates how these requirements are connected across four business functions.'), '',
              '§2. Establish the enterprise result.',
              'Enterprise Objective: Deliver a reliable and responsive customer service experience.',
              'Begin with the enterprise key result: "Increase customer enquiries resolved within 48 hours from 65% to 90% by year-end."',
              ask('Which functions must contribute to achieving this result?') + ' ' + F('Draw attention to Customer Service, IT, Operations and Workforce Planning.'), '',
              '§3. Walk through the Contribution Map.',
              F('Use Customer Service to demonstrate the distinction between the functional contribution, human performance and organisational enablement.') + ' ' + F('Emphasise that the enterprise target covers all customer enquiries, while the employee resolution measure covers enquiries within employees’ control.') + ' ' + F('Then examine the remaining functions, highlighting how their contributions support the same enterprise result.'), ''] + ml +
          [F('Performance targets are illustrative.') + ' ' + F('The purpose is to establish sufficient evidence to determine whether the required contribution is being delivered and whether the organisation is enabling it.'), '',
           '§4. Examine the measures.',
           ask('Why do we measure both employee incident-response performance and system availability in IT?') + ' ' + F('Guide participants to recognise that employees may fulfil their responsibilities while the underlying organisational capability remains inadequate.'), '',
           '§5. Consolidate the learning.',
           says('The Contribution Map gives us an integrated view of what must be delivered and enabled across the organisation.', 'It establishes the performance requirements against which contributions can be assessed.')])

s = new['map_read']
title(s, '4.1 · Interpreting the Contribution Map')
cards3(s, ['1 · Function', '2 · People (HOS)', '3 · Organisation (ESRG)'],
       ['What must the function deliver?', 'What must individuals and teams deliver?', 'What conditions must the organisation provide?'])
set_notes(s, ['§4.1.2 — INTERPRETING THE CONTRIBUTION MAP', '', TEACH, '',
              '§1. Teach the three elements, one card at a time.',
              says('The Contribution Map establishes the relationship between three elements.'),
              '1 · Functional contribution: what must the function deliver? ' + F('Each function has a defined contribution to the enterprise key result.'),
              '2 · Human Operating System: what must individuals and teams deliver? ' + F('The map identifies the results and behaviours required from people, supported by measurable performance expectations wherever practical.'),
              '3 · Organisational Operating System: what conditions must the organisation provide? ' + F('The map identifies the relevant ESRG requirements necessary to enable the human contribution.'), '',
              '§2. Close on the ESRG measures.',
              says('Only the ESRG elements that materially affect delivery are included.', 'Some require numerical performance measures, such as system availability or staffing coverage.', 'Others require evidence that an arrangement or capability exists and is functioning, such as documented approval authority or established escalation procedures.')])

s = new['faces_ex']
title(s, '4.2 · Managing the contributions through FACES')
icon_rows(s, [('F  ', 'Foster: commitments and obstacles discussed openly'), ('A  ', 'Align: tied to the enterprise Key Result'), ('C  ', 'Transparency: people and constraints visible'),
              ('E  ', 'Accountability: owners with the authority'), ('S  ', 'Development: capability and conditions improve')])
fx = []
for r in C.FACES_FN_ROWS:
    letter, name = re.match(r'<strong>([A-Z]) &mdash; ([^<]+)</strong>', r[0]).groups(); cs, it, op, hr = r[1:]
    fx += ['%s · %s.' % (letter, name), 'Customer Service: ' + F(cs), 'IT: ' + F(it), 'Operations: ' + F(op), 'HR / Workforce Planning: ' + F(hr), '']
INTERP = [('Foster', 'Performance commitments and delivery obstacles are discussed openly, with consistent follow-through from employees and leaders.'),
          ('Align', 'Human contributions and organisational conditions are connected to the requirements of the enterprise key result.'),
          ('Create Measurement Transparency', 'Performance evidence makes both employee contributions and organisational constraints visible.'),
          ('Embed Shared Accountability', 'Responsibilities for delivering contributions and resolving constraints are assigned to the people and functions with the relevant authority.'),
          ('Support Continuous Development', 'Performance findings lead to improvements in employee capability and organisational operating conditions, with subsequent results used to evaluate progress.')]
set_notes(s, ['§4.2 — MANAGING THE CONTRIBUTIONS THROUGH FACES (Applying FACES across the four functions)', '', TEACH, '',
              '§1. Connect FACES to the Contribution Map.',
              says('The Contribution Map™ has established the contributions required from Customer Service, IT, Operations and Workforce Planning, including the expected human performance and the organisational conditions necessary to enable delivery.', 'The following example demonstrates how each FACES function is achieved in practice across the four business functions.'),
              says('In the first part of this example, we established what each function must contribute and what the organisation must enable.', 'We now examine how these contributions are managed.', 'FACES provides the five management functions required to sustain delivery.'),
              says('The next requirement is to manage these contributions consistently throughout execution.', 'FACES provides the five management functions through which performance expectations are aligned, evidence is examined, accountability is maintained and improvement is supported.'), '',
              '§2. Demonstrate how FACES operates across functions.',
              F('Begin with Customer Service and work through the five FACES functions.') + ' ' + F('Then select IT to demonstrate how the same management functions apply to different performance responsibilities.') + ' ' + F('Emphasise that FACES applies to both employee contributions and organisational enabling conditions.'), ''] + fx +
          ['§3. Examine shared accountability.',
           ask('If Customer Service employees are meeting their performance commitments but system outages continue to delay customer resolution, who must take responsibility for the performance gap?') + ' ' + F('Guide participants to distinguish employee accountability from the organisational responsibility for system reliability.'), '',
           '§4. Connect management action to improvement.',
           ask('What should happen when performance reviews repeatedly identify the same approval delays or staffing shortages?') + ' ' + F('Emphasise that performance management must lead to appropriate decisions, assigned responsibilities and follow-through.'), '',
           '§5. Interpret the worked example (the five rows on the slide).',
           'Say: "FACES operates through five connected disciplines."'] + ['%s: %s' % (a, F(b)) for a, b in INTERP] +
          [says('Together, these disciplines ensure that the contributions established in the ' + SR + ' are actively managed towards the enterprise result.'), '',
           '§6. Consolidate the learning.',
           says('The Contribution Map establishes the contributions, performance measures and enabling conditions required for enterprise results.', 'FACES establishes the management disciplines through which those contributions are aligned, evaluated, acted upon and continuously strengthened.')])

# ═══ Section 5 ═══
opener(S[22], C.S5_TITLE, C.S5_SLO)
set_notes(S[22], opener_notes(5, 'Application — In Practice', C.S5_TITLE, tx(C.S5_HERO_F), C.S5_SLO,
                              ['§1. Frame the application.',
                               'Say: "You have already established your capstone organisation’s strategic direction, strategic objectives and Key Results. In this application, your group uses that work to develop two connected components of your Enterprise Performance Management System: Activity 1, the ' + SR + ', and Activity 2, the FACES application."',
                               F('Order: both activities are group work and Capstone work.') + ' ' + F('Activity 1 needs the group’s confirmed Enterprise OKRs from Unit 3, part 4.2.') + ' ' + F('Activity 2 uses the business functions the group names in Activity 1.'),
                               F('How to guide the work: brief each activity from the worked example in Section 4, then let the groups work through the steps on their own pages.') + ' ' + F('One member acts as scribe; every member types the agreed entries into their own page.')]))

s = S[26]  # Activity 1
title(s, '5.1 · Activity 1: Build your ' + SR)
cards3(s, ['Step 1 · Key Result', 'Step 2 · Functions', 'Six fields'],
       ['Choose an Objective, then one of its Key Results.', 'Name three or four functions that must contribute.', 'Contribution, people and ESRG, each with a measure.'])
set_runs(shp(s, 'Text 14'), 'Group work · Capstone work')
set_notes(s, ['§5.1 — ACTIVITY 1: BUILD YOUR S2R® CONTRIBUTION MAP™ (CAPSTONE WORK)', '', TEACH, '',
              '§1. Brief the activity.',
              F('Briefing: each group selects one enterprise Key Result from its Unit 3 Objectives and Key Results that requires contributions from several business functions.') + ' ' + F('It then identifies three or four business functions and completes the six fields for each.') + ' ' + F('Point the groups to the field guidance and to the worked example in part 4.1.'), '',
              '§2. Step 1 · Select the enterprise Key Result.',
              F('The group’s Objectives and Key Results from Unit 3 are in two drop-down lists on each participant’s page: the Objective first, then its Key Results.') + ' ' + F('For the purposes of the programme, the group selects one enterprise Key Result from its capstone strategy that requires contributions from several business functions.'), '',
              '§3. Step 2 · Teach the six fields.'] + ['%s: %s' % (tx(a), F(b)) for a, b in C.FIELD_GUIDE] +
          ['Say: "Identify three or four business functions that must contribute to the selected Key Result. For each function, complete the six fields."',
           'The template on each page asks, for each function: ' + ' '.join(F(q) for q in ['What must this function deliver?', 'What must individuals or teams deliver?', 'How will their contribution be measured?', 'What Environment, Systems, Resources or Governance conditions are required?', 'How will the effectiveness or readiness of these conditions be assessed?']), '',
           '§4. Read the completion check.', 'Before proceeding, the group confirms that:'] + ['%d. %s' % (i + 1, F(c)) for i, c in enumerate(C.CHECK1)] +
          ['Output 1: the completed ' + SR + '.', '',
           PORTAL, 'Participants work as a group in Activity 1 of Section 5 of the participant file. This is Capstone work. The group agrees each entry and one member acts as scribe. Every member then types the agreed entries into their own page, in the session or after it.',
           '1. ' + F('The group chooses an Objective from its Unit 3 OKRs, then one of its Key Results that requires contributions from several business functions, from the two drop-down lists on each page.'),
           '2. ' + F('The group identifies three or four business functions that must contribute to that Key Result.'),
           '3. ' + F('For each function, the group completes the six fields: business function, required functional contribution, Human Operating System required contribution and performance measure, required ESRG contribution and its measure or evidence.'),
           '4. ' + F('The group reads the completion check, then goes to Activity 2.')])

s = S[24]  # Activity 2 (Carol's FACES step slide)
title(s, '5.2 · Activity 2: Apply FACES to your Contribution Map™')
icon_rows(s, [('F  ', 'Foster a Results Culture'), ('A  ', 'Align Performance'), ('C  ', 'Create Measurement Transparency'),
              ('E  ', 'Embed Shared Accountability'), ('S  ', 'Support Continuous Development')])
set_notes(s, ['§5.2 — ACTIVITY 2: APPLY FACES TO YOUR CONTRIBUTION MAP™ (CAPSTONE WORK)', '', TEACH, '',
              '§1. Brief the activity.',
              F('Briefing: each group applies FACES to the contributions it mapped in Activity 1.') + ' ' + F('For each business function it describes one practical management arrangement for each FACES function.') + ' ' + F('Point the groups to the application questions and to the worked example in part 4.2.'), '',
              F('The group has established the contributions required to achieve its enterprise Key Result.') + ' ' + F('It now applies FACES to determine how these contributions will be managed.') + ' ' + F('The business functions named in Activity 1 are shown in the FACES template on each participant’s page.'), '',
              '§2. Step 3 · Teach the application question for each function.'] +
          ['%s · %s: %s The template asks: %s' % (q[0], q[1], F(q[2]), F(q[3])) for q in C.FACES_Q] + ['',
           '§3. Read the completion check.', 'Before confirming, the group checks that:'] + ['%d. %s' % (i + 1, F(c)) for i, c in enumerate(C.CHECK2)] +
          ['Say: "%s"' % ' '.join(F(c) for c in C.CLOSE), '',
           PORTAL, 'Participants work as a group in Activity 2 of Section 5 of the participant file. This is Capstone work. The group agrees each entry and one member acts as scribe. Every member then types the agreed entries into their own page, in the session or after it.',
           '1. ' + F('For each business function named in Activity 1, the group describes one practical management arrangement for each of the five FACES functions.'),
           '2. ' + F('The group reads the completion check.'),
           '3. ' + F('Each participant reads the record and selects Confirm. Print gives a copy.'),
           F('The confirmed record feeds the team’s Capstone Blueprint.'), '',
           '§AFTER THE SESSION:', 'The record reaches you with each participant’s submission, under the heading of Section 5. The confirmed record feeds the Unit 6 boxes of the team’s Capstone Blueprint: the Contribution Map and the FACES application.'])

s = S[27]  # closing
set_runs(shp(s, 'Text 2'), tx(C.CLOSING_Q))
set_notes(s, ['§UNIT SUMMARY (UNIT 6 SYNTHESIS)', '', TEACH, '',
              '§1. Recap the unit, section by section.'] +
          ['- %d · %s: %s' % (i + 1, a.split(' — ')[0], tx(b)) for i, (a, t_, b) in enumerate(C.SUMMARY[:4])] +
          ['- 5 · Application: ' + tx(C.SUMMARY_F5), '',
           '§2. Ask the closing question on the slide.',
           'Ask: "%s" Take an answer from each group.' % tx(C.CLOSING_Q), '',
           '§3. Close the unit.',
           'Say: "A performance management system that only measures has not yet become a management system. The moment measurement becomes the basis for a decision, a conversation, and an accountable intervention — that is the moment execution begins."', '',
           '§4. Say what happens on the portal.',
           'Say: "You now complete the unit on the portal: your four personal reflections and your group’s record in Section 5. When your page is complete, you submit the unit."', '',
           PORTAL, 'Participants complete every part of the participant file and submit the unit to the facilitator from the end of Section 5.', '',
           '§AFTER THE SESSION:', 'Each submission reaches you in the facilitator dashboard. The personal reflections go to each participant’s Learning Portfolio; the group’s confirmed record feeds the team’s Capstone Blueprint.'])

# ═══ order: delete the slides whose content left; place the new slides ═══
ORDER = [S[0], S[1], S[2], S[3], S[4], S[5], new['two_sys'], new['esrg'], new['interact'], new['cm_logic'], new['cm_connect'], S[6], new['faces_two'], new['faces_short'], S[10],
         S[11], S[12], new['bridges_weak'], S[13], new['acc_two'],
         S[16], S[17], new['hot'], S[19], new['refl3'],
         S[20], new['map_ex'], new['map_read'], new['faces_ex'],
         S[22], S[26], S[24], S[27]]
DEL = [S[i] for i in (7, 8, 9, 14, 15, 18, 21, 23, 25)]
sldIdLst = prs.slides._sldIdLst
by_id = {int(el.get('id')): el for el in sldIdLst}
oid = {id(s): s.slide_id for s in ORDER + DEL}
for d in DEL:
    el = by_id[oid[id(d)]]; rId = el.get(qn('r:id'))
    prs.part.drop_rel(rId); sldIdLst.remove(el)
for el in list(sldIdLst): sldIdLst.remove(el)
for s in ORDER: sldIdLst.append(by_id[oid[id(s)]])
assert len(prs.slides) == len(ORDER) == 33
prs.save(out)
print('deck written:', out, len(ORDER), 'slides')
