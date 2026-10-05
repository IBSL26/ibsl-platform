"""Place journal insights inside the presenter notes of a unit deck (no stand-alone block).

Usage: python3 embed_insights.py <deck_in.pptx> <spec.json> <deck_out.pptx> [report.json]

Spec entries (one per insight), each with "slide" and "after" (the end of the line it follows):
  - an insight already in the notes:  "find" = start of its old "Label: text" paragraph;
    optional "title" (title line) and "drop" (a pointer sentence to remove);
  - a new insight:  "text" and "source", optional "title".
Each insight is written as: title line (if any), text, "Source: ..." line.
The old block and its "Professional insights" heading are removed. Nothing else in the notes
changes; the script checks that every other line is still there, in the same order, and stops if not.
Slide text is never touched. Generalised from embed_insights_u01.py (Unit 1).
"""
import json, re, sys, copy
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

HEAD = re.compile(r"^professional insights?:?$", re.I)
T, BR, PPR, END = qn("a:t"), qn("a:br"), qn("a:pPr"), qn("a:endParaRPr")
NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

def segs(p):
    out = []
    for ch in p:
        if ch.tag == BR:
            out.append((ch, "\v"))
        elif ch.find(T) is not None:
            out.append((ch, ch.find(T).text or ""))
    return out

def ptxt(p):
    return "".join(t for _, t in segs(p))

def blank(p):
    return ptxt(p).strip() == ""

def lines(txt):
    return [l.strip() for l in re.split(r"[\n\v]", txt) if l.strip()]

def del_range(p, a, b):
    pos = 0
    for ch, t in segs(p):
        s, e = pos, pos + len(t)
        pos = e
        if e <= a or s >= b:
            continue
        if ch.tag == BR:
            p.remove(ch); continue
        keep = t[:max(0, a - s)] + t[max(0, b - s):]
        if keep == "":
            p.remove(ch)
        else:
            ch.find(T).text = keep

def split_at(p, off):
    newp = etree.SubElement(p.getparent(), qn("a:p")); p.addnext(newp)
    if p.find(PPR) is not None:
        newp.append(copy.deepcopy(p.find(PPR)))
    pos, moving = 0, False
    for ch in list(p):
        if ch.tag == PPR:
            continue
        if ch.tag == END:
            newp.append(ch); continue
        t = "\v" if ch.tag == BR else (ch.find(T).text or "" if ch.find(T) is not None else "")
        s, e = pos, pos + len(t)
        pos = e
        if moving or s >= off:
            moving = True; newp.append(ch)
        elif e > off:
            right = copy.deepcopy(ch)
            ch.find(T).text = t[:off - s]; right.find(T).text = t[off - s:]
            newp.append(right); moving = True
    return newp

def mk(after, text):
    x = f'<a:p {NS}><a:endParaRPr lang="en-GB" dirty="0"/></a:p>' if text == "" else f'<a:p {NS}><a:r><a:rPr lang="en-GB" dirty="0"/><a:t/></a:r></a:p>'
    el = etree.fromstring(x)
    if text:
        el.find(qn("a:r")).find(T).text = text
    after.addnext(el)
    return el

def paras(tb):
    return [c for c in tb if c.tag == qn("a:p")]

def do_slide(n, slide, entries, report):
    tb = slide.notes_slide.notes_text_frame._txBody
    before = lines("\n".join(ptxt(p) for p in paras(tb)))
    mine, marked = [], []
    ps = paras(tb)
    for e in entries:
        if "find" not in e:
            e["_body"], e["_src"] = e["text"].strip(), "Source: " + e["source"].strip()
            continue
        hit = [i for i, p in enumerate(ps) if ptxt(p).startswith(e["find"])]
        if len(hit) != 1:
            sys.exit(f"Slide {n}: '{e['find']}' found {len(hit)} times")
        i = hit[0]
        src = ptxt(ps[i + 1]).strip()
        if not src.startswith("Source: Journal"):
            sys.exit(f"Slide {n}: no source line under '{e['find']}'")
        body = ptxt(ps[i])[len(e["find"]):].strip()
        mine += lines(ptxt(ps[i])) + [src]
        if e.get("drop"):
            if e["drop"] not in " " + body:
                sys.exit(f"Slide {n}: drop text not found in '{e['find']}'")
            body = (" " + body).replace(e["drop"], "").strip()
        e["_body"], e["_src"] = body, src
        marked += [ps[i], ps[i + 1]]
    if marked:
        first = min(ps.index(m) for m in marked)
        head = None
        for j in range(first - 1, -1, -1):
            t = ptxt(ps[j])
            if t.strip() == "":
                continue
            last = lines(t)[-1]
            if HEAD.match(last):
                head = ps[j]; mine.append(last)
                if lines(t) == [last]:
                    marked.append(ps[j])
                else:
                    cut = t.rstrip().rfind(last)
                    del_range(ps[j], len(t[:cut].rstrip("\n\v ")), len(t))
            break
        if head is None:
            sys.exit(f"Slide {n}: no 'Professional insights' heading above the entries")
        idx = sorted(ps.index(m) for m in marked)
        lo, hi = idx[0], idx[-1]
        for p in ps[lo:hi + 1]:
            if p in marked or blank(p):
                tb.remove(p)
            else:
                sys.exit(f"Slide {n}: unexpected text inside the insights block: {ptxt(p)[:60]}")
        ps = paras(tb)
        start = lo
        while start > 0 and blank(ps[start - 1]):
            start -= 1
        end = lo
        while end < len(ps) and blank(ps[end]):
            end += 1
        run = ps[start:end]
        nothing_after = end == len(ps) or ptxt(ps[end])[:1] in ("\n", "\v")
        for p in (run if nothing_after else run[1:]):
            tb.remove(p)
    last_at, added = {}, []
    for e in entries:
        ps = paras(tb)
        if e["after"] in last_at:
            anchor = last_at[e["after"]]
        else:
            hit = [(p, ptxt(p)) for p in ps if e["after"] in ptxt(p) and p not in added]
            total = sum(h[1].count(e["after"]) for h in hit)
            if total != 1:
                sys.exit(f"Slide {n}: anchor found {total} times: '{e['after']}'")
            anchor, t = hit[0]
            end = t.index(e["after"]) + len(e["after"])
            rest = t[end:]
            if rest.strip(" \n\v") == "":
                del_range(anchor, end, len(t))
            elif rest.lstrip(" ")[:1] in ("\n", "\v"):
                k = end + len(rest) - len(rest.lstrip(" \n\v"))
                # keep the indent of the next line: only line breaks (and spaces before them) go
                k = end + len(re.match(r"[ \n\v]*[\n\v]", rest).group(0))
                del_range(anchor, end, k)
                split_at(anchor, end)
            else:
                sys.exit(f"Slide {n}: anchor is in the middle of a line: '{e['after']}' + '{rest[:30]}'")
        block = [""] + ([e["title"]] if e.get("title") else []) + [e["_body"], e["_src"], ""]
        for txt in block:
            anchor = mk(anchor, txt); added.append(anchor)
        last_at[e["after"]] = added[-2]
    for p in [a for a in added if blank(a)]:
        ps = paras(tb)
        i = ps.index(p)
        nxt = ps[i + 1] if i + 1 < len(ps) else None
        prv = ps[i - 1] if i > 0 else None
        if nxt is None or blank(nxt) or ptxt(nxt)[:1] in ("\n", "\v") or (prv is not None and blank(prv)):
            tb.remove(p)
    after = lines("\n".join(ptxt(p) for p in paras(tb)))
    def minus(seq, take):
        seq = list(seq)
        for x in take:
            if x not in seq:
                sys.exit(f"Slide {n}: check failed, line missing: {x[:70]}")
            seq.remove(x)
        return seq
    new_lines = []
    for e in entries:
        new_lines += ([e["title"]] if e.get("title") else []) + [e["_body"], e["_src"]]
    if minus(before, mine) != minus(after, new_lines):
        sys.exit(f"Slide {n}: check failed, the other lines of the notes changed")
    if any(HEAD.match(l) for l in after):
        sys.exit(f"Slide {n}: a 'Professional insights' heading is still there")
    flat = []
    for p in paras(tb):
        for l in (lines(ptxt(p)) or [""]):
            flat.append((l, p in added))
    for e in entries:
        i = next(k for k, (l, a) in enumerate(flat) if a and l == e["_body"])
        s = i - 1 if e.get("title") else i
        prev = [l for l, a in flat[:s] if l][-1:]
        nxt = [l for l, a in flat[i + 2:] if l][:1]
        report.append({"slide": n, "new": "find" not in e, "title": e.get("title", ""), "body": e["_body"], "source": e["_src"],
                       "before": prev, "after": nxt, "dropped": e.get("drop", "").strip()})

def main(deck, spec, out, rep=None):
    d = json.load(open(spec, encoding="utf-8"))
    prs = Presentation(deck)
    slides = list(prs.slides)
    report, by = [], {}
    for e in d["entries"]:
        by.setdefault(e["slide"], []).append(e)
    for n in sorted(by):
        do_slide(n, slides[n - 1], by[n], report)
    prs.save(out)
    if rep:
        json.dump(report, open(rep, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{len(report)} insights placed on {len(by)} slides ({sum(r['new'] for r in report)} new) -> {out}")

if __name__ == "__main__":
    main(*sys.argv[1:5])
