# -*- coding: utf-8 -*-
"""Unit 9 · the two performance views (16 October 2026): Carol's deck brought into line with the PREVIEW pages.
python3 deck_u09pv.py <her deck cf941f0d…> <facilitator page out\\unit3_m2_lens8_f.html> <pictures folder> <out deck>
Slides are reused from her deck (same design, same icons), retitled and refilled; old 1.2, 3.2 and the Collective Intelligence
Challenge slide come out; new slides are copies of her own slides. Notes: heading, teaching notes (the facilitator detail moved off
the page), the page content, what participants complete, and her source detail and professional insights where they still apply."""
import sys, os, re, copy, html
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

SRC, FPAGE, PICS, OUT = sys.argv[1:5]
prs = Presentation(SRC)
S = list(prs.slides)              # S[0] is slide 1 of her deck
O = {i + 1: s for i, s in enumerate(S)}
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

# ── page text (facilitator page) for the "Content" block of the notes ──
FT = open(FPAGE, encoding='utf-8').read()
def plain_lines(h):
    h = re.sub(r'<script.*?</script>|<style.*?</style>', '', h, flags=re.S)
    h = re.sub(r'<(br|/p|/div|/li|/h\d|/tr|/summary|/label)[^>]*>', '\n', h)
    h = re.sub(r'</t[dh]>', ' · ', h)
    t = html.unescape(re.sub(r'<[^>]+>', '', h))
    return [re.sub(r'\s+', ' ', l).strip(' ·') for l in t.split('\n') if l.strip(' ·\r\t ')]
def div_end(t, i):
    d = 0
    for m in re.compile(r'<div\b|</div>').finditer(t, i):
        d += 1 if m.group(0) == '<div' else -1
        if d == 0: return m.end()
def part(title_start):
    k = FT.index('<div class="acc-t">' + title_start); s = FT.index('<div class="acc-b"><div class="cb">', k)
    body = FT[s:div_end(FT, s)]
    body = re.sub(r'<div class="fac-note-full[^"]*">.*?</div>\s*</div>', '', body, flags=re.S)  # guidance boxes go to TEACHING NOTES
    return plain_lines(body)
def guide(label):
    k = FT.find('&#9670; ' + label)
    if k < 0: k = FT.index('&#9670; ' + label.replace('—', '&mdash;'))
    s = FT.rindex('<div class="fac-note-full', 0, k)
    return plain_lines(FT[s:div_end(FT, s)])[1:]

# ── notes ──
BAD = re.compile(r'\b(COMPASS|minutes?|mins?|Timing|triads?|Clog Clinic|KR/KPI|Brigade|Unclogging Tool|Listen for|Challenge any|Check that|Push back|Watch for)\b')
def old_tail(n, keep_source=True):
    t = O[n].notes_slide.notes_text_frame.text.replace('\x0b', '\n')
    out = []
    for head in (['Source detail (S2R® manuscript):'] if keep_source else []) + ['Professional insights:']:
        i = t.find(head)
        if i < 0: continue
        j = min([x for x in (t.find(h2, i + 1) for h2 in ('Source detail (S2R® manuscript):', 'Professional insights:')) if x > i] + [len(t)])
        paras = [p for p in t[i + len(head):j].split('\n') if p.strip()]
        paras = [p for p in paras if not BAD.search(p)]
        if paras: out += ['', head, ''] + paras
    return out
def set_notes(slide, lines):
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in lines:
        p = etree.SubElement(body, '{%s}p' % A)
        if not line.strip():
            e = etree.SubElement(p, '{%s}endParaRPr' % A); e.set('lang', 'en-GB'); continue
        r = etree.SubElement(p, '{%s}r' % A); rp = etree.SubElement(r, '{%s}rPr' % A); rp.set('lang', 'en-GB')
        etree.SubElement(r, '{%s}t' % A).text = line
def notes(slide, heading, teach=None, content=None, does=None, tail=None):
    L = [heading]
    if teach: L += ['', 'TEACHING NOTES:'] + teach
    if content: L += ['', 'Content:'] + content
    if does: L += ['', 'Participant file — what participants complete:'] + ['- ' + d for d in does]
    if tail: L += tail
    for l in L: assert not BAD.search(l), (heading, l)
    set_notes(slide, L)

# ── slide surgery ──
def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def shp(s, name):
    m = [sh for sh in s.shapes if sh.name == name]; assert len(m) == 1, name; return m[0]
def put(s, name, *runs):
    """Write the runs of the first paragraph (bold run, plain run …) keeping their formatting."""
    tf = shp(s, name).text_frame
    for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
    rs = tf.paragraphs[0].runs
    assert len(rs) >= len(runs), (name, [r.text for r in rs], runs)
    for r, t in zip(rs, runs): r.text = t
    for r in rs[len(runs):]: r._r.getparent().remove(r._r)
def drop(s, names):
    for sh in [x for x in s.shapes if x.name in names]: sh._element.getparent().remove(sh._element)
def move(s, name, dx=0, dy=0, w=None, h=None, left=None, top=None):
    sh = shp(s, name)
    if left is not None: sh.left = left
    if top is not None: sh.top = top
    sh.left += dx; sh.top += dy
    if w is not None: sh.width = w
    if h is not None: sh.height = h
def duplicate(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId] = dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID + 'embed', RID + 'link', RID + 'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    nt, st = dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'), qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag == qn('p:sp')]: nt.append(copy.deepcopy(el))
    return dst
OVAL = ('''<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="900" name="Part"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="868680" cy="868680"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="C6A24C"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="C6A24C"/></a:solidFill></a:ln><a:effectLst/></p:spPr><p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-GB" sz="2400" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:latin typeface="Calibri"/></a:rPr><a:t>1.1</a:t></a:r></a:p></p:txBody></p:sp>''')
def oval(slide, label, left, top):
    el = parse_xml(OVAL); nid = max(int(e.get('id')) for e in slide._element.iter(qn('p:cNvPr'))) + 1
    el.find('.//' + qn('p:cNvPr')).set('id', str(nid)); slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]; sh.left, sh.top = left, top; sh.name = 'Part ' + label
    sh.text_frame.paragraphs[0].runs[0].text = label
def reflection_slide(title, items, note_lines):
    s = duplicate(O[7])
    drop(s, [sh.name for sh in s.shapes if sh.name != 'Text 0'])
    put(s, 'Text 0', title)
    left, right, top0, bottom = Inches(1.85), Inches(12.85), Inches(1.6), Inches(7.2)
    gap = Inches(0.3); width = right - left
    hs = []
    for _, f in items:
        w, h = Image.open(os.path.join(PICS, f)).size; hs.append(width * h / w)
    scale = min(1, (bottom - top0 - gap * (len(items) - 1)) / sum(hs))
    y = top0
    for (lab, f), h in zip(items, hs):
        pw = int(width * scale); ph = int(h * scale)
        pic = s.shapes.add_picture(os.path.join(PICS, f), left, int(y), width=pw); pic.name = 'Reflection ' + lab
        oval(s, lab, Inches(0.6), int(y + ph / 2 - Inches(0.475)))
        y += ph + gap
    set_notes(s, note_lines)
    return s
def list_rows(s, rows, n_have):
    """Fill the icon list of a slide: rows = [(bold, plain)]; extra rows of the template are removed."""
    for i in range(n_have):
        names = ('Shape %d' % (2 * i + 1), 'Image %d' % i, 'Text %d' % (2 * i + 2))
        if i < len(rows):
            b, p = rows[i]
            if p is None: put(s, names[2], b)
            else: put(s, names[2], b, p)
        else: drop(s, names)
def cards(s, title, items, banner):
    put(s, 'Text 0', title)
    for (t, b), (nt, nb) in zip(items, (('Text 3', 'Text 4'), ('Text 7', 'Text 8'), ('Text 11', 'Text 12'))):
        put(s, nt, t); put(s, nb, b)
    if banner: put(s, 'Text 14', banner)
def lower_bodies(s, dy=Emu(380000)):
    for nb in ('Text 4', 'Text 8', 'Text 12'):
        sh = shp(s, nb); sh.top = sh.top + dy; sh.height = sh.height - dy
def section_slide(s, title, slo):
    put(s, 'Text 3', title); put(s, 'Text 6', slo[0]); put(s, 'Text 8', slo[1])

def P(*t): return list(t)
def gl(label): return guide(label)

# ═══ New slides (copies of her slides, made before any of them is edited) ═══
N_los = duplicate(O[21]); N_44 = duplicate(O[11]); N_45 = duplicate(O[18]); N_52 = duplicate(O[7]); N_53 = duplicate(O[11])

# ── 1 · Title
notes(O[1], 'UNIT 9 · STRATEGIC UNCLOGGING', content=['Module 3 · Influence · Unit 9', '', 'Unit intent (Facilitator Guide):'] + gl('Unit Intent') +
      ['', 'Unit overview (facilitator file): When strategy stalls, the failure usually sits in the system that carries the plan. This unit equips leaders to diagnose the blockages preventing strategic intent from flowing into real-world execution — and to intervene with precision using the 7 Strategic Unclogging Levers.'])
# ── 3 · Unit journey (session overview)
notes(O[3], 'FACILITATOR GUIDE · SESSION OVERVIEW', content=['Key Facilitation Questions'] + gl('Key Facilitation Questions') + ['', 'Tone & Watch Points'] + gl('Tone &amp; Watch Points'))

# ═══ Section 1 ═══
notes(O[4], 'SECTION 1 · AWARENESS — WHAT', teach=gl('FACILITATOR GUIDANCE — Awareness') + [
    'The central reframe: a clog is a signal. Clogs are evidence that organisational performance and people performance have not yet adapted to the strategy. Leaders who understand this stop asking "Why are we failing?" and start asking "Where are our people and our organisational conditions not yet ready to carry the strategy?"',
    '1.1 · Build the case before the two sides: the strategy is clear, KPIs are with individuals and teams, the organisational conditions are agreed, and still a Key Result stops moving. Name that as a Strategic Clog, teach the four ABCV elements, then where clogs form: in the tasks individuals execute and in the organisational conditions around them.',
    '1.2 · Teach the two dimensions of performance first. Then open the seven clogs one card at a time: the system reality, the contribution reality, the patterns, the early warnings, the impact and the diagnostic question. Close with the ABCV connection, the three lines of sight and the customer service example.'],
    content=['What is a Strategic Clog?', 'Naming the internal blockages that prevent strategic intent from flowing into execution — and introducing the 7-clog diagnostic map aligned to ABCV Directional Integrity.', '', 'Section learning outcomes:',
             '1. Recognise the systemic blockages that widen the gap between strategic intent and execution.', '2. Identify the observable symptoms of friction in execution.'], tail=old_tail(4))
notes(O[5], '1.1 — WHAT IS A STRATEGIC CLOG? (Core concept)', content=part('1.1 &mdash;'),
      does=['✎ Reflection — Strategic Clog in Your Organisation — Think of a strategy initiative that stalled in your organisation. Was the plan wrong, or was something in the system blocking its flow? What was the blockage? Could it have been named earlier? Which organisational condition did the individuals’ tasks depend on, and did it support their execution?'],
      tail=old_tail(5))
put(O[7], 'Text 0', '1.2 · The 7-Clog Diagnostic Map')
notes(O[7], '1.2 — THE 7-CLOG DIAGNOSTIC MAP (ABCV-aligned · Select each clog to explore)', content=part('1.2 &mdash;'),
      does=['✎ Reflection — Which Clogs Are Active? — Rate each clog in your current strategy portfolio (1 = absent, 5 = active). Then name the clog most active in your organisation, and the organisational condition it sits in.'],
      tail=old_tail(7))
put(N_los, 'Text 0', '1.2.2 · Three performance lines of sight')
list_rows(N_los, [('1  Performance outcome  ', 'What result is falling short?'), ('2  Performance behaviour  ', 'What are people doing, or unable to do?'), ('3  Performance condition  ', 'What is causing or reinforcing it?')], 4)
notes(N_los, '1.2.2 — THE THREE PERFORMANCE LINES OF SIGHT', teach=[
    'Teach the three levels of performance evidence in order. The outcome names the result falling short and the size of the gap. The behaviour names what individuals and teams are doing, or unable to do. The condition names the organisational arrangement, capability, incentive or leadership practice behind that behaviour; the seven clogs are tested here.',
    'Land the reason for the discipline: it keeps leaders from treating every disappointing result as an individual performance failure, and it keeps people accountable for what is within their control.',
    'Walk the customer service example of 1.2.3: 65% of enquiries resolved within 48 hours against 90%; cases transferred without end-to-end ownership; four clogs in the evidence table; the organisational, people and strategic interpretations; the management response.'])
R1 = reflection_slide('Section 1 Reflections', [('1.1', 'ref_1.1.png'), ('1.2', 'ref_1.2.png')], ['SECTION 1 REFLECTIONS', '', 'Participants complete these reflections in their own file after the teaching.',
    '1.1 · Strategic Clog in Your Organisation: a strategy initiative that stalled, the blockage behind it, whether it could have been named earlier, and the organisational condition the individuals’ tasks depended on.',
    '1.2 · Which Clogs Are Active?: a rating of each of the seven clogs, and the clog most active in their organisation with the organisational condition it sits in.'])

# ═══ Section 2 (parked: only the guidance question) ═══
t8 = O[8].notes_slide.notes_text_frame
for p in t8.paragraphs:
    full = ''.join(r.text for r in p.runs)
    if '“Where is the system not yet designed to carry the strategy?”' in full:
        p.runs[0].text = full.replace('“Where is the system not yet designed to carry the strategy?”', '“Where are our people and our organisational conditions not yet ready to carry the strategy?”')
        for r in p.runs[1:]: r.text = ''
put(O[9], 'Text 2', 'Where are our people and our organisational conditions not yet ready to carry the strategy?')
R2 = reflection_slide('Section 2 Reflections', [('2.1', 'ref_2.1.png'), ('2.3', 'ref_2.3.png')], ['SECTION 2 REFLECTIONS', '', 'Participants complete these reflections in their own file after the teaching.',
    '2.1 · Where Is Your System Not Yet Ready?: the element of their operating system most likely to create friction under their current strategic direction, and why.',
    '2.3 · Your Lever: the primary lever and the second lever they would pull for the clog they identified in Awareness, and why this combination.'])

# ═══ Section 3 ═══
section_slide(O[12], 'Where Do Executive Decisions and Practices Contribute to Performance Clogs?', ('Reveal how executive decisions and practices can create or reinforce performance clogs.', 'Examine how decisions within one executive function affect contributions and performance across the organisation.'))
notes(O[12], 'SECTION 3 · EXTRAPOLATING — WHERE', teach=gl('FACILITATOR GUIDANCE — Extrapolating'),
      content=['Where Do Executive Decisions and Practices Contribute to Performance Clogs?', 'Examining how decisions and management practices within executive functions influence people performance, organisational conditions and strategic results.', '', 'Section learning outcomes:',
               '1. Reveal how executive decisions and practices can create or reinforce performance clogs.', '2. Examine how decisions within one executive function affect contributions and performance across the organisation.'], tail=old_tail(12, keep_source=False))
s = O[13]; put(s, 'Text 0', '3.1 · The CXO performance clog contributors')
ROLES = [('CEO  ', 'Enterprise Direction'), ('CFO  ', 'Financial Stewardship'), ('COO  ', 'Operational Discipline'), ('CHRO  ', 'People and Culture'), ('CTO/CIO  ', 'Technology and Information'),
         ('CMO  ', 'Market and Customer Insight'), ('CCO  ', 'Commercial Performance'), ('CPO  ', 'Procurement and Supply'), ('CRO  ', 'Risk and Governance'), ('CSO  ', 'Strategic Integration')]
list_rows(s, ROLES[:7], 7)
for i in range(3):  # three more rows: copies of the last row
    for nm in ('Shape 13', 'Image 6', 'Text 14'):
        el = copy.deepcopy(shp(s, nm)._element); s.shapes._spTree.append(el)
        s.shapes[-1].name = nm + '_x%d' % i
    put(s, 'Text 14_x%d' % i, *ROLES[7 + i])
# two columns of five rows
step = Emu(960000); top = Emu(1450000); colw = Emu(4750000); col2 = Emu(6250000)
rows = [('Shape %d' % (2 * i + 1), 'Image %d' % i, 'Text %d' % (2 * i + 2)) for i in range(7)] + [('Shape 13_x%d' % i, 'Image 6_x%d' % i, 'Text 14_x%d' % i) for i in range(3)]
for k, (sh_, im, tx) in enumerate(rows):
    c, r = divmod(k, 5); x0 = Emu(548640) + (col2 - Emu(548640)) * c; y = top + step * r
    move(s, sh_, left=x0, top=y + Emu(150000)); move(s, im, left=x0 + Emu(122301), top=y + Emu(272301))
    move(s, tx, left=x0 + Emu(809244), top=y, w=colw, h=Emu(789204))
notes(s, '3.1 — THE CXO PERFORMANCE CLOG CONTRIBUTORS (Select each executive role to explore)', teach=[
    'Begin with each participant’s own executive role, then the two roles closest to it. For each card, read the decision or practice, then trace it through the three lines of sight: the result affected, how people respond, and the clog that follows.',
    'Name the clog on each card: CEO, Leadership Avoidance; CFO, Decision Bottlenecks; COO, Structural Misfit; CHRO, Misaligned Metrics & Incentives; CTO/CIO, Technology / Infrastructure Lag; CMO, Structural Misfit; CCO, Misaligned Metrics & Incentives; CPO, Misaligned Metrics & Incentives; CRO, Decision Bottlenecks; CSO, Structural Misfit.',
    'Close with the Closing Insight: executive decisions shape the conditions in which people and teams perform.'],
    content=part('3.1 &mdash;'), does=['✎ Reflection — My Executive Contribution to Performance (Portfolio work): my executive function; one decision or management practice; 1 · Performance Outcome; 2 · Performance Behaviour; 3 · Performance Condition.'],
    tail=old_tail(13, keep_source=False))
R3 = reflection_slide('Section 3 Reflections', [('3.1', 'ref_3.1.png')], ['SECTION 3 REFLECTIONS', '', 'Participants complete this reflection in their own file after the teaching, as Portfolio work.',
    '3.1 · My Executive Contribution to Performance: one decision, policy or management practice within their own executive function that may be affecting another team or function, traced through the Three Performance Lines of Sight to the clog it creates or reinforces.'])

# ═══ Section 4 ═══
section_slide(O[15], 'Cross Leadership Clog Ownership', ('Mobilise leaders across functions to restore performance momentum together.', 'Establish a shared leadership rhythm for surfacing and resolving friction across people and organisational performance.'))
notes(O[15], 'SECTION 4 · INTEGRATION — COLLECTIVE', teach=gl('FACILITATOR GUIDANCE — Integration'),
      content=['Cross Leadership Clog Ownership', 'Building the collective discipline of normalising the unclog, where detecting and clearing performance friction becomes a shared leadership habit.', '', 'Section learning outcomes:',
               '1. Mobilise leaders across functions to restore performance momentum together.', '2. Establish a shared leadership rhythm for surfacing and resolving friction across people and organisational performance.'], tail=old_tail(15, keep_source=False))
s = O[16]; put(s, 'Text 0', '4.1 · Normalising the unclog: from crisis to habit')
list_rows(s, [('REVIEW  ', 'the evidence, and confirm the clog'), ('REALIGN  ', 'ownership, and the changes on both sides'), ('RESTORE  ', 'activate the lever, and verify recovery'), ('Every executive  ', 'makes emerging friction visible')], 4)
notes(s, '4.1 — NORMALISING THE UNCLOG: FROM CRISIS TO HABIT (The collective discipline)', teach=[
    'Teach friction as a normal feature of organisational life. Without a shared discipline, problems are discussed informally, handled in silos or carried from review to review.',
    'Teach the consistent path: REVIEW, REALIGN, RESTORE. Then the collective leadership principle: every executive makes emerging friction visible; the leaders who control the contributing conditions act together, with clear individual ownership of each commitment.'],
    content=part('4.1 &mdash;'), tail=old_tail(16, keep_source=False))
s = O[17]; put(s, 'Text 1', '4.2 · REVIEW: establish the performance reality')
put(s, 'Text 2', 'REVIEW is complete when the executive team agrees on the performance gap, the supporting evidence and the clog requiring intervention.')
put(s, 'Text 3', 'Outcome, behaviour, condition: the Three Performance Lines of Sight.')
notes(s, '4.2 — REVIEW: ESTABLISH THE PERFORMANCE REALITY', teach=[
    'Teach REVIEW as a shared, evidence-based view of the result falling short and the organisational condition behind it, built on the three lines of sight.',
    'Teach the "see something, say something" principle here: the review conversation lets leaders surface friction across functional boundaries, through observable evidence, including where the condition sits in another executive portfolio.'],
    content=part('4.2 &mdash;'), tail=old_tail(17, keep_source=False))
cards(O[18], '4.3 · REALIGN: collective executive ownership',
      [('What must change?', 'The arrangement, practice or decision sustaining the clog.'), ('Who must act together?', 'One coordinating owner, and the contributing executives.'), ('What must be agreed?', 'Decisions, responsibilities, resources and timelines.')],
      'Consider both: people performance and organisational performance.')
lower_bodies(O[18])
notes(O[18], '4.3 — REALIGN: ESTABLISH COLLECTIVE EXECUTIVE OWNERSHIP', teach=[
    'Teach the three leadership decisions in order. Then take the group to both dimensions of performance before the response is final: what must change for the people delivering, and what must change in structures, processes, systems, resources, decision rights or governance.',
    'Close on clear individual accountability inside collective ownership: one leader coordinates, the other executives stay accountable for their agreed contributions.'],
    content=part('4.3 &mdash;'), tail=old_tail(18, keep_source=False))
put(N_44, 'Text 0', '4.4 · RESTORE: activate the Seven Unclogging Levers')
notes(N_44, '4.4 — RESTORE: ACTIVATE THE SEVEN UNCLOGGING LEVERS', teach=[
    'Teach RESTORE as the step that turns the agreed response into organisational action: select the lever, activate the intervention, verify recovery.',
    'Distinguish the REALIGN stage of S2R® (the collective decision and ownership step) from Lever 06 · Realign (the intervention that adjusts metrics and rewards).',
    'Land the closing line: removing a clog requires evidence that the organisational condition has changed, verified through the normal performance review rhythm.'],
    content=part('4.4 &mdash;'))
cards(N_45, '4.5 · Worked example: Review–Realign–Restore',
      [('REVIEW', 'Decision Bottlenecks: routine maintenance purchases wait for senior approval.'), ('REALIGN', 'The CFO coordinates and the COO contributes: revised approval authority and controls.'), ('RESTORE', 'Lever 02 · Streamline, verified through approval time, downtime, delivery and controls.')],
      'Increase customer orders delivered on time from 72% to 95% by year-end.')
notes(N_45, '4.5 — WORKED EXAMPLE: REVIEW–REALIGN–RESTORE (A manufacturing case)', teach=[
    'Walk the example step by step: the evidence table of REVIEW, the leadership decisions of REALIGN with the people and organisational implications, the activation of Streamline in RESTORE and the four verification measures. Close on the Completed Unclogging Record.',
    'Then guide participants to apply the same process to their Capstone organisation in Section 5.'],
    content=part('4.5 &mdash;'))
R4 = reflection_slide('Section 4 Reflections', [('4', 'ref_4.png')], ['SECTION 4 REFLECTIONS', '', 'Participants complete this personal reflection in their own file after the teaching, as Portfolio work.',
    'My Leadership Contribution to Strategic Unclogging: REVIEW (the clog, its effect on people’s performance and on organisational performance); REALIGN (their own contribution to the clog, what they must change and the leaders they need to engage); RESTORE (the Unclogging Lever, the action and the evidence of improvement).',
    '', 'Facilitator Closing Note — Integration'] + gl('Facilitator Closing Note — Integration'))

# ═══ Section 5 ═══
section_slide(O[19], 'The Strategic Unclogging Clinic', ('Classify an original organisational challenge using the Seven Performance Clogs and select an appropriate established Unclogging Lever.', 'Develop a collective executive response through Review–Realign–Restore, addressing people performance and organisational performance.'))
notes(O[19], 'SECTION 5 · APPLICATION — IN PRACTICE', teach=gl('FACILITATOR GUIDANCE — The Clog Clinic'),
      content=['The Strategic Unclogging Clinic', 'Returning to an original work-based challenge, identifying the performance clog, selecting the corresponding Strategic Unclogging Lever and applying Review–Realign–Restore within the Capstone organisation.', '', 'Section learning outcomes:',
               '1. Classify an original organisational challenge using the Seven Performance Clogs and select an appropriate established Unclogging Lever.', '2. Develop a collective executive response through Review–Realign–Restore, addressing people performance and organisational performance.',
               '', 'How Section 5 works:',
               '1. Each group selects one challenge from the list it documented before CESI began and links it to a strategic result (5.1).',
               '2. The group identifies the performance clog (5.2) and selects the Strategic Unclogging Lever (5.3).',
               '3. The group applies Review–Realign–Restore (5.4), then tests its response against the Diagnostic Quality Check (5.5).',
               '4. The group agrees each entry and one member acts as scribe. Every member then types the agreed entries into their own page, in the session or after it.',
               '5. In part 5.6 each member reads the group’s record and selects Confirm. The confirmed record feeds the team’s Capstone Blueprint.'])
cards(O[20], '5.1 · Return to your original work-based challenge',
      [('The challenge', 'One challenge from the list documented before CESI began.'), ('The connection', 'The Objective, and the Key Result or KPI it affects.'), ('The evidence', 'The current result, or a realistic illustrative figure.')],
      'Choose a challenge with a clear connection to disappointing results.')
DO5 = 'Group work, Capstone work: the group agrees each entry, one member acts as scribe, and every member types the agreed entries into their own page.'
notes(O[20], '5.1 — RETURN TO YOUR ORIGINAL WORK-BASED CHALLENGE (Capstone Work)', teach=[
    'Each group returns to the organisational challenges its members documented before CESI began, and selects one that is still relevant to its Capstone organisation and clearly linked to disappointing results.', DO5],
    content=part('5.1 &mdash;'))
s = N_52; put(s, 'Text 0', '5.2 · Identify the performance clog')
list_rows(s, [('01 · Structural Misfit  ', 'Roles, ownership, coordination'), ('02 · Capability Gaps  ', 'Skills, tools, capacity'), ('03 · Decision Bottlenecks  ', 'Authority, approvals, escalation'), ('04 · Misaligned Metrics & Incentives  ', 'Measures and rewards'),
              ('05 · Cultural Resistance  ', 'Beliefs, norms, habits'), ('06 · Technology / Infrastructure Lag  ', 'Systems, information'), ('07 · Leadership Avoidance  ', 'Unresolved decisions, trade-offs')], 7)
notes(s, '5.2 — IDENTIFY THE PERFORMANCE CLOG (Capstone Work)', teach=[
    'The challenge describes what has been experienced; the clog names the organisational condition believed to be constraining performance. The group selects one clog and states the observations behind it.',
    'This is the initial diagnosis. The group tests it in REVIEW and corrects the clog and the lever if the evidence points elsewhere.'], content=part('5.2 &mdash;'))
put(N_53, 'Text 0', '5.3 · Select the Strategic Unclogging Lever')
notes(N_53, '5.3 — SELECT THE STRATEGIC UNCLOGGING LEVER (Capstone Work)', teach=[
    'One clog may have more than one suitable lever; the group selects the primary lever for the specific condition in its case and states how it addresses the clog.',
    'Selecting the lever does not start the intervention: the lever is tested and activated through Review–Realign–Restore.'], content=part('5.3 &mdash;'))
s = O[21]; put(s, 'Text 0', '5.4 · Apply Review–Realign–Restore')
list_rows(s, [('REVIEW  ', 'test the diagnosis against the evidence'), ('REALIGN  ', 'agree what must change, and who acts'), ('RESTORE  ', 'apply the lever, with owners and dates'), ('Evidence  ', 'of recovery, on both sides')], 4)
notes(s, '5.4 — APPLY REVIEW–REALIGN–RESTORE (Capstone Work)', teach=[
    'REVIEW: the group uses the three lines of sight on its own challenge and checks whether the evidence confirms the clog selected in 5.2.',
    'REALIGN: as the simulated executive team, the group sets out what must change for people performance and for organisational performance, and agrees the coordinating owner, the contributing functions, the agreed change and each executive’s contribution.',
    'RESTORE: the group specifies the intervention, the action and completion, the evidence on both sides and when the executive team will review progress.'], content=part('5.4 &mdash;'))
s = O[23]; put(s, 'Text 0', '5.5 · Diagnostic quality check')
list_rows(s, [('01  ', 'Linked to a specific strategic result'), ('02  ', 'The clog names the condition behind the symptom'), ('03  ', 'The lever fits; both sides are considered'), ('04  ', 'Specific, owned, time-bound and evidenced')], 4)
notes(s, '5.5 — DIAGNOSTIC QUALITY CHECK (Before the response is finalised)', teach=[
    'The group tests its own response against the four criteria and returns to the relevant step where a criterion is not met.'], content=part('5.5 &mdash;'))
cards(O[24], 'Section 5 · Capstone portfolio output',
      [('Challenge and clog', 'The original challenge, the clog, and the reason for the diagnosis.'), ('The lever', 'The selected lever and its intended purpose.'), ('Review–Realign–Restore', 'With executive ownership and evidence of recovery.')],
      None)
lower_bodies(O[24])
notes(O[24], 'SECTION 5 · CAPSTONE PORTFOLIO OUTPUT (before the exercise)', teach=['Show the group what it will submit before the exercise begins: one Strategic Unclogging Clinic Record with the original challenge, the clog and the reason, the lever and its purpose, and the Review–Realign–Restore responses with executive ownership and evidence of recovery.', 'At the end, in part 5.6, each member reads the group’s record on their own page and selects Confirm. The confirmed record feeds boxes 9A to 9E of the team’s Capstone Blueprint.', '', 'Facilitator Closing Note — Application'] + gl('Facilitator Closing Note — Application'),
      content=['The group submits one Strategic Unclogging Clinic Record containing:'] + ['%d. %s' % (i + 1, x) for i, x in enumerate(['The original work-based challenge selected from its pre-CESI preparation.', 'The identified Seven Performance Clog and reason for the diagnosis.', 'The selected Seven Strategic Unclogging Lever and its intended purpose.', 'The completed REVIEW, REALIGN and RESTORE responses, with executive ownership and evidence of recovery.'])] + ['The completed record is incorporated into the existing Capstone portfolio.'])

# ── 25 · Unit Summary
k = FT.index('<div class="acc-t">Unit Summary</div>'); s0 = FT.index('<div class="acc-b"><div class="cb">', k)
summ = plain_lines(FT[s0:div_end(FT, s0)])
notes(O[25], 'UNIT SUMMARY (Unit 9 synthesis)', content=summ, tail=old_tail(25))

# ═══ Order: drop old 1.2 (6), 3.2 (14) and the Collective Intelligence Challenge (22) ═══
order = [O[1], O[2], O[3], O[4], O[5], O[7], N_los, R1, O[8], O[9], O[10], O[11], R2, O[12], O[13], R3,
         O[15], O[16], O[17], O[18], N_44, N_45, R4, O[19], O[24], O[20], N_52, N_53, O[21], O[23], O[25]]
lst = prs.slides._sldIdLst
sid = {id(s): s.slide_id for s in list(O.values()) + [N_los, N_44, N_45, N_52, N_53, R1, R2, R3, R4]}
by_id = {int(el.get('id')): el for el in list(lst)}
for el in list(lst): lst.remove(el)
for s in order: lst.append(by_id[sid[id(s)]])
for s in (O[6], O[14], O[22]):
    prs.part.drop_rel(by_id[sid[id(s)]].get(RID + 'id'))
prs.save(OUT)
print('saved', OUT, len(order), 'slides')
