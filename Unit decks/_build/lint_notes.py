"""Checks on the presenter notes of a deck (Carol's standing rules). Usage: python3 lint_notes.py <deck.pptx> [first_slide]"""
import re, sys
from pptx import Presentation
p = Presentation(sys.argv[1]); first = int(sys.argv[2]) if len(sys.argv) > 2 else 1
rules = [
 ('contrast', re.compile(r'\brather than\b|\binstead of\b|\bin place of\b|—\s*not\b|, not\b|\bnot\b[^.?!\n]{0,60}\bbut\b|\bis not\b[^.?!\n]{0,40}[—;:] ?it is', re.I)),
 ('lens', re.compile(r'\blens(es)?\b', re.I)),
 ('time', re.compile(r'\b\d+\s*(–|-|to)?\s*\d*\s*(minutes?|mins?|hours?|hrs?)\b', re.I)),
 ('prework', re.compile(r'pre-?work|before the session|in advance of the session', re.I)),
 ('portal-steer', re.compile(r'\b(open|go to|turn to|log ?in to|click|select|expand)\b[^.\n]{0,40}\b(portal|participant file|your page|the page)\b', re.I)),
 ('automation', re.compile(r'automat|auto-?generat|\bAI\b', re.I)),
 ('mark', re.compile(r'§')),
 ('S2R', re.compile(r'Strategy2Results(?!®)|\bS2R(?!®)')),
 ('american', re.compile(r'\b(organiz|recogniz|prioritiz|behavior|program\b|center\b|analyz|realiz|crystalliz|optimiz)', re.I)),
 ('SIP', re.compile(r'\bSIP\b')),
 ('removed-content', re.compile(r'\brat(e|ed|ing|ings) (their|your|the) organisation|\(1–5\)|1 to 5|four numbers|private rating|Critical Insight|Probing question|maturity|Headline|Storyline', re.I)),
 ('loose-block', re.compile(r'^(CONTENT:|SOURCE DETAIL|FACILITATOR GUIDANCE|FACILITATOR SCRIPT)')),
 ('dimension-misuse', re.compile(r'(four|4) SiP dimensions|SiP dimension', re.I)),
]
n = 0
for i, s in enumerate(p.slides, 1):
    if i < first: continue
    for ln, line in enumerate(s.notes_slide.notes_text_frame.text.replace('\x0b', '\n').split('\n'), 1):
        for name, rx in rules:
            m = rx.search(line)
            if m:
                n += 1; a = max(0, m.start() - 70)
                print(f'[{name}] slide {i} line {ln}: …{line[a:m.end() + 50]}')
print('hits:', n)
