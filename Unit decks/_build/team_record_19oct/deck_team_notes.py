# -*- coding: utf-8 -*-
"""Deck presenter notes, Units 1 and 3 to 9: the shared group record (19 Oct 2026).
The group chooses its scribe for each assignment; the entries appear on every member's page; each member selects Agree.
Only the notes change; slides stay as they are. usage: python3 deck_team_notes.py <deck in> <deck out> [--dry]"""
import sys
from pptx import Presentation
F = 'The group chooses its scribe for this work. The scribe types what the group agrees, and the entries appear on every member’s page.'
R = [
 ('The group agrees each entry and one member acts as scribe. Every member then types the agreed entries into their own page of the participant file, in the session or after it.', F),
 ('The group agrees each entry and one member acts as scribe. Every member then types the agreed entries into their own page, in the session or after it.', F),
 ('The group agrees each entry and one member acts as scribe. Every member types the agreed entries into their own page, then selects Confirm. The confirmed record feeds the team’s Capstone Blueprint.',
  F + ' Each member then selects Agree. When every member has agreed, the record feeds the team’s Capstone Blueprint.'),
 ('Say: "This is group work. Your group agrees each entry. One member acts as scribe. Every member then types the agreed entries into their own page, in the session or after it."',
  'Say: "This is group work. Your group agrees each entry and chooses its scribe for this work. The scribe types, and the entries appear on every member’s page."'),
 ('is group work for the Capstone. The group agrees each entry and one member acts as scribe.', 'is group work for the Capstone. The group chooses its scribe for this work.'),
 ('Every member then types the agreed entries into their own page.', 'The scribe types what the group agrees, the entries appear on every member’s page, and each member then selects Agree.'),
 ('One member acts as scribe and types the agreed wording.', 'The group chooses its scribe, who types the agreed wording.'),
 ('Every member then types the agreed entries into their own page, in the session or after it."', 'The entries appear on every member’s page, and each member then selects Agree."'),
 ('The group agrees each entry and one member acts as scribe.', 'The group chooses its scribe for this work.'),
 ('One member acts as scribe.', 'Your group chooses its scribe for this work.'),
 ('You then select Confirm. This is Capstone work, and your confirmed record feeds your team’s Capstone Blueprint.', 'You then select Agree. This is Capstone work, and when every member has agreed, your record feeds your team’s Capstone Blueprint.'),
 ('Each participant reads the record and selects Confirm.', 'Each participant reads the record and selects Agree.'),
 ('One member acts as scribe; every member types the agreed entries into their own page.', 'The group chooses its scribe; the entries appear on every member’s page.'),
 ('reads the group’s record at the foot of the section and selects Confirm. The confirmed record feeds', 'reads the group’s record at the foot of the section and selects Agree. When every member has agreed, the record feeds'),
 ('read your group’s record at the foot of the section and select Confirm. Your confirmed record feeds', 'read your group’s record at the foot of the section and select Agree. When every member has agreed, your record feeds'),
 ('reads the group’s record and selects Confirm. The confirmed record feeds', 'reads the group’s record and selects Agree. When every member has agreed, the record feeds'),
 ('reads the group’s record on their own page and selects Confirm. The confirmed record feeds', 'reads the group’s record on their own page and selects Agree. When every member has agreed, the record feeds'),
 ('the group agrees each entry, one member acts as scribe, and every member types the agreed entries into their own page.', 'the group agrees each entry and chooses its scribe for this work, and the entries appear on every member’s page.'),
 ('The group agrees each entry, and every member types the agreed entries into their own page.', 'The group agrees each entry; the scribe types it, and it appears on every member’s page.'),
 ('One scribe, agreed entries, every member’s own page.', 'One scribe per assignment; the agreed entries appear on every member’s page.'),
 ('Confirm the Charter record: it becomes Section 1 of your team’s Capstone Blueprint.', 'In the Charter record each member selects Agree: when every member has agreed, it becomes Section 1 of your team’s Capstone Blueprint.'),
 ('The confirmed Keep, Improve, Start and Stop feed the team\'s Capstone Blueprint.', 'When every member has agreed, the Keep, Improve, Start and Stop feed the team\'s Capstone Blueprint.'),
]
def fix(t):
    for a, b in R:
        t = t.replace(a, b)
    return t
src, out = sys.argv[1], sys.argv[2]; dry = '--dry' in sys.argv
p = Presentation(src); n = 0
for i, s in enumerate(p.slides, 1):
    if not s.has_notes_slide: continue
    for para in s.notes_slide.notes_text_frame.paragraphs:
        runs = para.runs
        if not runs: continue
        t = ''.join(r.text for r in runs); new = fix(t)
        if new != t:
            n += 1
            if dry: print(i, '|', t[:120], '\n   ->', new[:160])
            runs[0].text = new
            for r in runs[1:]: r.text = ''
print('paragraphs changed:', n)
if not dry: p.save(out)
