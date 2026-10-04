"""Add the source-manuscript detail block to the presenter notes of a unit deck.

Usage: python3 source_notes.py <deck.pptx> <source/uNN.json>
Run after building a deck with unitNN.js and BEFORE enrich.py.
Safe to re-run at any point: an existing "Source detail" block is removed and
written again, and the block always sits before "Professional insights:".
Existing notes and slide text are not touched.

JSON: {"_check": {"<slide no>": "<opening slide text>"}, "<slide no>": ["line", "", "line", ...]}
An empty string is a blank line between groups.
"""
import copy, json, sys
from pptx import Presentation
from pptx.oxml.ns import qn

MARK = "Source detail (S2R® manuscript):"
NEXT = "Professional insights:"


def shape_text(sh):
    tb = sh._element.find(qn("p:txBody"))
    if tb is None:
        return ""
    return "\n".join("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in tb.iter(qn("a:p")))


def main(deck, spec):
    d = json.load(open(spec, encoding="utf-8"))
    chk = d.pop("_check", {})
    p = Presentation(deck)
    slides = list(p.slides)
    for key, lines in d.items():
        s = slides[int(key) - 1]
        text = " | ".join(t for t in (shape_text(sh) for sh in s.shapes) if t.strip())
        want = chk.get(key, "")
        if want and not text.replace("\n", " / ").startswith(want):
            sys.exit(f"Slide {key} does not match: expected '{want}', found '{text[:40]}'")
        tf = s.notes_slide.notes_text_frame
        paras = list(tf.paragraphs)
        # remove an earlier block: from the blank line before MARK up to the blank line before NEXT (or the end)
        start = next((i for i, q in enumerate(paras) if q.text.strip() == MARK), None)
        if start is not None:
            a = start - 1 if start > 0 and not paras[start - 1].text.strip() else start
            nxt = next((i for i, q in enumerate(paras) if i > start and q.text.strip() == NEXT), None)
            b = len(paras) if nxt is None else (nxt - 1 if not paras[nxt - 1].text.strip() else nxt)
            for q in paras[a:b]:
                q._p.getparent().remove(q._p)
            paras = list(tf.paragraphs)
        # anchor: the blank line before "Professional insights:", when that block exists
        nxt = next((i for i, q in enumerate(paras) if q.text.strip() == NEXT), None)
        anchor = None
        if nxt is not None:
            anchor = paras[nxt - 1]._p if nxt > 0 and not paras[nxt - 1].text.strip() else paras[nxt]._p
        for line in ["", MARK, ""] + list(lines):
            para = tf.add_paragraph()
            para.text = line
            if anchor is not None:
                anchor.addprevious(para._p)
    p.save(deck)
    print(f"{deck}: source detail added to {len(d)} slides")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
