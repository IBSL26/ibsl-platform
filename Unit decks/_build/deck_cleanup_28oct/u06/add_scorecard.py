import sys,copy,re
from pptx import Presentation
from pptx.oxml.ns import qn
src,out=sys.argv[1],sys.argv[2]
p=Presentation(src)
S=list(p.slides)
assert S[19].shapes[0].text_frame.text.startswith('2.2 · Accountability') and S[32].shapes[0].text_frame.text.startswith('5.4 · Activity 4')
def dup(src_slide, pos):
    new=p.slides.add_slide(src_slide.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    rmap={}
    for rId,rel in src_slide.part.rels.items():
        if rel.reltype.endswith('/image'):
            rmap[rId]=new.part.relate_to(rel.target_part,rel.reltype)
    for el in src_slide.shapes._spTree.iterchildren():
        if el.tag in (qn('p:nvGrpSpPr'),qn('p:grpSpPr')): continue
        e=copy.deepcopy(el)
        for node in e.iter():
            for a in (qn('r:embed'),qn('r:link'),qn('r:id')):
                v=node.get(a)
                if v in rmap: node.set(a,rmap[v])
        new.shapes._spTree.append(e)
    lst=p.slides._sldIdLst; ids=list(lst); el=ids[-1]; lst.remove(el); lst.insert(pos,el)
    return new
def settext(sh,*parts):
    paras=sh.text_frame.paragraphs; runs=paras[0].runs
    for i,r in enumerate(runs): r.text=parts[i] if i<len(parts) else ''
    for extra in paras[1:]: extra._p.getparent().remove(extra._p)
def notes(slide,text,src=None):
    ns=slide.notes_slide
    if ns.notes_text_frame is None:
        st=ns.shapes._spTree
        for el in list(st):
            if el.tag not in (qn('p:nvGrpSpPr'),qn('p:grpSpPr')): st.remove(el)
        for el in src.notes_slide.shapes._spTree.iterchildren():
            if el.tag not in (qn('p:nvGrpSpPr'),qn('p:grpSpPr')): st.append(copy.deepcopy(el))
    tf=ns.notes_text_frame; tf.clear(); first=True
    HEAD=re.compile(r"^[A-Z0-9][A-Z0-9 ·,&—’'():®./\-–?]+:?$"); STEP=re.compile(r'^\d+\.\s[A-Z][^"]*\.$')
    for l in text.strip('\n').split('\n'):
        para=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        r=para.add_run(); r.text=l
        if l and (HEAD.match(l) or STEP.match(l)): r.font.bold=True
N=open('notes.txt',encoding='utf-8').read().split('@@@')
N={k.strip().split('\n',1)[0]:k.strip().split('\n',1)[1] for k in N if k.strip()}
# A: 2.3a from 5.1 layout (slide 31), inserted after slide 20 (index 20)
a=dup(S[30],20); sh=list(a.shapes)
settext(sh[0],'2.3a · Performance scorecard: three categories')
settext(sh[4],'Strategy'); settext(sh[5],'OKRs: progress on the multi-year strategy. 30–50% at senior levels.')
settext(sh[9],'Operational'); settext(sh[10],'KPIs: the operational engine at standard. 30–50% at most levels.')
settext(sh[14],'Behavioural'); settext(sh[15],'Values and behaviours: how results are achieved. 20–40%.')
settext(sh[17],'How results are achieved matters as much as what results are achieved.')
notes(a,N['A'],S[30])
# B: 2.3b from 5.2 layout (slide 32), after A (index 21)
b=dup(S[31],21); sh=list(b.shapes)
tb=[x for x in sh if x.has_text_frame and x.text_frame.text.strip()]
settext(tb[0],'2.3b · Developing a performance scorecard')
rows=[('1','Start from the enterprise Key Results'),('2','Add the Operational KPIs'),('3','Add the Behavioural measures'),('4','Weight the three categories to 100%'),('5','Make it transparent to the role holder')]
for x,(l,t) in zip(tb[1:],rows): settext(x,l+'  ',t)
notes(b,N['B'],S[31])
# C: 5.3 from WARD slide (slide 33), inserted before it: WARD now index 34 (0-based) after two inserts
c=dup(S[32],34); sh=list(c.shapes)
tb=[x for x in sh if x.has_text_frame and x.text_frame.text.strip()]
settext(tb[0],'5.3 · Activity 3: CEO Performance Scorecard')
rows=[('S','Strategy: Key Result measures and their weight'),('O','Operational: KPIs with targets and their weight'),('B','Behavioural: values and behaviours, how assessed'),('=','The three weights add up to 100%')]
for x,(l,t) in zip(tb[1:],rows): settext(x,l+'  ',t)
notes(c,N['C'],S[32])
# notes edits on existing slides
def edit(slide,old,new):
    tf=slide.notes_slide.notes_text_frame; done=False
    for para in tf.paragraphs:
        t=''.join(r.text for r in para.runs)
        if old in t:
            para.runs[0].text=t.replace(old,new)
            for r in para.runs[1:]: r.text=''
            done=True
    assert done,old
s16=S[15]
edit(s16,'2. 2.2: compliance and commitment architecture. The four dimensions, then accountability for people performance and organisational performance (two slides).','2. 2.2: compliance and commitment architecture. The four dimensions, then accountability for people performance and organisational performance (two slides).\n3. 2.3: performance scorecard development. The three categories, then how a scorecard is developed (two slides).')
edit(s16,'The section closes on the 2.2 slide','The section closes on the 2.3b slide')
s20=S[19]
edit(s20,'4. Close Section 2 and bridge to Section 3.','4. Bridge to 2.3.')
edit(s20,'Say: "Performance management bridges strategy and execution when it builds commitment to results. Section 3 examines where performance management becomes vulnerable across executive perspectives, and how its practice is evolving."','Say: "Commitment needs measures each person can see. The performance scorecard is where the connection between enterprise results and the people who deliver them reaches the individual."')
p.save(out); print('saved',len(p.slides),'slides')
