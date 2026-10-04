"""Add journal insights to the presenter notes of Carol's own Unit 1 deck.

Usage: python3 insights_u01.py <deck.pptx> <u01_insights.json>

Unlike enrich.py, this never deletes anything. For each slide it inserts the entries
straight after the last "Professional insights" heading already in the notes (where: placeholder),
or appends them at the end of the notes, adding the heading when the notes have none (where: end).
Safe to re-run: an entry whose label and source line are already in the notes is skipped.
Slide text is not touched.
"""
import json, re, sys, copy
from pptx import Presentation
from pptx.oxml.ns import qn

MARK = "Professional insights:"
HEAD = re.compile(r"^professional insights?:?$", re.I)

def shape_text(sh):
    tb = sh._element.find(qn("p:txBody"))
    if tb is None:
        return ""
    return "\n".join("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in tb.iter(qn("a:p")))

def ptext(para):
    return "".join(r.text for r in para.runs)

def main(deck, spec):
    d = json.load(open(spec, encoding="utf-8"))
    chk = d.pop("_check", {}); d.pop("_about", None)
    prs = Presentation(deck)
    slides = list(prs.slides)
    added = skipped = 0
    for key, item in d.items():
        s = slides[int(key) - 1]
        text = " | ".join(t for t in (shape_text(sh) for sh in s.shapes) if t.strip())
        want = chk.get(key, "")
        if want and want not in text[:160]:
            sys.exit(f"Slide {key} does not match: expected '{want}', found '{text[:60]}'")
        tf = s.notes_slide.notes_text_frame
        paras = list(tf.paragraphs)
        whole = "\n".join(ptext(p) for p in paras)
        todo = [e for e in item["entries"] if not (f"{e['label']}:" in whole and f"Source: {e['source']}" in whole)]
        skipped += len(item["entries"]) - len(todo)
        if not todo:
            continue
        # heading paragraphs: last non-empty line of the paragraph is the heading
        heads = [p for p in paras if ptext(p).strip() and HEAD.match(ptext(p).strip().split("\n")[-1].strip())]
        lines = []
        if item["where"] == "placeholder":
            if not heads:
                sys.exit(f"Slide {key}: no 'Professional insights' heading found")
            anchor = heads[-1]._p
        else:
            anchor = paras[-1]._p
            if not heads:
                lines += ["", MARK]
        for e in todo:
            lines += ["", f"{e['label']}: {e['text']}", f"Source: {e['source']}"]
        for line in lines:
            np_ = tf.add_paragraph()          # appended at the end …
            np_.text = line
            el = np_._p
            if el.getprevious() is not anchor:  # … then moved to sit after the anchor
                anchor.addnext(el)
            anchor = el
        added += len(todo)
    prs.save(deck)
    print(f"{deck}: {added} insights added, {skipped} already present")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
