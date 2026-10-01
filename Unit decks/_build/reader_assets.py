"""Build the deck-reader files for one or more unit decks.

Usage: python3 reader_assets.py <out_dir> <deck.pptx> [<deck.pptx> ...]
For each deck writes <out_dir>/Module-<m>/unit-<NN>/ with s01.jpg … and manifest.json
(slide titles and presenter notes). Upload each unit folder into the matching
Module folder of the facilitator_decks bucket. Run after enrich.py.
"""
import json, os, re, shutil, subprocess, sys, tempfile, glob
from pptx import Presentation
from pptx.oxml.ns import qn

def module_of(unit):
    return 1 if unit == 1 else 2 if unit <= 4 else 3 if unit <= 9 else 4

def shape_text(sh):
    tb = sh._element.find(qn("p:txBody"))
    if tb is None:
        return ""
    return "\n".join("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in tb.iter(qn("a:p")))

def build(out_dir, deck):
    base = os.path.basename(deck)
    m = re.match(r"Unit (\d\d) - (.+)\.pptx$", base)
    unit, title = int(m.group(1)), m.group(2)
    dest = os.path.join(out_dir, f"Module-{module_of(unit)}", f"unit-{unit:02d}")
    os.makedirs(dest, exist_ok=True)
    for old in glob.glob(os.path.join(dest, "s*.jpg")):
        os.remove(old)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, deck],
                       check=True, capture_output=True)
        pdf = os.path.join(tmp, base[:-5] + ".pdf")
        subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=82", "-scale-to-x", "1600",
                        "-scale-to-y", "-1", pdf, os.path.join(tmp, "s")], check=True)
        imgs = sorted(glob.glob(os.path.join(tmp, "s-*.jpg")), key=lambda p: int(re.findall(r"(\d+)\.jpg$", p)[0]))
        slides = []
        prs = Presentation(deck)
        if len(imgs) != len(prs.slides):
            sys.exit(f"{base}: {len(imgs)} images for {len(prs.slides)} slides")
        for i, (img, s) in enumerate(zip(imgs, prs.slides), 1):
            name = f"s{i:02d}.jpg"
            shutil.move(img, os.path.join(dest, name))
            texts = [t.strip() for t in (shape_text(sh) for sh in s.shapes) if t.strip()]
            first = texts[0].split("\n")[0] if texts else f"Slide {i}"
            if re.fullmatch(r"\d+", first) and len(texts) > 1:   # section dividers open with the number
                first = " · ".join(x.split("\n")[0] for x in texts[1:3])
            notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
            slides.append({"n": i, "img": name, "title": first[:90], "notes": notes})
    manifest = {"unit": unit, "title": title, "module": module_of(unit), "slides": slides}
    with open(os.path.join(dest, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
    print(f"{base}: {len(slides)} slides -> {dest}")

if __name__ == "__main__":
    for d in sys.argv[2:]:
        build(sys.argv[1], d)
