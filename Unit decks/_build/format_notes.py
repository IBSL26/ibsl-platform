"""Put the generated presenter notes into Carol's own copy of the deck.

Usage: python3 format_notes.py <carol_deck.pptx> <generated_deck.pptx> <deck_out.pptx> [first_slide_to_replace]

Carol's deck is the master: its slides are not touched, and the notes of the slides before
<first_slide_to_replace> (default 5) stay exactly as she wrote them. From that slide on, the notes are replaced
with the notes of the generated deck:
  - one paragraph for each line;
  - a line that starts with § is a heading: the mark is removed and the line is set in bold.
The two decks must have the same slides in the same order (the script stops if the slide text differs).
One correction is made in Carol's own notes: on the Facilitator Guide slide "using the list above" becomes
"using the list below", because the list of the five sections now sits under the steps.
Step 4 of that slide's notes ("Keep the four Key Facilitation Questions in front of you…") is removed, on her instruction.
"""
import copy, sys
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

def slide_text(slide):
    out = []
    for sh in slide.shapes:
        tb = sh._element.find(qn("p:txBody"))
        if tb is None:
            continue
        for p in tb.iter(qn("a:p")):
            t = "".join(x.text or "" for x in p.iter(qn("a:t"))).strip()
            if t:
                out.append(t)
    return out

def notes_body(slide):
    """The txBody of the notes placeholder."""
    return slide.notes_slide.notes_text_frame._txBody

def notes_text(slide):
    return slide.notes_slide.notes_text_frame.text.replace("\x0b", "\n")

def paragraph(line):
    p = etree.Element(f"{{{A}}}p")
    bold = line.startswith("§")
    text = line[1:] if bold else line
    if not text.strip():
        e = etree.SubElement(p, f"{{{A}}}endParaRPr"); e.set("lang", "en-US"); e.set("dirty", "0")
        return p
    r = etree.SubElement(p, f"{{{A}}}r")
    rpr = etree.SubElement(r, f"{{{A}}}rPr"); rpr.set("lang", "en-US")
    if bold:
        rpr.set("b", "1")
    rpr.set("dirty", "0")
    t = etree.SubElement(r, f"{{{A}}}t"); t.text = text
    return p

def main(base, gen, out, first=5):
    b, g = Presentation(base), Presentation(gen)
    bs, gs = list(b.slides), list(g.slides)
    if len(bs) != len(gs):
        sys.exit(f"slide count differs: {len(bs)} in Carol's deck, {len(gs)} in the generated deck")
    for i, (x, y) in enumerate(zip(bs, gs), 1):
        if slide_text(x) != slide_text(y):
            sys.exit(f"slide {i}: the slide text differs between the two decks; stop and look before replacing notes")
    replaced = 0
    for i, (x, y) in enumerate(zip(bs, gs), 1):
        if i < first:
            continue
        text = notes_text(y)
        if "§" not in text.split("\n")[0]:
            sys.exit(f"slide {i}: generated notes do not start with a heading")
        body = notes_body(x)
        for p in body.findall(qn("a:p")):
            body.remove(p)
        for line in text.split("\n"):
            body.append(paragraph(line.rstrip()))
        replaced += 1
    # the one correction inside Carol's own notes
    fixed = 0
    for i, x in enumerate(bs, 1):
        if i >= first:
            break
        for t in notes_body(x).iter(qn("a:t")):
            if t.text and "using the list above" in t.text:
                t.text = t.text.replace("using the list above", "using the list below"); fixed += 1
    # Carol, 5 Oct: "Yes remove that step" — step 4 of her slide 3 notes named the Key Facilitation Questions she had deleted.
    cut = 0
    for i, x in enumerate(bs, 1):
        if i >= first:
            break
        for para in notes_body(x).findall(qn("a:p")):
            runs = para.findall(qn("a:r"))
            for k, r in enumerate(runs):
                t = r.find(qn("a:t"))
                if t is not None and t.text and "\n4. Keep the four Key Facilitation Questions in front of you." in t.text:
                    t.text = t.text.split("\n4. Keep the four Key Facilitation Questions in front of you.")[0]
                    tail = "".join((y.find(qn("a:t")).text or "") for y in runs[k + 1:])
                    if not tail.strip().endswith("role contribution (Section 5)."):
                        sys.exit("slide %d: the end of step 4 is not where it was; look before cutting" % i)
                    for y in runs[k + 1:]:
                        para.remove(y)
                    cut += 1
    b.save(out)
    print(f"notes replaced on slides {first} to {len(bs)} ({replaced} slides); Carol's notes kept on slides 1 to {first - 1}; 'above' -> 'below' corrections: {fixed}; step 4 removed on the Facilitator Guide slide: {cut}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4]) if len(sys.argv) > 4 else 5)
