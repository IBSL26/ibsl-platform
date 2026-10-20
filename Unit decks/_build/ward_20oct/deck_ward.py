# -*- coding: utf-8 -*-
"""WARD in the Unit 6 and Unit 8 decks (20 Oct 2026). python3 deck_ward.py <u6 in> <u8 in> <u6 out> <u8 out>"""
import sys
sys.path.insert(0, '.')
from decklib import *
from ward_text import *
from pptx import Presentation
s6, s8, o6, o8 = sys.argv[1:5]
def drop(slide, names):
    for n in names: e = shp(slide, n)._element; e.getparent().remove(e)
def move_last_to(prs, pos):
    lst = prs.slides._sldIdLst; new = list(lst)[-1]; lst.remove(new); lst.insert(pos, new)
# ── Unit 6 ──
p = Presentation(s6); S = list(p.slides); assert len(S) == 33 and S[31].shapes[0].text_frame.text.startswith('5.2 · Activity 2')
sub(S[29], [("Application · Each group builds two connected components of the Enterprise Performance Management System for the strategy in its Capstone Blueprint: the S2R® Contribution Map™ and the FACES application.",
             "Application · Each group builds three connected components of the Enterprise Performance Management System for the strategy in its Capstone Blueprint: the S2R® Contribution Map™, the FACES application and the WARD check-up for its enterprise Key Result."),
            ("your group uses that work to develop two connected components of your Enterprise Performance Management System: Activity 1, the S2R® Contribution Map™, and Activity 2, the FACES application.\"",
             "your group uses that work to develop three connected components of your Enterprise Performance Management System: Activity 1, the S2R® Contribution Map™, Activity 2, the FACES application, and Activity 3, the WARD check-up for your enterprise Key Result.\""),
            ("Order: both activities are group work and Capstone work. Activity 1 needs the group’s confirmed Enterprise OKRs from Unit 3, part 5.2. Activity 2 uses the business functions the group names in Activity 1.",
             "Order: all three activities are group work and Capstone work. Activity 1 needs the group’s confirmed Enterprise OKRs from Unit 3, part 5.2. Activity 2 uses the business functions the group names in Activity 1. Activity 3 sets the WARD check-up for the Key Result chosen in Activity 1, using the conditions named in the Contribution Map.")])
sub(S[31], [("2. The group reads the completion check.\n3. Each participant reads the record and selects Agree. Print gives a copy.\nThe confirmed record feeds the team’s Capstone Blueprint.\n\nAFTER THE SESSION:\nThe record reaches you with each participant’s submission, under the heading of Section 5. The confirmed record feeds the Unit 6 boxes of the team’s Capstone Blueprint: the Contribution Map and the FACES application.",
             "2. The group reads the completion check, then goes to Activity 3.")])
n = duplicate(p, S[31])
put(n, 'Text 0', '5.3 · Activity 3: WARD check-up')
put(n, 'Text 2', 'W  ', 'Waypoints: when the Key Result is checked')
put(n, 'Text 4', 'A  ', 'Alert: how far off raises a red flag')
put(n, 'Text 6', 'R  ', 'Recovery: for individuals and conditions')
put(n, 'Text 8', 'D  ', 'Decider: who decides and acts')
drop(n, ['Shape 9', 'Image 4', 'Text 10'])
exn = '\n'.join('%s: %s' % (a, b) for a, b in EX)
set_notes(n, '''5.3 — ACTIVITY 3: WARD · YOUR KEY RESULT’S VITAL CHECK-UP (CAPSTONE WORK)

HOW TO TEACH IT:

1. Open with the check-up.
Say: "''' + INTRO.replace('your group', 'your group') + '''"

2. Teach the four letters.
W · Waypoints: ''' + WARD[0][2] + '''
A · Alert: ''' + WARD[1][2] + '''
R · Recovery: ''' + WARD[2][2] + '''
D · Decider: ''' + WARD[3][2] + '''

3. Take the worked example through WARD.
Enterprise Key Result from part 4.1: ''' + EX_KR + '''
''' + exn + '''
Ask: "At the end of June the result is 74%. What happens on Monday morning, and who decides?"
Guide the groups to see that the waypoint, the alert and the recovery answer the question before the review takes place.

4. Read the completion check.
1. Each waypoint has a date and a level, and the waypoints lead to the target by the deadline.
2. The alert is specific enough that every member reads the same result the same way.
3. The recovery covers the individuals and the organisational conditions named in the Contribution Map.
4. One executive is named as the decider.
Say: "In Unit 8 your group tests its WARD: the alert fires, and you run the recovery."

ON THE PORTAL, AFTER THE TEACHING:
Participants work as a group in Activity 3 of Section 5 of the participant file. This is Capstone work. The group chooses its scribe for this work. The scribe types what the group agrees, and the entries appear on every member’s page.
1. The group sets two or three waypoints for its enterprise Key Result: a date and a level for each.
2. The group sets the alert, the recovery for the individuals and for the organisational conditions, and the decider.
3. The group reads the completion check.
4. Each participant reads the record and selects Agree. Print gives a copy.
The agreed record feeds the team’s Capstone Blueprint.

AFTER THE SESSION:
The record reaches you with each participant’s submission, under the heading of Section 5. The agreed record feeds the Unit 6 boxes of the team’s Capstone Blueprint: the Contribution Map in 6A, the FACES application in 6B and the WARD check-up in 6C.''')
move_last_to(p, 32)
sub(list(p.slides)[33], [("- 5 · Application: Each group built an S2R® Contribution Map™ for one of its own enterprise Key Results and applied FACES to manage each function’s contribution.",
                           "- 5 · Application: Each group built an S2R® Contribution Map™ for one of its own enterprise Key Results, applied FACES to manage each function’s contribution and set the WARD check-up for the Key Result.")])
p.save(o6)
# ── Unit 8 ──
q = Presentation(s8); T = list(q.slides); assert len(T) == 24 and T[23].shapes[0].text_frame.text.startswith('5.1 ·')
sub(T[22], [("2. Writes a decision point for recovery for each gap: the decision the executive team takes when the gap appears, and who takes it.\nThe group chooses",
             "2. Writes a decision point for recovery for each gap: the decision the executive team takes when the gap appears, and who takes it.\nIn part 5.2 each group tests the WARD check-up it set for its enterprise Key Result in Unit 6: the alert fires, and the group runs the recovery.\nThe group chooses"),
            ("and sets a decision point for recovery for each.\"", "and sets a decision point for recovery for each. Then your group tests the WARD check-up it set in Unit 6: the alert fires, and you run the recovery.\"")])
m = duplicate(q, T[23])
put(m, 'Text 0', '5.2 · Testing your WARD')
put(m, 'Text 2', '01  ', 'The red flag  ·  one waypoint, beyond the alert')
put(m, 'Text 4', '02  ', 'What explains the gap  ·  individuals · conditions')
put(m, 'Text 6', '03  ', 'The recovery  ·  what, who decides, when checked again')
drop(m, ['Shape 7', 'Image 3', 'Text 8'])
ex8 = '\n'.join('%s: %s' % (a, b) for a, b in EX8)
set_notes(m, '''5.2 — TESTING YOUR WARD (CAPSTONE WORK)

HOW TO TEACH IT:

1. Recall the check-up.
Say: "In Unit 6 your group set the WARD check-up for its enterprise Key Result: the waypoints, the alert, the recovery and the decider. It was set while the Key Result was on track. Today the alert fires."

2. Fire the alert.
Say: "Choose one of your waypoints and write the result your executive team reads at that review, beyond the alert."

3. Run the recovery.
Ask: "What explains the gap in the individuals and in the organisational conditions? Which recovery do you run, who decides, and when do you check the result again?"
Guide the groups to use the evidence in their Contribution Map and the recovery they set in advance.

4. Give the example from the Unit 6 worked example.
Enterprise Key Result: ''' + EX_KR + ''' Waypoint: 80% by the end of June. Alert: more than three percentage points below a waypoint.
''' + ex8 + '''

ON THE PORTAL, AFTER THE TEACHING:
Participants work as a group in part 5.2 of their own file. This is Capstone work. The group’s WARD from Unit 6 is shown on each page; if it is missing, the group selects Bring in from my Unit 6 page.
1. The group writes the red flag: the waypoint and the result that raises the alert.
2. The group states what explains the gap in the performance of individuals and in the organisational conditions.
3. The group states the recovery: what is done, who decides and when the result is checked again.
4. Each participant reads the record and selects Agree.

AFTER THE SESSION:
The record reaches you with each participant’s submission, under the heading of Section 5. The agreed record feeds the Unit 8 boxes of the team’s Capstone Blueprint: the four consequences in 8A to 8D and the WARD test in 8E.''')
q.save(o8)
print('saved', len(p.slides), len(q.slides))
