# Rewrites the notes of every slide as one paragraph per line and sets bold:
# § lines (headings), numbered step titles inside HOW TO TEACH IT / HOW TO USE THIS SLIDE / HOW TO OPEN THE SECTION,
# and the Say: / Ask: / Listen for: / Deepen, say: labels.
# Usage: python3 bold_notes.py in.pptx [out.pptx]
import sys,re
from pptx import Presentation
LAB=re.compile(r'^(Say:|Ask:|Listen for, then confirm:|Listen for:|Deepen, say:|Add:)')
TEACH=('HOW TO TEACH IT','HOW TO USE THIS SLIDE','HOW TO OPEN THE SECTION')
p=Presentation(sys.argv[1])
for s in p.slides:
    tf=s.notes_slide.notes_text_frame
    lines=tf.text.replace('\r','').split('\n')
    for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
    first=tf.paragraphs[0]
    for r in list(first.runs): r._r.getparent().remove(r._r)
    teach=False; listmode=False
    for i,l in enumerate(lines):
        para=first if i==0 else tf.add_paragraph()
        head=l.startswith('§'); l=l.lstrip('§')
        if head:
            if l.startswith(TEACH): teach=True; listmode=False
            elif 'LEARNING OUTCOMES' in l: listmode=True
            else: teach=False
        elif not l.strip(): listmode=False
        step = teach and not listmode and not head and re.match(r'^\d+\. ',l) and len(l)<110 and '"' not in l
        m=LAB.match(l)
        if head or step:
            r=para.add_run(); r.text=l; r.font.bold=True
        elif m:
            r=para.add_run(); r.text=m.group(1); r.font.bold=True
            r2=para.add_run(); r2.text=l[len(m.group(1)):]
        else:
            r=para.add_run(); r.text=l
p.save(sys.argv[2] if len(sys.argv)>2 else sys.argv[1])
