# -*- coding: utf-8 -*-
"""Unit 6 deck: Carol's own deck is the master. This script changes the Section 5 slides only, to match the live pages of 9 October 2026.
Usage: python3 carol_deck_s5.py <Carol's deck> <participant page> <facilitator page> <out.pptx>
Slides 1 to 22 (her Sections 1 to 4) are not touched: no shape, no note.
Changed: the Section 5 opener (title and notes); "The six design tasks" becomes Step 1 · Scoring Logic; a copy of her 1.2 FACES slide becomes Step 2 · FACES;
a copy of her 4.1 EXECUTION slide becomes Step 3 · EXECUTION; the case slide becomes Step 4 · PM Scorecards; the Task 3 slide leaves;
the closing slide asks the closing question of the facilitator page (the 30-day commitment has left the unit).
Notes of the six changed slides follow the form of the Unit 2 to 5 decks: heading, HOW TO TEACH IT, numbered steps with Say / Ask, ON THE PORTAL, AFTER THE TEACHING."""
import sys, re, copy, html
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree
import u6_content as C
deck, P, F, out = sys.argv[1:5]
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
def plain(path):
    t = open(path, encoding='utf-8').read()
    t = re.sub(r'<style.*?</style>|<script.*?</script>|<!--.*?-->', ' ', t, flags=re.S)
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', t)))
vp, vf = plain(P), plain(F)
def on_pages(s, where='pf'):
    for w in where: assert s in (vp if w == 'p' else vf), ('not on the %s page: ' % w) + s
    return s

prs = Presentation(deck)
S = list(prs.slides)
assert len(S) == 27, len(S)
def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def first(s): return texts(s)[0].text_frame.text
assert first(S[24]) == 'The six design tasks' and first(S[23]).startswith('5.1 · Case') and first(S[25]).startswith('Task 3') and first(S[6]).startswith('1.2 · The FACES') and first(S[21]).startswith('4.1 · The EXECUTION')
assert [sh.text_frame.text for sh in texts(S[22])][3] == 'MyHealth Live Portal — PM Design Simulation' and 'Unit Summary' in texts(S[26])[0].text_frame.text

def set_text(shape, *run_texts):
    """New text in a text box, keeping the look of its runs. One string for each run of the first paragraph."""
    ps = shape.text_frame.paragraphs; runs = ps[0].runs
    assert len(ps) == 1 or all(not p.runs for p in ps[1:]), shape.name
    assert len(runs) == len(run_texts), (shape.name, [r.text for r in runs], run_texts)
    for r, t in zip(runs, run_texts): r.text = t
def set_notes(slide, lines):
    """Notes as one paragraph for each line. A line that starts with § is a heading or a step title, set in bold."""
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in lines:
        p = etree.SubElement(body, '{%s}p' % A)
        bold = line.startswith('§'); text = line[1:] if bold else line
        if not text.strip():
            e = etree.SubElement(p, '{%s}endParaRPr' % A); e.set('lang', 'en-GB'); e.set('dirty', '0'); continue
        r = etree.SubElement(p, '{%s}r' % A); rp = etree.SubElement(r, '{%s}rPr' % A); rp.set('lang', 'en-GB')
        if bold: rp.set('b', '1')
        rp.set('dirty', '0'); etree.SubElement(r, '{%s}t' % A).text = text
def duplicate(src):
    """A copy of a slide, with its background and pictures, added at the end of the deck."""
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId] = dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID + 'embed', RID + 'link', RID + 'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    left = [el.get(RID + 'embed') for el in dst._element.iter() if el.get(RID + 'embed')]
    assert all(x in dst.part.rels for x in left), left
    # the notes page: the same placeholders as the notes page of the slide it was copied from
    nt, st = dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'), qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag == qn('p:sp')]: nt.append(copy.deepcopy(el))
    assert dst.notes_slide.notes_text_frame is not None
    return dst

GROUP = 'Participants work as a group in Step %d of Section 5 of the participant file. This is Capstone work.'
# ── Section 5 opener ──
s23 = S[22]; set_text(texts(s23)[3], on_pages(C.S5_TITLE))
slo = [texts(s23)[5].text_frame.text, texts(s23)[6].text_frame.text]; [on_pages(x) for x in slo]
set_notes(s23, ['§SECTION 5 · APPLICATION — IN PRACTICE', '', '§HOW THIS SECTION RUNS:',
  on_pages('This is group work. Your group designs the PM architecture for the strategy in your Capstone Blueprint, in four steps:', 'p'),
  on_pages('Step 1 · Scoring Logic. Your group sets what each rating means, the evidence that confirms it and the leadership response it triggers in the monthly review. It then states the factors that may legitimately affect a rating, the lines of sight of progress it uses and its review cadence. The lessons from parts 1.3 and 1.4 sit above the boxes.', 'p'),
  on_pages('Step 2 · FACES. Your group describes how its PM architecture will serve each of the five FACES functions. The lesson from part 1.2 sits above each box.', 'p'),
  on_pages('Step 3 · EXECUTION. Your group describes how each of the nine EXECUTION principles is designed into its PM architecture. The lesson from part 4.1 sits above each box.', 'p'),
  on_pages('Step 4 · PM Scorecards. Your group designs a PM scorecard for the CEO and one for the CFO. It uses two of the Key Results it confirmed in Unit 3 as the basis for measurement.', 'p'),
  on_pages('You then read your group’s record at the foot of this section and select Confirm. This is Capstone work, and your confirmed record feeds your team’s Capstone Blueprint.', 'p'),
  '', '§SECTION LEARNING OUTCOMES:', '1. ' + slo[0], '2. ' + slo[1], '', '§HOW TO OPEN THE SECTION:', '',
  '§1. Say what the section does.', 'Say: "In this section your group designs the PM architecture for the strategy in your Capstone Blueprint.',
  'You set the scoring logic. You describe the FACES and the EXECUTION of your architecture.', 'You design a PM scorecard for the CEO and one for the CFO."', '',
  '§2. Say how the group works.', 'Say: "All four steps are group work.', 'Your group agrees each entry.', 'One member acts as scribe.', 'Every member then types the agreed entries into their own page, in the session or after it."', '',
  '§3. Say where the work goes.', 'Say: "This is Capstone work. Your group’s confirmed record feeds your team’s Capstone Blueprint."',
  'Step 4 needs the group’s Enterprise OKRs. Each group confirms them in Unit 3, part 4.2, first.'])
# ── Step 1 · Scoring Logic (her "six design tasks" slide; the six points in the order she gave them) ──
s1 = S[24]; t = texts(s1); set_text(t[0], 'Step 1 · Scoring Logic')
six = [on_pages('Performance level that earns this rating', 'p'), on_pages('Evidence required to confirm it', 'p'), on_pages('Contextual factors that might legitimately affect a rating', 'p'),
       on_pages('Lines of sight of progress', 'p'), 'Review cadence and why', on_pages('Leadership response in the monthly review', 'p')]
for i, x in enumerate(six): set_text(t[i + 1], '%d  ' % (i + 1), x)
set_notes(s1, ['§SECTION 5 · STEP 1 — SCORING LOGIC (CAPSTONE WORK)', '', '§HOW TO TEACH IT:', '',
  '§1. Brief the step.', 'Say: "Your group designs the scoring logic of its PM architecture. There are six things to agree. They are on the slide."', '',
  '§2. Teach the rating scale.', 'Say: "For each rating from 1 to 5, your group agrees three things:', 'the performance level that earns the rating,', 'the evidence required to confirm it,', 'and the leadership response it should trigger in the monthly review."',
  'The lesson above the three boxes is the rating card of part 1.4.', 'Watch for: a performance level with no measure in it. Ask: "What would a manager see or count?"', 'Watch for: one leadership response given for two ratings. Ask: "What changes for the person at each rating?"', '',
  '§3. Teach the contextual factors.', 'Say: "Agree the contextual factors that might legitimately affect a rating."', 'Examples on the page: the size of the unit, the volume of work, a change outside the person’s control.', '',
  '§4. Teach the lines of sight of progress.', 'Say: "Describe how your PM architecture uses each of the Three Sights of Progress: Strategy, Operational, and Behavioural & Values Alignment."', 'The lesson above each box is part 1.3.', '',
  '§5. Teach the review cadence.', 'Say: "Describe your review cadence and why your group chose it."', 'Recall the Ongoing principle of part 4.1: ' + on_pages(C.EXEC_O_ADDED), '',
  '§ON THE PORTAL, AFTER THE TEACHING:', GROUP % 1,
  '1. ' + on_pages('For each rating from 1 to 5, the group agrees the performance level that earns it, the evidence required to confirm it and the leadership response it should trigger in the monthly review.', 'f'),
  '2. ' + on_pages('The group agrees the contextual factors that might legitimately affect a rating.', 'f'), '3. ' + on_pages('The group describes the lines of sight of progress its PM architecture uses.', 'f'), '4. ' + on_pages('The group describes its review cadence and why it chose it.', 'f')])
# ── Step 4 · PM Scorecards (her case slide) ──
s4 = S[23]; t = texts(s4); set_text(t[0], 'Step 4 · PM Scorecards')
set_text(t[1], 'Key Results'); set_text(t[2], on_pages('Choose two of your Key Results as the basis for measurement.', 'p'))
set_text(t[3], 'CEO'); set_text(t[4], on_pages('Design the PM scorecard for the CEO.', 'p'))
set_text(t[5], 'CFO'); set_text(t[6], on_pages('Design the PM scorecard for the CFO in the same way.', 'p'))
set_text(t[7], on_pages('The four weights add up to 100%.', 'p'))
set_notes(s4, ['§SECTION 5 · STEP 4 — PM SCORECARDS (CAPSTONE WORK)', '', '§HOW TO TEACH IT:', '',
  '§1. Open with the group’s Key Results.', 'Say: "Your page shows the Enterprise OKRs your group confirmed in Unit 3.', 'Choose two of your Key Results as the basis for measurement."', '',
  '§2. Teach the scorecard.', 'Say: "Each scorecard has four rows: Key Result 1, Key Result 2, the operational measures, and the behavioural and values measures.', 'Each row carries a measure and a weight.', 'The four weights add up to 100%."',
  'Recall part 2.3: ' + on_pages(C.LVT_ONE_LINE), 'Recall part 1.3 for the weight of each sight.', '',
  '§3. Design the scorecard for the CEO, then for the CFO.', 'Say: "For the CEO, state what the CEO is measured on for each Key Result. Then add the operational measures and the behavioural and values measures.', 'Design the PM scorecard for the CFO in the same way."',
  'Recall part 3.1: the hot zone of the CEO is Enterprise Visibility, and the hot zone of the CFO is Financial Signals.', 'Watch for: the same measure on both scorecards. Ask: "What is each role best placed to track?"', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', GROUP % 4,
  '1. ' + on_pages('The group reads its Enterprise OKRs from Unit 3, shown on each page.', 'f'), '2. ' + on_pages('The group chooses two of its Key Results as the basis for measurement.', 'f'),
  '3. ' + on_pages('For the CEO, the group states what the role is measured on for each Key Result, then the operational measures and the behavioural and values measures, each with a weight.', 'f'),
  '4. ' + on_pages('The group designs the PM scorecard for the CFO in the same way.', 'f'), '5. ' + on_pages('Each participant reads the record and selects Confirm. Print gives a copy.', 'f'), '',
  '§AFTER THE SESSION:', 'The record reaches you with each participant’s submission, under the heading Designing the PM Architecture.',
  'The confirmed record feeds boxes 6A to 6D of the team’s Capstone Blueprint: the scoring logic in 6A, FACES in 6B, EXECUTION in 6C and the two scorecards in 6D.'])
# ── Step 2 · FACES (a copy of her 1.2 slide) and Step 3 · EXECUTION (a copy of her 4.1 slide) ──
s2 = duplicate(S[6]); set_text(texts(s2)[0], 'Step 2 · FACES')
fb = re.findall(r'<strong>(.*?)</strong>', re.search(r'class="acc-t">1\.2 &mdash;.*?<div class="fac-note-full">', open(F, encoding='utf-8').read(), re.S).group(0)); assert len(fb) == 5, fb
set_notes(s2, ['§SECTION 5 · STEP 2 — FACES (CAPSTONE WORK)', '', '§HOW TO TEACH IT:', '',
  '§1. Brief the step.', 'Say: "Your group describes how its PM architecture will serve each of the five FACES functions."', '',
  '§2. Recall each function with its guide.', 'The guide above each box is the lesson of part 1.2.'] +
  ['%s · %s: %s' % (l, n, html.unescape(re.sub(r'<[^>]+>', '', b))) for (l, n), b in zip(C.FACES, fb)] + ['',
  '§3. Set the standard.', 'Say: "For each function, name what in your architecture makes it happen: a rule, a routine, a report or an owner."', 'Watch for: a function restated in other words. Ask: "What in the architecture makes it happen?"', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', GROUP % 2, '1. ' + on_pages('The group reads the guide in each card. It is the lesson from part 1.2.', 'f'), '2. ' + on_pages('For each of the five functions, the group describes how its PM architecture will serve it.', 'f')])
s3 = duplicate(S[21]); set_text(texts(s3)[0], 'Step 3 · EXECUTION')
set_notes(s3, ['§SECTION 5 · STEP 3 — EXECUTION (CAPSTONE WORK)', '', '§HOW TO TEACH IT:', '',
  '§1. Recall the distinction.', 'Say: "' + on_pages('FACES tells us what a PM system must deliver.') + '', on_pages('EXECUTION tells us what must be designed into the system for it to deliver.') + '"',
  'Say: "' + on_pages('Step 2 states what your PM architecture must deliver. Step 3 states what your group designs into it so that it delivers.', 'p') + '"', '',
  '§2. Recall each principle with its self-check question.', 'The guide above each box is the lesson of part 4.1.'] +
  ['%s · %s — %s: %s' % (l, w, n, on_pages(q)) for (l, w, n, _, q) in C.EXEC] + ['',
  '§3. Set the standard.', 'Say: "For each principle, describe how it is designed into your PM architecture."', 'Watch for: an entry that repeats the Step 2 answer. Ask: "How is it built into the system?"', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', GROUP % 3, '1. ' + on_pages('The group reads the guide in each card. It is the lesson from part 4.1.', 'f'), '2. ' + on_pages('For each of the nine principles, the group describes how it is designed into its PM architecture.', 'f')])
# ── Unit Summary: the closing question of the facilitator page ──
s27 = S[26]; t = texts(s27)
QUESTION = on_pages('Which part of your PM architecture was hardest for your group to agree, and what does that tell you about the PM system in your own organisation?', 'f')
set_text(t[1], QUESTION); set_text(t[2], on_pages('Specificity is the test of whether the unit has converted into intent.', 'f'))
blocks = re.findall(r'<div style="font-size:9px;font-weight:700;letter-spacing:2px;text-transform:uppercase;color:[^"]*;margin-bottom:8px;">(.*?)</div><p[^>]*>(.*?)</p>', open(F, encoding='utf-8').read(), re.S); assert len(blocks) == 5, len(blocks)
set_notes(s27, ['§UNIT SUMMARY (UNIT 6 SYNTHESIS)', '', '§HOW TO TEACH IT:', '', '§1. Recap the unit, section by section.'] +
  ['- %d · %s: %s' % (i + 1, html.unescape(a).split(' — ')[0], html.unescape(re.sub(r'<[^>]+>', '', b))) for i, (a, b) in enumerate(blocks)] + ['',
  '§2. Ask the closing question on the slide.', 'Ask: "' + QUESTION + '"', 'Take an answer from each group.', '',
  '§3. Close the unit.', 'Say: "A performance management system that only measures has not yet become a management system.', 'The moment measurement becomes the basis for a decision, a conversation, and an accountable intervention — that is the moment execution begins."', '',
  '§4. Say what happens on the portal.', 'Say: "You now complete the unit on the portal: your reflections and your group’s record in Section 5.', 'When your page is complete, you submit the unit."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', 'Participants complete every part of the participant file and submit the unit to the facilitator from the end of Section 5.', '',
  '§AFTER THE SESSION:', 'Each submission reaches you in the facilitator dashboard.', 'The group’s confirmed record feeds boxes 6A to 6D of the team’s Capstone Blueprint.'])
# ── order: 1 to 23, Step 1, Step 2, Step 3, Step 4, Unit Summary; the Task 3 slide leaves ──
lst = prs.slides._sldIdLst; ids = list(lst)
assert len(ids) == 29
old = {i + 1: ids[i] for i in range(27)}; new2, new3 = ids[27], ids[28]
order = [old[i] for i in range(1, 24)] + [old[25], new2, new3, old[24], old[27]]
drop = old[26]
for el in ids: lst.remove(el)
for el in order: lst.append(el)
prs.part.drop_rel(drop.get(RID + 'id'))
prs.save(out)
# ── check the result ──
chk = Presentation(out); assert len(chk.slides) == 28
for i, s in enumerate(chk.slides, 1):
    if i < 23: continue
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                assert r.font.size is not None and r.font.size >= Pt(24), (i, r.text, r.font.size)
    n = s.notes_slide.notes_text_frame.paragraphs[0]; assert n.runs and n.runs[0].font.bold, i
print('written', out, '· 28 slides ·', [(' | '.join(sh.text_frame.text for sh in chk.slides[i].shapes if sh.has_text_frame and sh.text_frame.text.strip())[:40]) for i in range(22, 28)])
