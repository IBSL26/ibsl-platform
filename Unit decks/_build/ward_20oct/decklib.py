# -*- coding: utf-8 -*-
"""Unit 4 (new 5.3 slide, notes) and Unit 5 (Step 2 highest-risk dimensions) on Carol's decks, 20 Oct 2026.
python3 deck_u4u5.py <unit 4 deck> <unit 5 deck> <out 4> <out 5>"""
import sys, copy
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
A='http://schemas.openxmlformats.org/drawingml/2006/main'
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
def ntext(s): return s.notes_slide.notes_text_frame.text.replace('\x0b','\n')
def set_notes(slide, text):
    body=slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in text.split('\n'):
        p=etree.SubElement(body,'{%s}p'%A)
        if not line.strip():
            e=etree.SubElement(p,'{%s}endParaRPr'%A); e.set('lang','en-GB'); continue
        r=etree.SubElement(p,'{%s}r'%A); rp=etree.SubElement(r,'{%s}rPr'%A); rp.set('lang','en-GB')
        etree.SubElement(r,'{%s}t'%A).text=line
def sub(slide, pairs):
    t=ntext(slide)
    for a,b in pairs:
        assert t.count(a)==1,(t.count(a),a[:70]); t=t.replace(a,b)
    set_notes(slide,t)
def shp(s,name):
    m=[x for x in s.shapes if x.name==name]; assert len(m)==1,name; return m[0]
def put(s,name,*runs):
    tf=shp(s,name).text_frame
    for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
    rs=tf.paragraphs[0].runs; assert len(rs)>=len(runs),(name,[r.text for r in rs])
    for r,t in zip(rs,runs): r.text=t
    for r in rs[len(runs):]: r._r.getparent().remove(r._r)
def duplicate(prs, src):
    dst=prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid={}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId]=dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg=src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID+'embed',RID+'link',RID+'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    nt,st=dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'),qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag==qn('p:sp')]: nt.append(copy.deepcopy(el))
    return dst
