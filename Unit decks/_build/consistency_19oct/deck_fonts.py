"""Rule 16: every run under 24pt on the listed slides goes to 24pt. Usage: python3 deck_fonts.py <deck> <slide,slide,...>"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn
f, sl = sys.argv[1], [int(x) for x in sys.argv[2].split(',')]
p = Presentation(f); n = 0
for i in sl:
    for el in p.slides[i - 1].shapes._spTree.iter():
        if el.tag in (qn('a:rPr'), qn('a:defRPr'), qn('a:endParaRPr')) and el.get('sz') and int(el.get('sz')) < 2400:
            el.set('sz', '2400'); n += 1
p.save(f); print(f, n)
