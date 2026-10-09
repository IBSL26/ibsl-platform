# -*- coding: utf-8 -*-
"""Final three-way scan: the Unit 6 deck against the two Unit 6 pages. Usage: python3 final_scan_u06.py <deck> <p.html> <f.html>
Reports only. Sections 1 to 4 of the deck are Carol's own slides: differences there are listed for her and are not errors of the build."""
import sys, re, html
from pptx import Presentation
from pptx.util import Pt
deck, P, F = sys.argv[1:4]
def norm(t): return re.sub(r'\s+([?.,;:])', r'\1', re.sub(r'\s+', ' ', t).replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"').replace('–', '-')).strip()
def plain(path):
    t = open(path, encoding='utf-8').read(); t = re.sub(r'<style.*?</style>|<script.*?</script>|<!--.*?-->', ' ', t, flags=re.S)
    return norm(html.unescape(re.sub(r'<[^>]+>', ' ', t)))
vp, vf = plain(P), plain(F); low = (vp + ' ' + vf).lower()
rp, rf = open(P, encoding='utf-8').read(), open(F, encoding='utf-8').read()
print('1. The two pages')
for lab, rx in (('part titles', r'class="acc-t">(.*?)<'), ('sub-lines', r'class="acc-meta">(.*?)<')):
    a = [html.unescape(x) for x in re.findall(rx, rp)][:12]; b = [html.unescape(x) for x in re.findall(rx, rf)][:12]
    print('   %s identical on both pages: %s' % (lab, a == b and len(a) == 12))
    if lab == 'part titles': titles = a
for lab, rx in (('Key learning outcomes', r'<span class="klo">(.*?)</span>'), ('section outcomes', r'<div class="slo-box">.*?</ul>')):
    a = re.findall(rx, rp, re.S); b = re.findall(rx, rf, re.S); print('   %s identical on both pages: %s (%d)' % (lab, [norm(html.unescape(re.sub(r'<[^>]+>', ' ', x))) for x in a] == [norm(html.unescape(re.sub(r'<[^>]+>', ' ', x))) for x in b], len(a)))
pats = [r'Save to Portfolio', r'\blens(es)?\b', r'\b\d+\s*(–|-|to)?\s*\d*\s*(minutes|mins?)\b', r'[Ss]uggested tim', r'[Ww]hiteboard|flip chart|on the board|at the table|[Tt]able discussion|in th(e|is) room', r'pre-?work', r'MyHealth|Rules of the Game|Collective Progress Signal|Three Elements of Enterprise|30-Day|Baseline Simulation',
        r'Strategy2Results(?!®)', r'\bS2R(?!®)', r'\bSIP\b', r'rather than|instead of|in place of']
for name, t in (('participant page', vp), ('facilitator page', vf)):
    hits = [m.group(0) for pat in pats for m in re.finditer(r'.{0,40}' + pat + r'.{0,30}', t)]; print('   leftovers on the %s: %d' % (name, len(hits)), hits[:5])
print('2. The deck')
prs = Presentation(deck); texts = []; small = {}
for i, s in enumerate(prs.slides, 1):
    runs = []
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            t = norm(''.join(r.text for r in para.runs))
            if t: runs.append(t)
            for r in para.runs:
                if r.text.strip() and r.font.size is not None and r.font.size < Pt(24): small[i] = small.get(i, 0) + 1
    texts.append(runs)
print('   slides:', len(texts), '· slides with text under 24pt:', small)
allt = [t for sl in texts for t in sl]
for t in titles:
    num, name = t.split(' — ', 1)
    hit = [x for x in allt if x.startswith(num + ' · ')]
    if not hit: print('   no slide numbered', num, 'for the part:', t)
    elif not any(x.lower() == (num + ' · ' + name).lower() for x in hit): print('   slide title differs · page: "%s" · slide: %s' % (t, hit))
nums = sorted(set(m.group(1) for x in allt for m in [re.match(r'(\d\.\d) · ', x)] if m)); pn = [t.split(' — ')[0] for t in titles]
print('   part numbers on slides that are not on the pages:', [n for n in nums if n not in pn])
klo = [norm(html.unescape(x)) for x in re.findall(r'<span class="klo">(.*?)</span>', rp)]
print('   Key learning outcomes on slide 2 word for word:', all(k in texts[1] for k in klo))
slo = [norm(html.unescape(re.sub(r'<[^>]+>', '', x))) for b in re.findall(r'<div class="slo-box">.*?</ul>', rp, re.S) for x in re.findall(r'<li>(.*?)</li>', b)]
print('   section outcomes on the section slides word for word:', all(any(o in sl for sl in texts) for o in slo))
print('   Section 5 slides (23 to 28): text not on a page word for word:')
for i in range(22, len(texts)):
    for t in texts[i]:
        if len(t) < 15 or re.match(r'^(Step \d · |Section \d|Unit Summary · |Section learning outcomes$|\d\s)', t): continue
        k = re.sub(r'^\S+\s+', '', t) if re.match(r'^[A-Z1-6]\s+\S', t) and i in (23, 24, 25) else t
        if k.rstrip('.').lower() not in low and k.lower().replace(': ', ' — ') not in low: print('      slide %d: %s' % (i + 1, t))
notes = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else '' for s in prs.slides]
print('   leftovers in the notes, by slide:')
for name, pat in (('times', r'\b\d+\s*(–|-|to)?\s*\d*\s*(minutes|mins?)\b|[Ss]uggested tim|\bTime:'), ('room wording', r'at the table|[Tt]able discussion|in th(e|is) room|\bvote\b|[Cc]lick a role'),
                  ('parts that have left', r'MyHealth|Rules of the Game|Collective Progress Signal|Three Elements of Enterprise|Baseline Simulation|Periodic Measures|next 30 days|30-day'),
                  ('contrast wording', r'rather than|instead of|in place of'), ('brand without ®', r'Strategy2Results(?!®)|\bS2R(?!®)')):
    per = [(i + 1, len(re.findall(pat, n))) for i, n in enumerate(notes)]; per = [p for p in per if p[1]]
    print('      %s: %s' % (name, per if per else 'none'))
print('   notes pages that open with a bold heading:', sum(1 for s in prs.slides if s.notes_slide.notes_text_frame.paragraphs[0].runs and s.notes_slide.notes_text_frame.paragraphs[0].runs[0].font.bold), 'of', len(texts))
