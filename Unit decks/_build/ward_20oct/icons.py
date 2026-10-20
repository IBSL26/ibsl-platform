# -*- coding: utf-8 -*-
"""Give the two new WARD slides fitting icons (white icons from the Unit 8 deck). python3 icons.py <u6> <u8> <u8 source deck>"""
import sys, io
from pptx import Presentation
from pptx.oxml.ns import qn
u6, u8, src = sys.argv[1:4]
S = list(Presentation(src).slides)
def blob(si, name): return [sh for sh in S[si].shapes if sh.name == name][0].image.blob
TREND, WARN, LENS, STETH = blob(17, 'Image 0'), blob(17, 'Image 1'), blob(6, 'Image 0'), blob(6, 'Image 1')
def seticon(slide, name, data):
    pic = [sh for sh in slide.shapes if sh.name == name][0]
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(data))
    pic._element.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
p = Presentation(u6); s = list(p.slides)[32]; assert s.shapes[0].text_frame.text.startswith('5.3 · Activity 3')
seticon(s, 'Image 0', TREND); seticon(s, 'Image 1', WARN); seticon(s, 'Image 2', STETH); p.save(u6)
q = Presentation(u8); s = list(q.slides)[24]; assert s.shapes[0].text_frame.text.startswith('5.2 · Testing')
seticon(s, 'Image 0', WARN); seticon(s, 'Image 1', LENS); seticon(s, 'Image 2', STETH); q.save(u8)
print('icons ok')
