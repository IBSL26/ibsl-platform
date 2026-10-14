# -*- coding: utf-8 -*-
"""Is every sentence of the facilitator page (sections 1 to 5) somewhere in the deck notes? Usage: check_notes_cover.py <f page> <deck>"""
import sys, re, html
from pptx import Presentation
f, deck = sys.argv[1:3]
t = open(f, encoding='utf-8').read()
t = t[t.index('<div class="mod-panel" id="mod1">'):t.index('Unit Summary</div>')]
t = re.sub(r'<style.*?</style>|<script.*?</script>', ' ', t, flags=re.S)
t = re.sub(r'</(p|li|td|th|div|h4)>', '\n', t)
t = html.unescape(re.sub(r'<[^>]+>', ' ', t))
def norm(s): return re.sub(r'[^a-z0-9%]+', ' ', s.lower()).strip()
notes = norm(' '.join(s.notes_slide.notes_text_frame.text for s in Presentation(deck).slides))
miss = []
for line in t.split('\n'):
    for sent in re.split(r'(?<=[.?!])\s+', line.strip()):
        n = norm(sent)
        if len(n) > 25 and n not in notes: miss.append(sent.strip())
print(len(miss), 'sentences of the facilitator page not in the notes')
for m in miss: print(' -', m[:160])
