# -*- coding: utf-8 -*-
"""Unit 4 (new 5.3 slide, notes) and Unit 5 (Step 2 highest-risk dimensions) on Carol's decks, 20 Oct 2026.
python3 deck_u4u5.py <unit 4 deck> <unit 5 deck> <out 4> <out 5>"""
import sys, copy
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
A='http://schemas.openxmlformats.org/drawingml/2006/main'
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
def ntext(s): return s.notes_slide.notes_text_frame.text.replace('\x0b','\n')
def set_notes(slide, text):
    body=slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in text.split('\n'):
        p=etree.SubElement(body,'{%s}p'%A)
        if not line.strip():
            e=etree.SubElement(p,'{%s}endParaRPr'%A); e.set('lang','en-GB'); continue
        r=etree.SubElement(p,'{%s}r'%A); rp=etree.SubElement(r,'{%s}rPr'%A); rp.set('lang','en-GB')
        etree.SubElement(r,'{%s}t'%A).text=line
def sub(slide, pairs):
    t=ntext(slide)
    for a,b in pairs:
        assert t.count(a)==1,(t.count(a),a[:70]); t=t.replace(a,b)
    set_notes(slide,t)
def shp(s,name):
    m=[x for x in s.shapes if x.name==name]; assert len(m)==1,name; return m[0]
def put(s,name,*runs):
    tf=shp(s,name).text_frame
    for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
    rs=tf.paragraphs[0].runs; assert len(rs)>=len(runs),(name,[r.text for r in rs])
    for r,t in zip(rs,runs): r.text=t
    for r in rs[len(runs):]: r._r.getparent().remove(r._r)
def duplicate(prs, src):
    dst=prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid={}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId]=dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg=src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID+'embed',RID+'link',RID+'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    nt,st=dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'),qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag==qn('p:sp')]: nt.append(copy.deepcopy(el))
    return dst
src4,src5,out4,out5=sys.argv[1:5]
# ── Unit 4 ──
p=Presentation(src4); S=list(p.slides); assert len(S)==29
assert S[27].shapes[0].text_frame.text.startswith('5.2b')
sub(S[20],[("2. 5.2: the briefing for the Capstone work. Each group takes its own Unit 3 Key Results through the same four steps.",
            "2. 5.2: the briefing for the Capstone work. Each group takes its own Unit 3 Key Results through the same four steps.\n3. 5.3: verification ownership and the executive blind spot. The group assigns each critical condition to an executive and names the checkpoint its team is most likely to overlook.")])
sub(S[27],[('Say when the work is confirmed.\nSay: "When every Key Result has been taken through the four steps, each of you confirms the record."',
            'Say what follows.\nSay: "When every Key Result has been taken through the four steps, your group goes on to 5.3."'),
           ("Close the section with the commitment question.","Lead into 5.3 with the question it answers."),
           ("4. A record shows the condition named at each step for every Key Result. Each participant confirms the record. Print gives a copy.",
            "4. A record shows the condition named at each step for every Key Result. The group then goes on to 5.3.")])
n=duplicate(p,S[27])
put(n,'Text 0','5.3 · Verification Ownership and Executive Blind Spots')
put(n,'Text 2','1 · Read your conditions  ','the Must-Be-True conditions your group named in 5.2')
put(n,'Text 4','2 · Verification ownership  ','the executive, the evidence and the date of the first check')
put(n,'Text 6','3 · Executive blind spot  ','the checkpoint your team is most likely to overlook, and the risk')
put(n,'Text 8','4 · Agree  ','every member agrees; the work feeds boxes 4A to 4F')
put(n,'Text 10','Every condition has an executive who checks it.')
set_notes(n,'''5.3 — VERIFICATION OWNERSHIP AND EXECUTIVE BLIND SPOTS (CAPSTONE WORK)

HOW TO TEACH IT:

1. Say why ownership matters.
Say: "Each Must-Be-True condition your group named in 5.2 holds only while someone checks it.
In 5.3 your group assigns each check to an executive."

2. Teach verification ownership.
Say: "For each critical condition, name the executive accountable for verifying it, the evidence that will confirm it still holds, and the date of the first check."
One executive owns each check, with evidence someone else could inspect and a date.

3. Teach the executive blind spot.
Say: "Section 3 showed that each executive role tends to overlook one checkpoint.
Look around your table: which checkpoint is your executive team most likely to overlook? What is the enterprise risk if it does?"

4. Ask the question 5.3 answers.
Ask: "Which one MBT condition, if left unverified, represents the greatest risk to our strategy’s integrity? Who is responsible for that verification?"

ON THE PORTAL, AFTER THE TEACHING:
Participants work in their groups in 5.3 of the participant file. This is Capstone work, typed by the group’s scribe.
1. The group reads its Must-Be-True conditions from 5.2, listed in 5.3.
2. For each critical condition, the group names the executive accountable, the evidence and the date of the first check.
3. The group selects the checkpoint its executive team is most likely to overlook, and states the enterprise risk.
4. Each participant reads the record and selects Agree. Print gives a copy of the record.

AFTER THE SESSION:
The record reaches you with each participant’s submission, under the heading ABCV–MBT: Stress-Testing the Key Results.
The agreed work feeds boxes 4A to 4F of the team’s Capstone Blueprint: the conditions in 4A to 4D, verification ownership in 4E and the executive blind spot in 4F.''')
sub(S[28],[("- 4 · Integration: each group takes the Key Results it confirmed in Unit 3 through four steps and names the condition that must remain true at each step.\n- 5 · Application: in Cause of Death each participant tests four failed strategies against the Industry Illusion.",
            "- 4 · Integration: in Cause of Death each participant tests four failed strategies against the Industry Illusion.\n- 5 · Application: each group takes the Key Results it confirmed in Unit 3 through four steps, names the condition that must remain true at each step, assigns each critical condition to an executive and names the checkpoint its team is most likely to overlook."),
           ("The confirmed Must-Be-True conditions feed boxes 4A to 4D of the team’s Capstone Blueprint.","The group’s agreed Section 5 work feeds boxes 4A to 4F of the team’s Capstone Blueprint.")])
lst=p.slides._sldIdLst; ids=list(lst); new=ids[-1]; lst.remove(new); lst.insert(28,new)
p.save(out4)
# ── Unit 5 ──
q=Presentation(src5); T=list(q.slides); assert len(T)==35 and T[33].shapes[0].text_frame.text.startswith('Step 2')
old2="Step 2 · ACE-IT. Your group reads its team OCEAVL profile. It then states"
new2="Step 2 · ACE-IT. Your group reads its team OCEAVL profile and names the two dimensions on which the team carries the highest risk for the strategy, with the response to each. It then states"
sub(T[31],[(old2,new2)])
put(T[33],'Text 12','First name your two highest-risk dimensions. Then state the behaviour that answers each risk.')
sub(T[33],[('''2. Ask the linking question.
Ask: "Which level in your team’s profile carries the highest risk for this change? Which ACE-IT behaviour answers it?"''',
'''2. Name the two highest-risk dimensions.
Say: "From your team’s profile, your group names the two dimensions on which your executive team carries the highest risk for the strategy, and the response your team adopts for each.
The profile gives the risk and the response at your team’s level."
Ask: "Which level in your team’s profile carries the highest risk for this change? Which ACE-IT behaviour answers it?"'''),
 ("2. The group reads the guide in each card. It is the lesson from part 1.3.4.\n3. For each behaviour,",
  "2. The group names the two dimensions on which its executive team carries the highest risk for the strategy, and the response it adopts for each.\n3. The group reads the guide in each card. It is the lesson from part 1.3.4.\n4. For each behaviour,"),
 ("4. Each participant reads the record and selects Agree.","5. Each participant reads the record and selects Agree."),
 ("The confirmed record feeds boxes 5B and 5C of the team’s Capstone Blueprint: the Shift, Stake and Step in 5B, the expected leadership behaviours in 5C.",
  "The agreed record feeds boxes 5A to 5C of the team’s Capstone Blueprint: the two highest-risk dimensions in 5A, the Shift, Stake and Step in 5B, the expected leadership behaviours in 5C.")])
sub(T[34],[("- 5 · Application: each group states its Shift, Stake and Step for the strategy, reads its team OCEAVL profile and states the ACE-IT behaviours it expects from the leadership team.",
            "- 5 · Application: each group states its Shift, Stake and Step for the strategy, reads its team OCEAVL profile, names its two highest-risk dimensions with the response to each, and states the ACE-IT behaviours it expects from the leadership team.")])
q.save(out5)
print('saved', len(p.slides), len(q.slides))
