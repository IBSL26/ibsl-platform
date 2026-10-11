# -*- coding: utf-8 -*-
"""Unit 7 deck: Carol's own deck is the master (23 slides, saved 11 October 2026 12:25). This script changes only the slides her amendments of 11 October touch.
Usage: python3 carol_deck_u07.py <Carol's deck> <participant page> <facilitator page> <pictures folder> <out.pptx>
Untouched, shape and note: slides 1 to 12, 14 and 17 of her deck (title, outcomes, journey, Sections 1 and 2, 3.1, 4.1).
Changed:
  - her "4.2 · Score → Evaluate → Resolve" slide moves to Section 1, after "1.4 · The Alignment Dialogue", as "1.4 · Score → Evaluate → Resolve";
  - a copy of her 4.1 slide becomes "1.4 · The GROW coaching model", and another copy becomes "4.2 · Practising the Alignment Dialogue";
  - her "Reflections" slide after 3.1 becomes "Section 3 Reflections", and her empty slide after 4.2 becomes "Section 4 Reflections" (pictures of the participant page);
  - the second outcome on the Section 3 and Section 4 slides, and the Section 5 slides follow the pages: the opener, the case slide (now two Objectives) and a copy of her 4.1 slide as the four steps; her "Role play: seven steps" slide leaves;
  - the closing slide loses the five-day window.
Notes of the changed slides follow the form of the uploaded Unit 2 to 6 decks: heading, HOW TO TEACH IT, numbered steps with Say / Ask, ON THE PORTAL, AFTER THE TEACHING.
Teaching comes first in every note; a question is put only on content that has been taught. No times."""
import sys, re, copy, html, json, os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image
import u7_content as C
deck, P, F, PICS, out = sys.argv[1:6]
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
def plain(path):
    t = open(path, encoding='utf-8').read()
    t = re.sub(r'<style.*?</style>|<script.*?</script>|<!--.*?-->', ' ', t, flags=re.S)
    t = re.sub(r'</?(?:em|strong|b|i)\b[^>]*>', '', t)
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', t)))
vp, vf = plain(P), plain(F)
def tx(s): return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', '', s))).strip()
def on_pages(s, where='pf'):
    s = tx(s)
    for w in where: assert s in (vp if w == 'p' else vf), ('not on the %s page: ' % w) + s
    return s

prs = Presentation(deck)
S = list(prs.slides)
assert len(S) == 23, len(S)
def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def first(s): t = texts(s); return t[0].text_frame.text if t else ''
assert first(S[7]) == '1.4 · The Alignment Dialogue' and first(S[13]).startswith('3.1 · CXO') and first(S[14]) == 'Reflections' and first(S[16]).startswith('4.1 · The four-step')
assert first(S[17]) == '4.2 · Score → Evaluate → Resolve' and not texts(S[18]) and texts(S[19])[3].text_frame.text == 'Role Play — Establishing Delivery Expectations'
assert first(S[20]).startswith('Role play · Case') and first(S[21]) == 'Role play: seven steps' and 'Unit Summary' in first(S[22])

def set_text(shape, *run_texts):
    ps = shape.text_frame.paragraphs; runs = ps[0].runs
    assert len(ps) == 1 or all(not p.runs for p in ps[1:]), shape.name
    assert len(runs) == len(run_texts), (shape.name, [r.text for r in runs], run_texts)
    for r, t in zip(runs, run_texts): r.text = t
def set_notes(slide, lines):
    """Notes as one paragraph for each line. A line that starts with § is a heading or a step title, set in bold."""
    body = slide.notes_slide.notes_text_frame._txBody
    for p in body.findall(qn('a:p')): body.remove(p)
    for line in lines:
        p = etree.SubElement(body, '{%s}p' % A)
        bold = line.startswith('§'); text = line[1:] if bold else line
        if not text.strip():
            e = etree.SubElement(p, '{%s}endParaRPr' % A); e.set('lang', 'en-GB'); e.set('dirty', '0'); continue
        r = etree.SubElement(p, '{%s}r' % A); rp = etree.SubElement(r, '{%s}rPr' % A); rp.set('lang', 'en-GB')
        if bold: rp.set('b', '1')
        rp.set('dirty', '0'); etree.SubElement(r, '{%s}t' % A).text = text
    for l in lines: assert not re.search(r'\b(minutes?|mins?|Timing|working days|five-day)\b', l), l
def note_paras(slide): return slide.notes_slide.notes_text_frame._txBody.findall(qn('a:p'))
def para_text(p): return ''.join(t.text or '' for t in p.iter(qn('a:t')))
def note_sub(slide, old, new):
    """One exact replacement inside the notes of a slide that is otherwise left as Carol wrote it."""
    hits = [t for t in slide.notes_slide.notes_text_frame._txBody.iter(qn('a:t')) if old in (t.text or '')]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, (len(hits), old[:60])
    hits[0].text = hits[0].text.replace(old, new)
def duplicate(src):
    """A copy of a slide, with its background and pictures, added at the end of the deck."""
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes): sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith('/image'): rid[rel.rId] = dst.part.relate_to(rel._target, rel.reltype)
    for sh in src.shapes: dst.shapes._spTree.append(copy.deepcopy(sh._element))
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None and dst._element.cSld.find(qn('p:bg')) is None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in dst._element.iter():
        for k in (RID + 'embed', RID + 'link', RID + 'id'):
            if el.get(k) in rid: el.set(k, rid[el.get(k)])
    left = [el.get(RID + 'embed') for el in dst._element.iter() if el.get(RID + 'embed')]
    assert all(x in dst.part.rels for x in left), left
    nt, st = dst.notes_slide._element.cSld.spTree, src.notes_slide._element.cSld.spTree
    for el in [e for e in nt if e.tag in (qn('p:sp'), qn('p:pic'))]: nt.remove(el)
    for el in [e for e in st if e.tag == qn('p:sp')]: nt.append(copy.deepcopy(el))
    assert dst.notes_slide.notes_text_frame is not None
    return dst
GOLD, WHITE = RGBColor(0xC6, 0xA2, 0x4C), RGBColor(0xFF, 0xFF, 0xFF)
REFL = json.load(open(os.path.join(PICS, 'reflections.json'), encoding='utf-8'))
def reflections(slide, title, parts, model_title):
    """A "Section N Reflections" slide in the form of the other decks: the title, then one picture of the participant page for each part, with the part number in a gold circle."""
    tt = texts(slide)
    if tt: set_text(tt[0], title)
    else:
        t = copy.deepcopy(model_title._element); slide.shapes._spTree.append(t)
        runs = t.findall('.//' + qn('a:r')); runs[0].find(qn('a:t')).text = title
        for r in runs[1:]: r.getparent().remove(r)
    top = Inches(1.72); pic_left, pic_w, dia, gap = Inches(1.85), Inches(10.85), Inches(0.95), Inches(0.32)
    for part in parts:
        path = os.path.join(PICS, REFL[part]['file']); w, h = Image.open(path).size; ph_ = int(pic_w * h / w)
        p = slide.shapes.add_picture(path, pic_left, top, width=pic_w)
        p._element.nvPicPr.cNvPr.set('descr', (REFL[part]['label'] + '. ' + REFL[part]['prompt']).strip()); p.name = 'Reflection ' + part
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), top + int((min(ph_, Inches(1.6)) - dia) / 2), dia, dia); o.name = 'Part ' + part
        o.fill.solid(); o.fill.fore_color.rgb = GOLD; o.line.color.rgb = GOLD; o.shadow.inherit = False
        tf = o.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        par = tf.paragraphs[0]; par.alignment = PP_ALIGN.CENTER; r = par.add_run(); r.text = part
        r.font.size = Pt(24); r.font.bold = True; r.font.name = 'Calibri'; r.font.color.rgb = WHITE
        top += ph_ + gap
    assert top - gap <= prs.slide_height - Inches(0.15), 'pictures run off the slide'
Q = lambda s: '"' + tx(s).strip('“”"') + '"'

# ═══ 1.4 · Score → Evaluate → Resolve (her 4.2 slide, moved to Section 1) ═══════════════════════════════
ser = S[17]; set_text(texts(ser)[0], '1.4 · Score → Evaluate → Resolve')
set_notes(ser, ['§1.4 — THE ALIGNMENT DIALOGUE · SCORE → EVALUATE → RESOLVE', '', '§HOW TO TEACH IT:', '',
  '§1. Say what the sequence is for.',
  'Say: "' + on_pages('When the alignment score is below 9, the leader treats the number as data — it signals that something is unresolved.', 'f'),
  on_pages('The Score → Evaluate → Resolve sequence transforms that signal into a resolved commitment.') + '"',
  'The scale carries no judgement of performance. It helps the leader quickly understand whether the expectation has moved beyond communication into genuine ownership.', '',
  '§2. Teach Score It.',
  'Say: "' + on_pages('The individual provides a numerical assessment of their level of alignment. The score acts as a signal of readiness for ownership.') + '"',
  'Read the four rows on the slide:',
  '9–10: ' + on_pages('Clear ownership and readiness to deliver') + '. ' + on_pages('Proceed; confirm the commitment and document') + '.',
  '7–8: ' + on_pages('Partial alignment — some uncertainty remains') + '. ' + on_pages('Move to Evaluate → surface what is unresolved') + '.',
  '5–6: ' + on_pages('Significant concerns affecting confidence in delivery') + '. ' + on_pages('Move to Evaluate → expect multiple concerns to surface') + '.',
  'Below 5: ' + on_pages('Expectation is not yet viable in its current form') + '. ' + on_pages('Do not proceed — the expectation must be redesigned') + '.', '',
  '§3. Teach Evaluate It.',
  'Say: "' + on_pages('If the score is below 9, the leader explores the reasons behind the rating. The objective is to understand what is preventing full alignment.', 'f') + '"',
  on_pages('The leader guides the conversation using diagnostic questions:', 'f'),
  Q('What influenced the score you gave?'), Q('What feels unclear about the expectation?'), Q('What concerns do you have about delivering this outcome?'), Q('Where do you anticipate obstacles or constraints?'),
  'Common concerns include: ' + '; '.join(on_pages(x).lower() for x in ('Unclear definition of the outcome or success criteria', 'Cross-functional dependencies not yet confirmed', 'Capacity or resource constraints affecting feasibility', 'Confidence in the delivery pathway or timeline')) + '.',
  'Say: "' + on_pages('The purpose of Evaluate is to make hidden concerns visible before execution begins.') + '"',
  'Leaders should be able to hear raw frustration without immediately defending the system or correcting the emotion. Credibility grows when leaders acknowledge the gap between intention and execution and work to close it.', '',
  '§4. Teach Resolve It.',
  'Say: "' + on_pages('Once the sources of misalignment are understood, the leader works with the individual to resolve them.', 'f'),
  'The guiding coaching question becomes: ‘' + on_pages('What would help move your alignment from where it is now to a 9 or 10?') + '’"',
  on_pages('This question shifts the conversation toward solutions. Resolution may involve:'),
  '; '.join(on_pages(x)[0].lower() + on_pages(x)[1:] for x in ('Clarifying the outcome definition or success criteria', 'Adjusting priorities or realistic timelines', 'Identifying additional support or resources required', 'Confirming collaboration across the Expectations Chain')) + '.',
  'Say: "' + on_pages('Once concerns are resolved, the leader asks the alignment question again to confirm the expectation has moved into clear ownership.', 'f') + '"',
  'Speaking before understanding is a common leadership failure. People perform better when difficulty does not automatically threaten their belonging or worth.', '',
  '§5. Teach what the leader captures after the dialogue.',
  'Say: "After the dialogue, the leader captures three reflections:',
  '1. Alignment Rating — Where does the commitment currently stand?', '2. Primary Concern Identified — What surfaced during the discussion?', '3. Action for Elevation — What will move alignment upward?"',
  'Documenting these reflections converts a conversation into a visible commitment record.', '',
  '§6. Check understanding.',
  'Ask: "When someone on your team gives you a 6 out of 10 on alignment, what does your instinct tell you to do? And what should you do?"',
  'Take two answers in the chat. Bring each one back to Evaluate It: understand the score before you act on it.', '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants read Score → Evaluate → Resolve in part 1.4 of the participant file.'])

# ═══ 1.4 · The GROW coaching model (a copy of her 4.1 slide) ════════════════════════════════════════════
grow = duplicate(S[16]); t = texts(grow); set_text(t[0], '1.4 · The GROW coaching model')
for i, (letter, name, mean, in_p, in_f, qs) in enumerate(C.GROW):
    set_text(t[i + 1], '%s · %s  ' % (letter, name), on_pages(mean).rstrip('.').lower())
lines = ['§1.4 — THE ALIGNMENT DIALOGUE · THE GROW COACHING MODEL', '', '§HOW TO TEACH IT:', '',
  '§1. Introduce the model.',
  'Say: "' + on_pages('GROW is the coaching model the leader uses when holding the Alignment Dialogue.', 'f'),
  'It is a widely used coaching model, popularised by Sir John Whitmore in Coaching for Performance.',
  on_pages('Its four stages are Goal, Reality, Options and Will.'),
  on_pages('At each stage the leader asks questions and the individual does the thinking, so the answers the individual reaches are their own.', 'f') + '"',
  'Responding with a question returns ownership of the response to the person and gives them space to examine their own position without defending against an external judgement. Questions open where corrections close.', '']
for i, (letter, name, mean, in_p, in_f, qs) in enumerate(C.GROW):
    lines += ['§%d. Teach %s · %s.' % (i + 2, letter, name), 'Say: "' + on_pages(mean) + ' ' + on_pages(in_f, 'f') + '"', 'The questions the leader asks:'] + [Q(on_pages(q)) for q in qs] + ['']
lines += ['§6. Show how the two fit together.',
  'Say: "' + on_pages('Score → Evaluate → Resolve is the sequence of the dialogue. GROW is how the leader coaches inside it.', 'f'),
  'Goal at Score It. Reality at Evaluate It. Options and Will at Resolve It."', '',
  '§7. Say where it is practised.', 'Say: "You practise this dialogue in a role play in Section 4."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants read the GROW coaching model in part 1.4 of the participant file.',
  C.REF + C.ACT['1.4']]
set_notes(grow, lines)

# ═══ Section 3: the second outcome, and the reflection slide ═════════════════════════════════════════════
s3 = S[12]; old3, new3 = C.SLO_EDITS[2]; assert texts(s3)[6].text_frame.text == old3; set_text(texts(s3)[6], on_pages(new3))
note_sub(s3, '2. ' + old3, '2. ' + new3)
note_sub(s3, '\n\nTime: the facilitator file gives no suggested time for this section.', '')
r3 = S[14]; reflections(r3, 'Section 3 Reflections', ['3.1'], None)
q31 = on_pages(REFL['3.1']['prompt'], 'p')
set_notes(r3, ['§SECTION 3 REFLECTIONS', '', '§HOW TO USE THIS SLIDE:', '',
  '§1. Show the slide and read the prompt aloud.', '3.1 Reflection — Your Expectation Gap: "' + q31 + '"', '',
  '§2. Say what the reflections are for.',
  'Point learners to the portal section that they will be required to complete with as much depth as possible.',
  'The reflections build the learning portfolio they receive at the end of the programme.', '',
  'Say: "You write this reflection on the portal after the teaching, with as much depth as you can."', '',
  '§3. Set the standard.',
  'Say: "Name the expectation, the person or function you need to establish it with, and the alignment dialogue you will have."',
  'Watch for: a reflection that names a team where a person is needed. Ask for one role.', '',
  '§4. Bridge to Section 4.',
  'Say: "You have seen where the expectation gap opens in your own role. Section 4 gives the practice every leader uses to close it."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants write this reflection in Section 3 of the participant file, at the foot of part 3.1.'])

# ═══ Section 4: the second outcome, 4.2 as practice, and the reflection slide ════════════════════════════
s4 = S[15]; old4, new4 = C.SLO_EDITS[3]; assert texts(s4)[6].text_frame.text == old4; set_text(texts(s4)[6], on_pages(new4))
note_sub(s4, '2. ' + old4, '2. ' + new4)
prac = duplicate(S[16]); t = texts(prac); set_text(t[0], '4.2 · Practising the Alignment Dialogue')
for sh, (a, b) in zip(t[1:], (('1 · Set up  ', 'one expectation, a leader and an owner'), ('2 · Demonstrate  ', 'the facilitator plays the leader'),
                             ('3 · Practise  ', 'participants lead, in the roles chosen'), ('4 · Debrief  ', 'which question moved the score'))): set_text(sh, a, b)
R = C.ROLEPLAY
lines = ['§4.2 — PRACTISING THE ALIGNMENT DIALOGUE (ROLE PLAY IN CLASS)', '', on_pages(R['purpose'], 'f'), '', '§HOW TO RUN IT:', '',
  '§1. Prepare three things.'] + [on_pages(x, 'f') for x in R['decide']] + ['']
names = ['Set it up.', 'Give the observers their task.', 'Round 1 · You demonstrate.', 'Pause.', 'Round 2 · Participants lead.', 'Debrief.', 'Point to the Portfolio work.']
for i, (nm, stp) in enumerate(zip(names, R['steps'])):
    full = on_pages(stp, 'f'); assert full.startswith(nm), (nm, full[:40])
    lines += ['§%d. %s' % (i + 2, nm), full[len(nm):].strip().replace('“', '"').replace('”', '"'), '']
lines += ['§Owner briefs for the example of part 4.1:'] + ['%s: %s' % (a, on_pages(b, 'f')) for a, b in R['briefs']] + ['',
  '§Watch for:'] + [on_pages(w, 'f').replace('“', '"').replace('”', '"') for w in R['watch']] + ['',
  '§ON THE PORTAL, AFTER THE TEACHING:', on_pages(C.ACT_42_LEAD, 'f')] + ['%d. %s' % (i + 1, on_pages(x, 'f')) for i, x in enumerate(C.ACT_42_STEPS)] + [on_pages(C.ACT_42_TAIL, 'f')]
set_notes(prac, lines)
r4 = S[18]
# her empty slide carries the dark background; the reflection slides of every deck are on the light one, so it takes the background of her "Reflections" slide
bg_old, bg_new = r4._element.cSld.find(qn('p:bg')), S[14]._element.cSld.find(qn('p:bg'))
assert bg_new is not None and not [e for e in bg_new.iter() if e.get(RID + 'embed')]
if bg_old is not None: r4._element.cSld.remove(bg_old)
r4._element.cSld.insert(0, copy.deepcopy(bg_new))
reflections(r4, 'Section 4 Reflections', ['4.1', '4.2'], texts(S[14])[0])
q41 = on_pages(REFL['4.1']['prompt'], 'p')
set_notes(r4, ['§SECTION 4 REFLECTIONS', '', '§HOW TO USE THIS SLIDE:', '',
  '§1. Show the slide and read the first prompt aloud.', '4.1 Reflection — Which Step Do You Most Often Skip?: "' + q41 + '"', '',
  '§2. Explain the Portfolio work of part 4.2.',
  'Say: "In part 4.2 you prepare an Alignment Dialogue with your board chair, using the GROW model.',
  'You choose the Key Result for which you may need that dialogue.',
  'The scenario is on your page: your board chair has disagreed with the current metric and timeframe of that Key Result and gives a 5.',
  'You write your side of the dialogue under Goal, Reality, Options and Will, to move your chair to a 10."', '',
  '§3. Say what the reflections are for.',
  'Point learners to the portal section that they will be required to complete with as much depth as possible.',
  'The reflections build the learning portfolio they receive at the end of the programme.', '',
  'Say: "You write the reflection and the Portfolio work on the portal after the teaching, with as much depth as you can."', '',
  '§4. Set the standard.',
  'Say: "Write the questions you would ask in your own words. A question you would never say aloud does not belong in your dialogue."',
  'Watch for: answers that argue the case for the metric. Ask which GROW question comes first.', '',
  '§5. Bridge to Section 5.',
  'Say: "You have practised the dialogue. In Section 5 your group establishes delivery expectations for two of its own Objectives."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants write the 4.1 reflection in Section 4 of the participant file, at the foot of part 4.1.',
  on_pages(C.ACT_42_LEAD, 'f')] + ['%d. %s' % (i + 1, on_pages(x, 'f')) for i, x in enumerate(C.ACT_42_STEPS)])

# ═══ Section 5 ═══════════════════════════════════════════════════════════════════════════════════════════
s5 = S[19]; t = texts(s5); set_text(t[3], on_pages(C.S5_TITLE))
slo = [on_pages(C.SLO_EDITS[4][1]), on_pages(C.SLO_EDITS[5][1])]; assert t[5].text_frame.text == C.SLO_EDITS[4][0] and t[6].text_frame.text == C.SLO_EDITS[5][0]
set_text(t[5], slo[0]); set_text(t[6], slo[1])
H = C.HOW_S5['p']
set_notes(s5, ['§SECTION 5 · APPLICATION — IN PRACTICE', '', '§HOW THIS SECTION RUNS:', on_pages(H['lead'], 'p')] + [on_pages(x, 'p') for x in H['steps']] + [on_pages(H['tail'], 'p'), '',
  '§SECTION LEARNING OUTCOMES:', '1. ' + slo[0], '2. ' + slo[1], '', '§HOW TO OPEN THE SECTION:', '',
  '§1. Say what the section does.', 'Say: "In this section your group establishes delivery expectations for two Objectives of the strategy in your Capstone Blueprint.',
  'For each Objective, you take one related Key Result through the four steps of the Expectation Architecture."', '',
  '§2. Say how the group works.', 'Say: "All four steps are group work.', 'Your group agrees each entry.', 'One member acts as scribe.', 'Every member then types the agreed entries into their own page, in the session or after it."', '',
  '§3. Say where the work goes.', 'Say: "This is Capstone work. Your group’s confirmed record feeds your team’s Capstone Blueprint."',
  'The drop-down lists of Step 1 hold the group’s Enterprise OKRs and its Success in Practice statements. Each group confirms them in Unit 3, part 4.2, and in Unit 2 first.'])
two = S[20]; t = texts(two); set_text(t[0], 'Section 5 · Capstone work'); set_text(t[1], 'Two Objectives. One Key Result for each. Four steps.')
set_text(t[2], 'Chosen from the Enterprise OKRs your group confirmed in Unit 3.')
set_notes(two, ['§SECTION 5 — TWO OBJECTIVES, ONE KEY RESULT FOR EACH (CAPSTONE WORK)', '', '§HOW TO TEACH IT:', '',
  '§1. Explain the choice.',
  'Say: "' + on_pages(C.HOW_STEP_P[0][0], 'p') + '"',
  'The lists hold the Enterprise OKRs the group confirmed in Unit 3 and the Success in Practice statements it confirmed in Unit 2.', '',
  '§2. Name the owner.',
  'Say: "' + on_pages(C.HOW_STEP_P[0][1], 'p'), 'The four steps establish the expectation with that owner."', '',
  '§3. Set the standard.',
  'Say: "Work on your real strategy. The closer you stay to it, the more transferable the skill."',
  'Watch for: an owner named as a department. Ask which executive role answers for the Key Result.', '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants work as a group in Step 1 of Section 5 of the participant file. This is Capstone work.',
  tx(C.COMPLETES_F[0][0])])
# "The four steps": a copy of her 4.1 slide (the same four steps); her "Role play: seven steps" slide leaves the deck
four = duplicate(S[16]); t = texts(four); set_text(t[0], 'Section 5 · The four steps')
for i, (p, d) in enumerate(zip(C.PARTS, ('the Future and the leader framing', 'what, by when, to what standard', 'Upline, Downline and External', 'the score and what resolves it'))):
    set_text(t[i + 1], '%d · %s  ' % (i + 1, on_pages(p[1])), d)
G = C.GUIDE_STEP_F; HS = C.HOW_STEP_P
set_notes(four, ['§SECTION 5 — THE FOUR STEPS (CAPSTONE WORK)', '', '§HOW TO TEACH IT:', '',
  '§1. Step 1 · Start With the Future.',
  'Say: "' + on_pages(HS[0][2], 'p'), on_pages('Write the leader framing: the sentence that you, as the leader, say to the owner to connect the Key Result to that Future.', 'p') + '"',
  on_pages(G[0][0], 'f'), '',
  '§2. Step 2 · Clarify the Delivery Expectation.',
  'Say: "' + ' '.join(on_pages(x, 'p') for x in HS[1]) + '"', on_pages(G[1][0], 'f'), '',
  '§3. Step 3 · Identify the Expectations Chain.',
  'Say: "' + ' '.join(on_pages(x, 'p') for x in HS[2]) + '"', on_pages(G[2][0], 'f'), '',
  '§4. Step 4 · Confirm Ownership.',
  'Say: "' + ' '.join(on_pages(x, 'p') for x in HS[3][:3]) + '"', on_pages(G[3][0], 'f'), on_pages(G[3][1], 'f'), '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants work as a group in the four steps of Section 5 of the participant file. This is Capstone work.'] +
  ['Step %d: %s' % (i + 1, ' '.join(tx(x) for x in C.COMPLETES_F[i])) for i in range(4)] +
  ['Each participant then reads the group’s record and selects Confirm. The confirmed record feeds the team’s Capstone Blueprint.'])
# closing slide: the five-day window leaves
cl = S[22]; t = texts(cl); assert 'five working days' in t[2].text_frame.text
set_text(t[2], 'Capture the score, the primary concern and the action for elevation when you hold that alignment dialogue.')
m = re.search(r'Close by asking each leader to name one expectation[^\n]*', ''.join(t.text or '' for t in cl.notes_slide.notes_text_frame._txBody.iter(qn('a:t')))); assert m and 'five working days' in m.group(0)
note_sub(cl, m.group(0), on_pages(C.CLOSE_F, 'f'))

# ═══ order of the slides ═════════════════════════════════════════════════════════════════════════════════
lst = prs.slides._sldIdLst; ids = list(lst); assert len(ids) == 26
by = {id(s): el for s, el in zip(list(prs.slides), ids)}
new_order = S[:8] + [ser, grow] + S[8:17] + [prac] + S[18:21] + [four, S[22]]
assert len(new_order) == 25 and len({id(s) for s in new_order}) == 25
gone = by[id(S[21])]; prs.part.drop_rel(gone.get(RID + 'id'))          # "Role play: seven steps"
for el in ids: lst.remove(el)
for s in new_order: lst.append(by[id(s)])
prs.save(out)
print(len(prs.slides), 'slides ->', out)
for i, s in enumerate(prs.slides, 1): print('%2d' % i, (' · '.join(x.text_frame.text for x in texts(s)[:2]) or [sh.name for sh in s.shapes][:2])[:90] if True else '')
