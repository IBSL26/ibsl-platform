"""Final three-way scan: the Unit 5 deck against the two Unit 5 pages. Usage: python3 final_scan.py <deck> <p.html> <f.html>
1. Key learning outcomes and section outcomes word for word on both pages. 2. Every piece of slide text of 15 characters or more is on a page
(deck-only labels are listed in DECK_ONLY). 3. No slide text under 24pt; every notes page opens with a bold heading.
4. Leftovers (times, "lens", Save to Portfolio, room wording, pre-work, the old unit name) in pages, slides and notes.
5. The two pages carry the same part titles and sub-lines. 6. Every part title is on a slide word for word, and every reflection question is in the notes."""
import sys, re, html
from pptx import Presentation
from pptx.util import Pt
deck, P, F = sys.argv[1:4]
def norm(t): return re.sub(r'\s+([?.,;:])', r'\1', re.sub(r'\s+', ' ', t).replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"').replace('–', '-')).strip()
def plain(path, js=True):
    t = open(path, encoding='utf-8').read()
    t = re.sub(r'<style.*?</style>', ' ', t, flags=re.S)
    if not js: t = re.sub(r'<script.*?</script>', ' ', t, flags=re.S)
    return norm(html.unescape(re.sub(r'<[^>]+>', ' ', t)).replace("\\'", "'"))
tp, tf = plain(P), plain(F); vp, vf = plain(P, False), plain(F, False)
low = (tp + ' ' + tf).lower()
prs = Presentation(deck); bad = 0
DECK_ONLY = [r'^Key learning outcomes$', r'^The unit journey$', r'^Section \d', r'^Section learning outcomes$', r'Facilitator deck$', r'^Module 3 · Influence · Unit 5$',
  r'^\d\.\d[ab]? · ', r'^Principle \d — ', r'^Step \d · ', r'^Unit Summary · The closing question$',
  # 1.2 overview: three of the five taglines are from Carol's original content (her instruction of 2 October 2026); the other two are on the pages
  r'^the tools that create alignment$', r'^the behavioural dynamics that influence execution$', r'^the observable standard for alignment$',
  # 1.3b: check, tool and job on one label (each word is on the pages)
  r'^(Mind · 3S · Explain|Heart · SCARF · Involve|Hands · STAT · Equip|Habit · ACE-IT · Reinforce)$']
print('1. Outcomes')
texts = []
for i, s in enumerate(prs.slides, 1):
    runs = []
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.font.size is not None and r.font.size < Pt(24): bad += 1; print(f'  slide {i}: {r.font.size.pt}pt "{r.text[:40]}"')
                if r.font.size is None and r.text.strip(): bad += 1; print(f'  slide {i}: no size set "{r.text[:40]}"')
                if r.text.strip(): runs.append(norm(r.text))
    texts.append(runs)
    nt = s.notes_slide.notes_text_frame.paragraphs[0]
    if not (nt.runs and nt.runs[0].font.bold): bad += 1; print(f'  slide {i}: notes do not open with a bold heading')
def need(label, t, where='pf'):
    global bad
    miss = [w for w in where if norm(t) not in (tp if w == 'p' else tf)]
    if miss: bad += 1; print(f'  DIFFERS · {label} · not on {"/".join("participant" if m == "p" else "facilitator" for m in miss)} page: {t[:110]}')
for t in texts[1][1:]: need('KLO', t)
n_slo = 0
for i, sl in enumerate(texts, 1):
    if 'Section learning outcomes' in sl:
        k = sl.index('Section learning outcomes'); [need(f'section outcome, slide {i}', t) for t in sl[k + 1:k + 3]]; n_slo += 1
        need(f'section title, slide {i}', sl[k - 1])
assert n_slo == 5 and len(texts[1]) == 5
print('2. Slide text on the pages')
seen = 0
for i, sl in enumerate(texts, 1):
    for t in sl:
        if len(t) < 15 or any(re.search(p, t) for p in DECK_ONLY): continue
        seen += 1
        k = t.rstrip('.').lower()
        if k not in low and k.strip('"') not in low:
            bad += 1; print(f'  slide {i}: not on a page: {t[:120]}')
print('   strings checked:', seen)
print('3. Leftovers')
pats = [r'Save to Portfolio', r'Go to Submit', r'\blens(es)?\b', r'\b\d+\s*(–|-|to)?\s*\d*\s*(minutes|mins?)\b', r'[Ss]uggested (section )?time', r'[Ww]hiteboard|flip chart|on the board|one sheet|[Cc]irculate',
        r'in th(e|is) room', r'pre-?work', r'as they are today', r'Aligning Hearts', r'KISS table', r'Strategy2Results(?!®)', r'\bS2R(?!®)', r'\bSIP\b', r'rather than|instead of|in place of']
notes = '\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides)
for name, t in (('participant page', vp), ('facilitator page', vf), ('deck slides', ' | '.join(' | '.join(x) for x in texts)), ('deck notes', notes)):
    hits = [m.group(0) for pat in pats for m in re.finditer(r'.{0,40}' + pat + r'.{0,30}', t)]
    print(f'   {name}: {len(hits)}', hits[:6]); bad += len(hits)
print('4. Page pair')
rp, rf = open(P, encoding='utf-8').read(), open(F, encoding='utf-8').read()
parts = {}
for lab, rx in (('part titles', r'class="acc-t[^"]*">(.*?)<'), ('sub-lines', r'class="acc-meta">(.*?)<')):
    a = [html.unescape(x) for x in re.findall(rx, rp)][:16]; b = [html.unescape(x) for x in re.findall(rx, rf)][:16]   # the sixteen parts of Sections 1 to 4
    same = a == b and len(a) == 16; print('   ', lab, 'identical on both pages:', same, '' if same else (a, b)); bad += (not same)
    if lab == 'part titles': titles = a
print('5. Part titles and reflection questions in the deck')
firsts = [sl[0] if sl else '' for sl in texts]; alltitles = [t for sl in texts for t in sl[:2]]
for t in titles:
    num, name = t.split(' — ', 1)
    hit = [x for x in alltitles if re.match(re.escape(num) + r'[ab]? · ', x) and (name in x or (num == '1.3' and x.startswith('1.3a · Change Management:')))]
    if not hit: bad += 1; print('   part title not on a slide word for word:', t)
qs = [norm(html.unescape(re.sub(r'<[^>]+>', ' ', x))) for x in re.findall(r'<label class="wp-label"[^>]*for="ref\d+">(.*?)</label>', rp)]
assert len(qs) == 15
nn = norm(notes)
for q in qs:
    if nn.count(q) < 2: bad += 1; print('   reflection question not in the notes twice (its part and its reflection slide):', q[:80])
print('   part titles on slides:', len(titles), '· reflection questions in the notes:', len(qs))
print('slides:', len(texts), '· problems:', bad)
