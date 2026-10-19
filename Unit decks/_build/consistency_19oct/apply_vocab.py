# -*- coding: utf-8 -*-
"""Usage: python3 -B apply_vocab.py <portal folder> <out folder>. Writes changed pages, decks and capstone_P.html to <out>,
plus vocab_log.json (unit, file, where, before, after) for the review file. Line endings and encoding kept."""
import sys, os, re, json, glob, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from vocab import U6, U7, U8, U9, CAPSTONE, DECK_EXACT
from pptx import Presentation
from pptx.oxml.ns import qn
src, out = sys.argv[1], sys.argv[2]; os.makedirs(out, exist_ok=True)
LOG = []; HITS = {}
ENT = [('’', '&rsquo;'), ('—', '&mdash;'), ('·', '&middot;'), ('–', '&ndash;')]
def variants(a, b):
    v = [(a, b)]
    a2, b2 = a, b
    for c, e in ENT: a2, b2 = a2.replace(c, e), b2.replace(c, e)
    if a2 != a: v.append((a2, b2))
    return v
def apply(t, table, tag):
    """Each pair replaces original text only: new text is held as a token until the end, so no pair rewrites another's result."""
    toks = []
    for a, b in table:
        for x, y in variants(a, b):
            n = t.count(x)
            if n:
                HITS[(tag, a)] = HITS.get((tag, a), 0) + n
                toks.append(y); t = t.replace(x, '\x00%d\x00' % (len(toks) - 1))
    return re.sub('\x00(\\d+)\x00', lambda m: toks[int(m.group(1))], t)
def sentences_diff(a, b):
    import difflib
    sa = re.split(r'(?<=[.:?!])\s+|\n', a); sb = re.split(r'(?<=[.:?!])\s+|\n', b)
    sm = difflib.SequenceMatcher(None, sa, sb, autojunk=False); res = []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op != 'equal': res.append((' '.join(sa[i1:i2]).strip(), ' '.join(sb[j1:j2]).strip()))
    return res
def vis(h):
    import html as H
    h = re.sub(r'<script.*?</script>|<style.*?</style>', lambda m: ' '.join(re.findall(r'"(?:title|body)":\s*"([^"]*)"', m.group(0))), h, flags=re.S)
    return '\n'.join(' '.join(l.split()) for l in H.unescape(re.sub(r'<[^>]+>', '\n', h)).split('\n') if l.strip())

PAGES = {6: ('unit3_m1_lens5', U6), 7: ('unit3_m1_lens6', U7), 8: ('unit3_m2_lens7', U8), 9: ('unit3_m2_lens8', U9)}
for u, (base, table) in PAGES.items():
    for s in 'pf':
        name = f'{base}_{s}.html'; raw = open(os.path.join(src, name), 'rb').read(); crlf = b'\r\n' in raw
        t = raw.decode('utf-8').replace('\r\n', '\n'); orig = t
        if u == 6:   # part 1.2 stays word for word, except the sentence that names the pair
            a = t.index('class="acc-t">1.2'); b = t.index('class="acc-t">1.3')
            keep = t[a:b]; keep2 = apply(keep, U6[:1], f'U{u}')
            t = apply(t[:a], table[1:], f'U{u}') + keep2 + apply(t[b:], table[1:], f'U{u}')
        else:
            t = apply(t, table, f'U{u}')
        if t != orig:
            open(os.path.join(out, name), 'wb').write((t.replace('\n', '\r\n') if crlf else t).encode('utf-8'))
            for x, y in sentences_diff(vis(orig), vis(t)): LOG.append(dict(unit=u, file='participant page' if s == 'p' else 'facilitator page', where='', before=x, after=y))

def deck_apply(prs, table, u, protect_slides=(), protect_prefix=()):
    for i, sl in enumerate(prs.slides, 1):
        for part, root in (('slide', sl.shapes._spTree), ('notes', sl.notes_slide.notes_text_frame._txBody if sl.has_notes_slide else None)):
            if root is None: continue
            if part == 'slide' and i in protect_slides: continue
            if part == 'notes' and i in protect_slides: continue
            for p in root.iter(qn('a:p')):
                runs = [r for r in p.iter(qn('a:t'))]
                full = ''.join(r.text or '' for r in runs)
                if not full or any(full.startswith(x) for x in protect_prefix): continue
                if full in DECK_EXACT.get(u, {}) and part == 'slide':
                    runs[0].text = DECK_EXACT[u][full]
                    for r in runs[1:]: r.text = ''
                    LOG.append(dict(unit=u, file='deck', where=f'slide {i} {part}', before=full, after=DECK_EXACT[u][full])); continue
                new = apply(full, table, f'U{u}')
                if new == full: continue
                # per-run first
                trial = [apply(r.text or '', table, f'_U{u}') for r in runs]
                if ''.join(trial) == new:
                    for r, x in zip(runs, trial): r.text = x
                else:
                    runs[0].text = new
                    for r in runs[1:]: r.text = ''
                LOG.append(dict(unit=u, file='deck', where=f'slide {i} {part}', before=full, after=new))
DECKS = {6: (U6[1:], (7, 8, 9), ('1.2 Reflection',)), 7: (U7, (), ()), 8: (U8, (), ()), 9: (U9, (), ())}
for u, (table, prot, pref) in DECKS.items():
    f = glob.glob(os.path.join(src, 'Unit decks', f'Unit {u:02d}*.pptx'))[0]
    prs = Presentation(f); n0 = len(LOG); deck_apply(prs, table, u, prot, pref)
    if len(LOG) > n0: prs.save(os.path.join(out, os.path.basename(f)))
# Capstone
raw = open(os.path.join(src, 'capstone_P.html'), 'rb').read(); crlf = b'\r\n' in raw
t = raw.decode('utf-8').replace('\r\n', '\n'); o = t; t = apply(t, CAPSTONE, 'CAP')
if t != o:
    open(os.path.join(out, 'capstone_P.html'), 'wb').write((t.replace('\n', '\r\n') if crlf else t).encode('utf-8'))
    for x, y in sentences_diff(o, t): LOG.append(dict(unit=0, file='Capstone page', where='box description', before=x[-260:], after=y[-260:]))
HITS = {k: v for k, v in HITS.items() if not k[0].startswith('_')}
used = {k[1] for k in HITS}
unused = [(n, a) for n, tb in (('U6', U6), ('U7', U7), ('U8', U8), ('U9', U9), ('CAP', CAPSTONE)) for a, b in tb if a not in used]
json.dump(dict(log=LOG, unused=unused), open(os.path.join(out, 'vocab_log.json'), 'w'), ensure_ascii=False, indent=1)
print('changes logged:', len(LOG)); print('unused pairs:', len(unused)); [print('  ', x) for x in unused]
