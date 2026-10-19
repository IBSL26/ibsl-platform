# -*- coding: utf-8 -*-
"""Reflection pictures: Unit 5 slide 31 (4.2), Unit 6 slide 15 (1.2-1.4) replaced from today's participant pages;
Unit 7 gets Section 1 and Section 2 Reflections slides (copies of its own Section 3 Reflections slide).
Usage: python3 deck_pics.py <decks folder (read+write)> <pics> <pages>"""
import sys, os, glob, copy, io, re, html
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
D, PICS, PAGES = sys.argv[1:4]
def deck(u): return glob.glob(os.path.join(D, f'Unit {u:02d}*.pptx'))[0]
def img(name):
    im = Image.open(os.path.join(PICS, name)); im = im.crop((0, 0, im.width, im.height - 12))
    b = io.BytesIO(); im.save(b, 'PNG'); b.seek(0); return b, im.width / im.height
def place(slide, pic_shape, oval, picname, top=None):
    """Put a new picture where pic_shape sits (same left and width), height from the picture; centre the oval on it."""
    b, ar = img(picname)
    left, width = pic_shape.left, pic_shape.width; top = pic_shape.top if top is None else top
    h = int(width / ar)
    new = slide.shapes.add_picture(b, left, top, width, h)
    new.name = pic_shape.name
    pic_shape._element.addprevious(new._element)
    pic_shape._element.getparent().remove(pic_shape._element)
    oval.top = int(top + h / 2 - oval.height / 2)
    return new, top + h
def pics_ovals(slide):
    ps = [s for s in slide.shapes if s.shape_type == 13]; ov = [s for s in slide.shapes if s.name.startswith('Part ')]
    return sorted(ps, key=lambda s: s.top), sorted(ov, key=lambda s: s.top)

# Unit 5, slide 31: the 4.2 picture
p = Presentation(deck(5)); s = p.slides[30]
ps, ov = pics_ovals(s); pairs = {o.text_frame.text.strip(): (pp, o) for pp, o in zip(ps, ov)}
place(s, *pairs['4.2'], 'u5_ref_4.2.png'); p.save(deck(5)); print('U5 s31 4.2 replaced')
# Unit 6, slide 15: 1.2, 1.3, 1.4 (stacked, gaps kept)
p = Presentation(deck(6)); s = p.slides[14]
ps, ov = pics_ovals(s); y = ps[0].top; gap = ps[1].top - (ps[0].top + ps[0].height)
for pp, o in zip(ps, ov):
    _, bottom = place(s, pp, o, f'u6_ref_{o.text_frame.text.strip()}.png', top=y); y = bottom + gap
p.save(deck(6)); print('U6 s15 replaced')

# Unit 7: new reflection slides from slide 17
pg = open(os.path.join(PAGES, 'unit3_m1_lens6_p.html'), encoding='utf-8').read()
def prompt(rid):
    i = pg.find(f'id="{rid}"'); blk = pg.rfind('class="ref-block"', 0, i); seg = pg[blk:i]
    lab = re.search(r'class="ref-label">(.*?)</div>', seg, re.S).group(1); pr = re.search(r'class="ref-prompt"[^>]*>(.*?)</p>', seg, re.S).group(1)
    clean = lambda x: ' '.join(html.unescape(re.sub('<[^>]+>', '', x)).split())
    return clean(lab).replace('✎ ', ''), clean(pr)
REF = {'1.1': 'ref1', '1.2': 'ref2', '1.3': 'ref3', '1.4': 'ref8', '2.1': 'ref4', '2.2': 'ref5'}
p = Presentation(deck(7)); tmpl = p.slides[16]
def new_slide(title, parts, after_idx, section, bridge, standard):
    s = p.slides.add_slide(tmpl.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for sh in tmpl.shapes: s.shapes._spTree.append(copy.deepcopy(sh._element))
    # the copied picture keeps the template image rel id: copy that relationship too
    for rel in tmpl.part.rels.values():
        if 'image' in rel.reltype: s.part.rels._rels[rel.rId] = rel
    t = [sh for sh in s.shapes if sh.has_text_frame and 'Reflections' in sh.text_frame.text][0]
    t.text_frame.paragraphs[0].runs[0].text = title
    pic = [sh for sh in s.shapes if sh.shape_type == 13][0]; oval = [sh for sh in s.shapes if sh.name.startswith('Part ')][0]
    y = 1463040; gap = 420000
    for k, part in enumerate(parts):
        if k == 0: pp, oo = pic, oval
        else:
            oo_el = copy.deepcopy(oval._element); s.shapes._spTree.append(oo_el); oo = [sh for sh in s.shapes if sh.name.startswith('Part ')][-1]
            pp_el = copy.deepcopy(pic._element) if False else None
            b, ar = img(f'u7_ref_{part}.png'); h = int(pic.width / ar)
            npic = s.shapes.add_picture(b, pic.left, y, pic.width, h); npic.name = 'Reflection ' + part
            oo.text_frame.paragraphs[0].runs[0].text = part; oo.name = 'Part ' + part; oo.top = int(y + h / 2 - oo.height / 2)
            y += h + gap; continue
        oo.text_frame.paragraphs[0].runs[0].text = part; oo.name = 'Part ' + part
        _, bottom = place(s, pp, oo, f'u7_ref_{part}.png', top=y); y = bottom + gap
    lines = [title.upper(), '', 'HOW TO USE THIS SLIDE:', '', '1. Show the slide and read each prompt aloud.']
    for part in parts:
        lab, pr = prompt(REF[part]); lines.append(f'{part} {lab}: "{pr}"')
    lines += ['', '2. Say what the reflections are for.', 'Point learners to the portal section that they will be required to complete with as much depth as possible.',
              'The reflections build the learning portfolio they receive at the end of the programme.', '',
              'Say: "You write these reflections on the portal after the teaching, with as much depth as you can."', '',
              '3. Set the standard.', f'Say: "{standard}"', '', f'4. {bridge}', '', 'ON THE PORTAL, AFTER THE TEACHING:',
              f'Participants write these reflections in Section {section} of the participant file, at the foot of parts ' + ' and '.join(parts) + '.']
    ns = s.notes_slide; tsp = tmpl.notes_slide.shapes._spTree; nsp = ns.shapes._spTree
    for el in list(nsp)[2:]: nsp.remove(el)
    for el in list(tsp)[2:]: nsp.append(copy.deepcopy(el))
    tf = ns.notes_text_frame; tf.text = lines[0]
    for l in lines[1:]: tf.add_paragraph().text = l
    lst = p.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(after_idx, el)
new_slide('Section 1 Reflections · 1.1 and 1.2', ['1.1', '1.2'], 10, 1, 'Continue with the second reflection slide of Section 1.', 'Name one real expectation and one condition, with the person concerned.')
new_slide('Section 1 Reflections · 1.3 and 1.4', ['1.3', '1.4'], 11, 1, 'Bridge to Section 2.', 'Name the people in your chain, and write the Resolve question as you would ask it.')
new_slide('Section 2 Reflections', ['2.1', '2.2'], 16, 2, 'Bridge to Section 3.', 'Name what your team does first, and the cost of the absent mechanism today.')
p.save(deck(7))
for i, s in enumerate(p.slides, 1):
    t = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame][:1]; print(i, t[0][:50] if t else '')
