# -*- coding: utf-8 -*-
"""Unit 8 deck, Carol's amendments of 13 October 2026. Works on the deck as pushed on 12 October (md5 2718c039…, 23 slides).
Usage: python3 carol_deck_u08_r5.py <deck> <participant page> <facilitator page> <pictures folder> <out.pptx>
Only the slides these amendments touch change:
  8  1.4 The Accountability Loop: each recovery achieved through GROW (card line + notes)
  11 2.1 Four consequences: execution velocity reflection added to the end of Carol's notes
  new Section 2 Reflections (2.1, the velocity gap), after slide 11
  13 3.1, 14 3.2, 17 4.1 notes: "the Alignment Brigade" replaced; 17 also its portal line (ref5 reframed)
  18 4.2 role play: the case (alignment falls), roles, notes
  19 4.2 criteria slide becomes "Analysing the conversation with GROW"
  20, 21 Section 4 Reflections: ref5 reframed, the 4.2 Portfolio work reworked
Every other slide stays exactly as saved."""
import sys, re, copy, html, json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image
import u8_content as C
deck, P, F, PICS, out = sys.argv[1:6]
exec(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'r5', 'helpers.py'), encoding='utf-8').read())
vp, vf = plain(P), plain(F)
prs = Presentation(deck)
S = list(prs.slides); assert len(S) == 23, len(S)
for i, a in {7: '1.4 · The Accountability Loop', 8: 'Section 1 Reflections', 10: '2.1 ·', 12: '3.1 ·', 13: '3.2 ·', 16: '4.1 ·', 18: '4.2 · Evaluation criteria', 19: 'Section 4 Reflections', 20: 'Section 4 Reflections'}.items():
    assert first(S[i]).startswith(a), (i, first(S[i]))
before = {i: (etree.tostring(s._element), etree.tostring(s.notes_slide._element)) for i, s in enumerate(S)}
REFS = json.load(open(os.path.join(PICS, 'reflections.json'), encoding='utf-8'))
def rq(part, label):
    return '%s Reflection — %s: "%s"' % (part[:3], label, REFS[part]['prompt'])

def refill(slide, title, items):
    """Clear a reflection slide (title kept) and stack the pictures with gold ovals, as in the other reflection slides."""
    drop_shapes(slide, [sh.name for sh in slide.shapes if sh.name != 'Text 0'])
    set_text(texts(slide)[0], title)
    left, right, top0, bottom = Inches(1.85), Inches(12.85), Inches(1.6), Inches(7.2)
    gap = Inches(0.22); width = right - left; hs = []
    for _, f in items:
        w, h = Image.open(os.path.join(PICS, f)).size; hs.append(width * h / w)
    scale = min(1, (bottom - top0 - gap * (len(items) - 1)) / sum(hs)); print(title, 'picture scale', round(scale, 2))
    y = top0
    for (lab, f), h in zip(items, hs):
        pw = int(width * scale); ph = int(h * scale)
        pic = slide.shapes.add_picture(os.path.join(PICS, f), left, int(y), width=pw); pic.name = 'Reflection ' + lab
        oval(slide, lab, Inches(0.6), int(y + ph / 2 - Inches(0.475)))
        y += ph + gap
    assert y - gap <= bottom + 10

# ═══ slide 8 · 1.4 the Accountability Loop: recovery through GROW ═══════════════════════════════════
s8 = S[7]
for name in ('Text 4', 'Text 8', 'Text 12'):
    sh = [x for x in s8.shapes if x.name == name][0]
    p0 = sh.text_frame.paragraphs[0]._p; p1 = copy.deepcopy(p0); p0.addnext(p1)
    r = p1.find(qn('a:r')); r.find(qn('a:t')).text = 'Achieved through GROW.'
    rp = r.find(qn('a:rPr')); rp.set('b', '1'); rp.find(qn('a:solidFill')).find(qn('a:srgbClr')).set('val', '1A5C2C')
    sp = etree.SubElement(p1.find(qn('a:pPr')), qn('a:spcBef')); etree.SubElement(sp, qn('a:spcPts')).set('val', '1200')
    p1.find(qn('a:pPr')).remove(sp); p1.find(qn('a:pPr')).insert(0, sp)
CA, GR = C.CAUSES14, C.GROW14
def growline(i): return 'Achieving it through GROW. ' + ' '.join('%s: %s' % (n, on(x)) for n, x in zip(C.GROW_NAMES, GR[i]))
set_notes(s8, ['§1.4 — THE ACCOUNTABILITY LOOP', '', '§HOW TO TEACH IT:', '',
  '§1. Name the default response.',
  'Say: "%s"' % on('Many organisations respond to performance gaps with one default response: they push harder. More urgency is applied. More pressure is introduced. Deadlines are reinforced. Pressure alone rarely solves execution challenges because pressure does not identify the underlying cause of the problem.'),
  'Say: "%s"' % on('When a gap is identified, the leader activates the Accountability Loop: a structured diagnostic process designed to restore traction and ownership without collapsing into unproductive performance pressure. Most performance gaps emerge from one of three underlying conditions, and each calls for its own recovery.'), '',
  '§2. Name how the recovery is achieved.',
  'Say: "%s"' % on('The leader achieves the recovery through the GROW model participants used for the Alignment Dialogue in Unit 7: Goal, Reality, Options and Will. The goal of the conversation is ownership of the KPI: an alignment score of 9 or 10.', 'f').replace('participants used', 'you used'), '',
  '§3. Teach the three root causes, one card at a time.',
  'Lack of accountability. Say: "%s Recovery: %s"' % (on(CA[0][2]), on(CA[0][3])), growline(0), '',
  'Lack of skill. Say: "%s Recovery: %s"' % (on(CA[1][2]), on(CA[1][3])), growline(1), '',
  'Misfit or structural constraint. Say: "%s Recovery: %s"' % (on(CA[2][2]), on(CA[2][3])), growline(2), '',
  'Then say: "Every recovery closes with the alignment question asked again: on a scale of 1 to 10, where is your alignment now?"',
  'Then say: "Leaders must distinguish between team members who cannot perform and those who will not. Inability needs development, and unwillingness needs accountability." (Journal, 7 January 2026, Professional Insights: Leadership)',
  'With the third card, say: "Avoid assuming that every negative outcome proves individual failure, or that every positive outcome proves sound practice." (Journal, 18 August 2026, Professional Insights: Strategy)', '',
  '§4. Keep the two conversations separate.',
  'Say: "%s"' % on('The performance conversation must remain separate from the role conversation. Performance discussions focus on progress toward outcomes. Role alignment discussions focus on placing individuals where their capabilities can generate the greatest organisational value. By maintaining this distinction, the organisation protects both fairness and clarity.'),
  'Say: "Accountability works when it is specific and behaviour-based. Vague, general assessments of character produce no specific correction." (Journal, 8 April 2026, Professional Insights: Culture)', '',
  '§5. Ask, now that the loop is taught.',
  'Ask: "Think of a performance gap in your area. Which of the three root causes sits behind it, and what evidence tells you so?" Take two or three answers in the chat.', '',
  '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants write the reflection at the foot of part 1.4: the stage of the Review–Realign–Restore Cycle a current gap sits in, and the root cause they would most likely diagnose.'])

# ═══ slide 11 · 2.1: Carol's notes stay; the reflection is added at their end ═══════════════════════
s11 = S[10]
assert 'velocity of the flight' in s11.notes_slide.notes_text_frame.text and 'ON THE PORTAL' not in s11.notes_slide.notes_text_frame.text
body = s11.notes_slide.notes_text_frame._txBody
tmp = prs.slides[0].notes_slide  # only to borrow set_notes' paragraph format
def add_lines(slide, lines):
    b = slide.notes_slide.notes_text_frame._txBody
    for line in lines:
        p = etree.SubElement(b, '{%s}p' % A); bold = line.startswith('§'); text = line[1:] if bold else line
        if not text.strip():
            e = etree.SubElement(p, '{%s}endParaRPr' % A); e.set('lang', 'en-GB'); e.set('dirty', '0'); continue
        r = etree.SubElement(p, '{%s}r' % A); rp = etree.SubElement(r, '{%s}rPr' % A); rp.set('lang', 'en-GB')
        if bold: rp.set('b', '1')
        rp.set('dirty', '0'); etree.SubElement(r, '{%s}t' % A).text = text
add_lines(s11, ['', '§ON THE PORTAL, AFTER THE TEACHING:',
  'Participants write the reflection at the foot of part 2.1: where they have seen the velocity gap in their own setting, the point where it opened (detecting the deviation, understanding its cause, making the decision or translating it into action), the consequence that followed and what the delay cost.'])

# ═══ new · Section 2 Reflections (a copy of the Section 1 Reflections slide) ═════════════════════════
R2 = duplicate(S[8])
refill(R2, 'Section 2 Reflections', [('2.1', 'ref_2.1.png')])
set_notes(R2, ['§SECTION 2 REFLECTIONS', '', '§HOW TO USE THIS SLIDE:', '',
  '§1. Show the slide and read the prompt aloud.', rq('2.1', 'Where You Have Seen the Velocity Gap'), '',
  '§2. Say what the reflection is for.', 'Say: "Execution velocity is the rate at which an organisation converts strategic intent into measurable results. In this reflection, you name one place in your own setting where the pace fell behind the plan, and the point where the gap opened."', '',
  '§3. Set the standard.', 'Say: "Name a real priority, the real point of delay and what it cost."', '',
  '§4. Bridge to Section 3.', 'Say: "You have seen what weak review discipline costs. Section 3 shows where the review of performance breaks down across functions."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', 'Participants write this reflection in Section 2 of the participant file, at the foot of part 2.1.'])

# ═══ slides 13, 14, 17 · "the Alignment Brigade" ═════════════════════════════════════════════════════
CI = 'the collective intelligence for execution discipline'
note_replace(S[12], [('The Alignment Brigade closes this gap by introducing', 'The leadership team, acting as %s, closes this gap by introducing' % CI),
                     ('Closing this gap requires the Alignment Brigade to reinforce', 'Closing this gap requires the leadership team, acting as %s, to reinforce' % CI),
                     ('The Alignment Brigade protects the rhythm', 'The leadership team, acting as %s, protects the rhythm' % CI)])
note_replace(S[13], [('No single role opens a fracture alone. The Alignment Brigade closes', 'No single role opens a fracture alone. The leadership team, acting as %s, closes' % CI)])
note_replace(S[16], [('allows the Alignment Brigade to observe', 'allows the leadership team, acting as %s, to observe' % CI),
                     ('and the Alignment Brigade can act on it', 'and the leadership team, acting as %s, can act on it' % CI),
                     ('Participants write their own three signals at the foot of each tab of part 4.1, from their own Key Results and KPIs.',
                      'Participants write their early warning signal and their execution learning at the foot of the second and third tabs of part 4.1, from their own Key Results and KPIs. At the foot of part 4.1, they name the conversation type most difficult for them and their team.')])

# ═══ slide 18 · 4.2 the role play ════════════════════════════════════════════════════════════════════
rp = S[17]
set_text(find(rp, 'Complaint resolution'), 'Complaint resolution has stalled at 36 hours against a 12-hour target, and alignment has fallen from 9.')
set_text(find(rp, 'Three roles'), 'CXO · Team member · Observer. Role play the conversation using GROW.')
sc = C.SCENARIO2.split('</p>')
set_notes(rp, ['§4.2 — THE ACCOUNTABILITY LOOP ROLE PLAY (PRACTICE IN CLASS)', '', '§HOW TO TEACH IT:', '',
  '§1. Give the case.', 'Say: "' + on(sc[0], 'f') + ' ' + on(sc[1], 'f') + '"', '',
  '§2. Assign the three roles.'] +
  ['%s: "%s"' % (a.replace(' Role', ''), on(b, 'f')) for a, b in C.ROLES42] +
  ['Give the team member this brief privately: "%s"' % on(C.TM_BRIEF, 'f'),
   'The role play case is built on a lack of skill. The Portfolio case participants analyse after class is built on a different root cause.',
   'Brief the CXO with this line: "The leader who receives honest feedback about a genuine problem and responds with honesty before justification, and action before communication, creates the conditions for real leadership." (Journal, 29 June 2026, Professional Insights: Leadership)', '',
  '§3. Run the conversation using GROW.', 'The CXO opens with the evidence and asks the alignment question, then works through Goal, Reality, Options and Will with the team member.',
  'If the CXO defaults to pressure, pause and redirect: "You have identified the gap. Now ask about the reality behind it before you move to options."', '',
  '§4. Close the conversation.', 'The CXO closes with the actions, the follow-up and the alignment question asked again, aiming close to 10.',
  'Say: "Accountability without a pathway back becomes punishment. Correct firmly where necessary, and deliberately restore trust, contribution and dignity when change is evident." (Journal, 16 September 2026, Professional Insights: Leadership)', '',
  '§5. Debrief.', 'The observer posts in the chat how the CXO used each stage of GROW and the root cause that surfaced.',
  'Ask the CXO: "If you ran this conversation again, what would you change first?"', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', 'After class, participants complete personal Portfolio work in part 4.2 of their own file. They read a scripted conversation on a different case, in which the root cause is never named. They deduce which of the three root causes is evident and the lines that rule out the other two, write how they would use GROW to resolve it and move the alignment close to 10, and judge how the CXO applied the Accountability Loop against the GROW model.'], watch_ok=True)

# ═══ slide 19 · the criteria slide becomes the GROW slide ═════════════════════════════════════════════
gs = S[18]
set_text(texts(gs)[0], '4.2 · Analysing the conversation with GROW')
drop_shapes(gs, ['Image 0', 'Image 2', 'Image 3'])
d = {sh.name: sh for sh in gs.shapes}
c4 = copy.deepcopy(d['Shape 7']._element); t4 = copy.deepcopy(d['Text 8']._element)
nid = max(int(e.get('id')) for e in gs._element.iter(qn('p:cNvPr')))
for k, el in enumerate((c4, t4)): el.find('.//' + qn('p:cNvPr')).set('id', str(nid + 1 + k))
c4.find('.//' + qn('p:cNvPr')).set('name', 'Shape 9'); t4.find('.//' + qn('p:cNvPr')).set('name', 'Text 10')
gs.shapes._spTree.append(c4); gs.shapes._spTree.append(t4)
d = {sh.name: sh for sh in gs.shapes}
ROWS = [('G', 'Goal  ', 'Ownership again: an alignment of 9 or 10.'), ('R', 'Reality  ', 'What stands behind the score and the gap.'),
        ('O', 'Options  ', 'The ways forward the owner generates.'), ('W', 'Will  ', 'Actions, owner and follow-up; the alignment question again.')]
top, step, th = Emu(1463040), Emu(1371600), Emu(1097280)
for k, ((circ, txt), (L, a, b)) in enumerate(zip((('Shape 1', 'Text 2'), ('Shape 5', 'Text 6'), ('Shape 7', 'Text 8'), ('Shape 9', 'Text 10')), ROWS)):
    T = d[txt]; T.top = top + step * k; T.height = th; set_text(T, a, b)
    Ci = d[circ]; Ci.top = top + step * k + (th - Ci.height) // 2
    tb = Ci._element.find(qn('p:txBody'))
    for p_ in tb.findall(qn('a:p')): tb.remove(p_)
    tb.find(qn('a:bodyPr')).attrib.update({'lIns': '0', 'tIns': '0', 'rIns': '0', 'bIns': '0', 'anchor': 'ctr'})
    p_ = etree.SubElement(tb, qn('a:p')); etree.SubElement(p_, qn('a:pPr')).set('algn', 'ctr')
    r_ = etree.SubElement(p_, qn('a:r')); rp_ = etree.SubElement(r_, qn('a:rPr')); rp_.attrib.update({'lang': 'en-GB', 'sz': '2400', 'b': '1'})
    f_ = etree.SubElement(rp_, qn('a:solidFill')); etree.SubElement(f_, qn('a:srgbClr')).set('val', 'FFFFFF')
    etree.SubElement(rp_, qn('a:latin')).set('typeface', 'Cambria'); etree.SubElement(r_, qn('a:t')).text = L
assert T.top + T.height <= Emu(6858000)
G = C.GROW42
set_notes(gs, ['§4.2 — ANALYSING THE CONVERSATION WITH GROW', '', '§HOW TO USE THIS SLIDE:', '',
  '§1. Recall GROW from Unit 7.', 'Say: "%s"' % on('The GROW model of Unit 7 is the basis for analysing the conversation. Each stage asks a question, and the individual does the thinking, so the answers reached are their own.'), '',
  '§2. Teach each stage as it works in the Accountability Loop.'] +
  ['%s · %s: "%s" Question: %s' % (g[0], g[1], on(g[2]), on(g[3])) for g in G] + ['',
  '§3. Brief the observer before the role play.', 'Say: "Note how the CXO used each stage of GROW: did the CXO agree the goal, explore the reality, let the owner generate the options, and close with the will and the alignment question asked again?"', '',
  '§4. Use GROW in the debrief.', 'The observer posts how the CXO used each stage. Ask the room which stage was hardest to hold, and why.',
  'Say: "Leaders who either conceal failure or weaponise it prevent learning. Mature accountability names what went wrong, understands why it happened and moves people toward responsible change." (Journal, 10 August 2026, Professional Insights: Leadership)',
  'Say: "Holding people accountable for outcomes without considering intent produces a culture of blame. Intent matters." (Journal, 17 March 2026, Professional Insights)', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', 'After class, in part 4.2 of their own file, participants deduce the root cause in a scripted conversation, use GROW to plan its resolution, and judge how the CXO applied the Accountability Loop.'], watch_ok=True)

# ═══ slides 20, 21 · Section 4 Reflections ════════════════════════════════════════════════════════════
refill(S[19], 'Section 4 Reflections', [('4.1', 'ref_4.1b.png'), ('4.1', 'ref_4.1c.png')])
set_notes(S[19], ['§SECTION 4 REFLECTIONS (1 OF 2)', '', '§HOW TO USE THIS SLIDE:', '',
  '§1. Show the slide and read the two prompts aloud.', rq('4.1b', 'Your Early Warning Signal'), rq('4.1c', 'Your Execution Learning'), '',
  '§2. Say what the reflections are for.', 'Say: "In part 4.1, you write your own EPIC signals, from your own Key Results and KPIs, with the evidence behind each. These reflections build your learning portfolio."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', 'Participants write their early warning signal and their execution learning at the foot of the second and third tabs of part 4.1.'])
refill(S[20], 'Section 4 Reflections', [('4.1', 'ref_4.1a.png'), ('4.2', 'ref_4.2.png')])
set_notes(S[20], ['§SECTION 4 REFLECTIONS (2 OF 2)', '', '§HOW TO USE THIS SLIDE:', '',
  '§1. Show the slide and read the prompt and the Portfolio work aloud.', rq('4.1a', 'The Conversation Type Most Difficult for Your Team'), '4.2 Portfolio work: the Accountability Loop case, read and analysed alone.', '',
  '§2. Say what the work is for.', 'Say: "In 4.2, you read a full Accountability Loop conversation on a different case from the role play. The conversation never names the root cause: you deduce which of the three it is, write how you would use GROW to resolve it and move the alignment close to 10, and judge how the CXO applied the loop against GROW. This is personal work, and it builds your learning portfolio."', '',
  '§3. Bridge to Section 5.', 'Say: "You have seen the enterprise view and practised the loop on one Key Result. Section 5 brings back the four consequences of weak review discipline, and your group guards its own strategy against them."', '',
  '§ON THE PORTAL, AFTER THE TEACHING:', 'Participants write the reflection at the foot of part 4.1, and complete the Portfolio work of part 4.2.'])

# ═══ order: the new Section 2 Reflections slide goes after slide 11 ═══════════════════════════════════
lst = prs.slides._sldIdLst; ids = list(lst); new = ids[-1]; lst.remove(new); ids[10].addnext(new)
# every slide not named above is untouched
CHANGED = {7, 10, 12, 13, 16, 17, 18, 19, 20}
for i, s in enumerate(S):
    if i not in CHANGED: assert (etree.tostring(s._element), etree.tostring(s.notes_slide._element)) == before[i], 'slide %d changed' % (i + 1)
allnotes = '\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides)
for bad in ('Brigade', 'criterion', 'Criterion', 'three criteria', 'Your Progress Signal'): assert bad not in allnotes, bad
prs.save(out); print(len(prs.slides), 'slides ->', out)
