import sys
from pptx import Presentation
p=Presentation(sys.argv[1]); out=[]
for i,s in enumerate(p.slides,1):
    out.append(f'\n################ SLIDE {i}')
    for para in s.notes_slide.notes_text_frame.paragraphs:
        t=''.join(r.text for r in para.runs).replace('\x0b','\n')
        b=bool(para.runs) and all(r.font.bold for r in para.runs)
        out.append(('**'+t+'**') if b else t)
open(sys.argv[2],'w',encoding='utf8').write('\n'.join(out))
