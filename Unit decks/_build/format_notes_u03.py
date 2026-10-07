"""Set the presenter notes of the Unit 3 deck in the form of Carol's Unit 2 deck.

Usage: python3 format_notes_u03.py <deck_in.pptx> <deck_out.pptx>

On every slide the notes are rebuilt as one paragraph for each line. A line that starts with § is a heading or a
step title: the mark is removed and the line is set in bold. Slides are not touched. Run after add_reflection_slides.py.
"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

def paragraph(line):
    p = etree.Element(f"{{{A}}}p")
    bold = line.startswith("§")
    text = line[1:] if bold else line
    if not text.strip():
        e = etree.SubElement(p, f"{{{A}}}endParaRPr"); e.set("lang", "en-GB"); e.set("dirty", "0")
        return p
    r = etree.SubElement(p, f"{{{A}}}r")
    rpr = etree.SubElement(r, f"{{{A}}}rPr"); rpr.set("lang", "en-GB")
    if bold:
        rpr.set("b", "1")
    rpr.set("dirty", "0")
    etree.SubElement(r, f"{{{A}}}t").text = text
    return p

def main(src, out):
    prs = Presentation(src)
    done = 0
    for i, s in enumerate(prs.slides, 1):
        text = s.notes_slide.notes_text_frame.text.replace("\x0b", "\n")
        if not text.split("\n")[0].startswith("§"):
            sys.exit(f"slide {i}: notes do not start with a heading")
        body = s.notes_slide.notes_text_frame._txBody
        for p in body.findall(qn("a:p")):
            body.remove(p)
        for line in text.split("\n"):
            body.append(paragraph(line.rstrip()))
        done += 1
    left = sum(1 for s in prs.slides if "§" in s.notes_slide.notes_text_frame.text)
    if left:
        sys.exit(f"{left} slides still carry the heading mark")
    prs.save(out)
    print(f"notes formatted on {done} slides -> {out}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
