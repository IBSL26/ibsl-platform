import sys,re
from pptx import Presentation
src,out=sys.argv[1],sys.argv[2]
p=Presentation(src)
HEAD=re.compile(r"^[A-Z0-9][A-Z0-9 ·,&—’'():®./\-–?]+:?$")
STEP=re.compile(r'^\d+\.\s')
for s in p.slides:
    tf=s.notes_slide.notes_text_frame
    text=tf.text
    lines=text.split('\n')
    # rebuild paragraphs
    tf.clear()
    first=True
    for l in lines:
        para=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        r=para.add_run(); r.text=l
        if l and (HEAD.match(l) or STEP.match(l)): r.font.bold=True
p.save(out)
